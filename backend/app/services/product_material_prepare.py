from __future__ import annotations

import asyncio
import csv
import hashlib
import json
import mimetypes
import re
from datetime import datetime
from html.parser import HTMLParser
from pathlib import Path
from typing import Any
from urllib.parse import quote_plus

from openpyxl import load_workbook
from PIL import Image, ImageOps
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.config import settings
from app.models import Product, ProductMaterialAsset
from app.pipeline.chrome_ctrl import chrome_execute_js, chrome_navigate, chrome_workflow
from app.pipeline.step1_collect import RAW_ASSETS_DIR, _download_material_zips
from app.services.product_pipeline_artifacts import register_pipeline_artifact
from app.services.upc_pool import refresh_upc_binding


IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".webp", ".gif", ".bmp", ".tif", ".tiff"}
VIDEO_EXTENSIONS = {".mp4", ".mov", ".m4v", ".avi", ".mkv", ".webm", ".wmv", ".mpeg", ".mpg"}
SPREADSHEET_EXTENSIONS = {".xlsx", ".xls", ".csv", ".tsv"}
HTML_EXTENSIONS = {".html", ".htm"}
PDF_EXTENSIONS = {".pdf"}
TEXT_EXTENSIONS = {".txt", ".md", ".json", ".xml"}
GIGAB2B_SEARCH_URL = (
    "https://www.gigab2b.com/index.php?route=product/search"
    "&search={item_code}&search_source=0&search_dimension=1"
)
GIGAB2B_DETAIL_URL = "https://www.gigab2b.com/index.php?route=product/product&product_id={product_id}"


class ProductMaterialPrepareError(RuntimeError):
    """Raised when supplier materials cannot be prepared safely."""


class _PlainTextHTMLParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []

    def handle_data(self, data: str) -> None:
        text = " ".join(str(data or "").split())
        if text:
            self.parts.append(text)


def _json_dumps(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True)


def _json_loads(value: str | None, fallback: Any) -> Any:
    if not value:
        return fallback
    try:
        return json.loads(value)
    except Exception:
        return fallback


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def classify_material_asset(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".zip":
        return "zip"
    if suffix in IMAGE_EXTENSIONS:
        return "image"
    if suffix in VIDEO_EXTENSIONS:
        return "video"
    if suffix in SPREADSHEET_EXTENSIONS:
        return "spreadsheet"
    if suffix in HTML_EXTENSIONS:
        return "html"
    if suffix in PDF_EXTENSIONS:
        return "pdf"
    if suffix in TEXT_EXTENSIONS:
        return "text"
    return "other"


def _append_usage(asset: ProductMaterialAsset, *values: str) -> None:
    usage = _json_loads(asset.downstream_usage_json, [])
    if not isinstance(usage, list):
        usage = []
    for value in values:
        if value and value not in usage:
            usage.append(value)
    asset.downstream_usage_json = _json_dumps(usage)


def _compact_cell(value: Any, limit: int = 300) -> str:
    text = " ".join(str(value if value is not None else "").split())
    return text if len(text) <= limit else text[: limit - 3].rstrip() + "..."


def _spreadsheet_preview(path: Path) -> dict[str, Any]:
    rows: list[list[str]] = []
    sheet_name: str | None = None
    if path.suffix.lower() == ".xlsx":
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            worksheet = workbook.worksheets[0]
            sheet_name = worksheet.title
            for row in worksheet.iter_rows(max_row=30, max_col=12, values_only=True):
                values = [_compact_cell(value) for value in row]
                if any(values):
                    rows.append(values)
        finally:
            workbook.close()
    elif path.suffix.lower() in {".csv", ".tsv"}:
        delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
        with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as handle:
            reader = csv.reader(handle, delimiter=delimiter)
            for index, row in enumerate(reader):
                if index >= 30:
                    break
                values = [_compact_cell(value) for value in row[:12]]
                if any(values):
                    rows.append(values)
    else:
        raise ValueError(f"暂不支持提取旧版表格事实: {path.suffix}")
    return {"sheet_name": sheet_name, "rows": rows}


def _text_preview(path: Path, asset_kind: str) -> str:
    text = path.read_text(encoding="utf-8", errors="replace")[:12000]
    if asset_kind == "html":
        parser = _PlainTextHTMLParser()
        parser.feed(text)
        text = "\n".join(parser.parts)
    compact = "\n".join(line.strip() for line in text.splitlines() if line.strip())
    return compact[:8000]


def write_information_material_facts(
    assets: list[ProductMaterialAsset],
    output_path: Path,
) -> dict[str, Any]:
    sources: list[dict[str, Any]] = []
    for asset in assets:
        if asset.package_type != "information" or asset.processing_status != "ready":
            continue
        path = Path(asset.path).expanduser()
        if not path.is_file() or asset.asset_kind == "zip":
            continue
        source: dict[str, Any] = {
            "asset_id": asset.id,
            "relative_path": asset.relative_path,
            "filename": asset.original_filename,
            "asset_kind": asset.asset_kind,
            "content_hash": asset.content_hash,
        }
        try:
            if asset.asset_kind == "spreadsheet":
                preview = _spreadsheet_preview(path)
                source.update({"extraction_status": "extracted", "table_preview": preview})
                _append_usage(asset, "product_detail_preview", "product_fact_source")
            elif asset.asset_kind in {"html", "text"}:
                preview_text = _text_preview(path, asset.asset_kind)
                source.update({
                    "extraction_status": "extracted" if preview_text else "empty",
                    "text_preview": preview_text,
                })
                _append_usage(asset, "product_detail_preview", "product_fact_source" if preview_text else "preview_only")
            elif asset.asset_kind in {"pdf", "video", "image"}:
                source["extraction_status"] = "preview_only"
                _append_usage(asset, "product_detail_preview", "preview_only")
            else:
                source["extraction_status"] = "unsupported"
                source["reason"] = "当前没有安全的结构化事实提取器"
                _append_usage(asset, "product_detail_preview", "unsupported_for_fact_extraction")
        except Exception as exc:
            source["extraction_status"] = "failed"
            source["reason"] = f"{type(exc).__name__}: {exc}"
            _append_usage(asset, "product_detail_preview", "fact_extraction_failed")
        sources.append(source)

    payload = {
        "schema": "fbm-supplier-material-facts.v1",
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "source_package": "information",
        "source_count": len(sources),
        "sources": sources,
    }
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return payload


def _relative_path(path: Path, root: Path) -> str:
    try:
        return str(path.resolve().relative_to(root.resolve()))
    except ValueError:
        raise ProductMaterialPrepareError(f"素材文件不在商品目录内: {path}")


def _image_dimensions(path: Path) -> tuple[int | None, int | None, str | None]:
    try:
        with Image.open(path) as image:
            image = ImageOps.exif_transpose(image)
            return int(image.width), int(image.height), None
    except Exception as exc:
        return None, None, f"图片无法读取: {type(exc).__name__}: {exc}"


async def resolve_gigab2b_product_page(item_code: str) -> dict[str, str]:
    normalized = str(item_code or "").strip()
    if not normalized:
        raise ProductMaterialPrepareError("缺少 GIGA Item Code")
    search_url = GIGAB2B_SEARCH_URL.format(item_code=quote_plus(normalized))

    async with chrome_workflow(f"giga-material-map:{normalized}"):
        if not await chrome_navigate(search_url, wait=1.0):
            raise ProductMaterialPrepareError(f"Chrome 无法打开 GIGA 搜索页: {search_url}")
        matches: list[dict[str, str]] = []
        card_count = 0
        for attempt in range(20):
            cards_json = await chrome_execute_js(
                r'''(function() {
                    return JSON.stringify(Array.from(document.querySelectorAll('.product-item[data-gmd-attr-product_id]')).map(function(card) {
                        return {
                            product_id: String(card.getAttribute('data-gmd-attr-product_id') || '').trim(),
                            text: String(card.innerText || card.textContent || '').trim()
                        };
                    }));
                })()''',
                timeout=20,
            )
            try:
                cards = json.loads(cards_json or "[]")
            except json.JSONDecodeError as exc:
                raise ProductMaterialPrepareError("GIGA 搜索页商品卡片解析失败") from exc

            matches = []
            card_count = len(cards) if isinstance(cards, list) else 0
            for card in cards if isinstance(cards, list) else []:
                if not isinstance(card, dict):
                    continue
                text = str(card.get("text") or "")
                code_match = re.search(r"Item\s*Code\s*:\s*([A-Za-z0-9_-]+)", text, flags=re.I)
                card_code = code_match.group(1).strip() if code_match else ""
                product_id = str(card.get("product_id") or "").strip()
                if card_code.casefold() == normalized.casefold():
                    matches.append({"item_code": card_code, "product_id": product_id})
            if matches:
                break
            if attempt < 19:
                await asyncio.sleep(1.0)

        if len(matches) != 1:
            raise ProductMaterialPrepareError(
                f"GIGA 搜索结果必须精确匹配 1 个商品，实际 {len(matches)} 个"
                f"（已加载卡片 {card_count} 个）: item_code={normalized}"
            )
        product_id = matches[0]["product_id"]
        if not product_id.isdigit():
            raise ProductMaterialPrepareError(f"GIGA product_id 不是纯数字: {product_id}")

        detail_url = GIGAB2B_DETAIL_URL.format(product_id=product_id)
        if not await chrome_navigate(detail_url, wait=1.0):
            raise ProductMaterialPrepareError(f"Chrome 无法打开 GIGA 商品详情页: {detail_url}")
        detail_codes: list[str] = []
        for attempt in range(20):
            detail_text = await chrome_execute_js(
                "(function(){return String(document.body && document.body.innerText || '');})()",
                timeout=20,
            )
            detail_codes = re.findall(r"Item\s*Code\s*:\s*([A-Za-z0-9_-]+)", detail_text or "", flags=re.I)
            if normalized.casefold() in {code.casefold() for code in detail_codes}:
                break
            if attempt < 19:
                await asyncio.sleep(1.0)
        else:
            raise ProductMaterialPrepareError(
                f"GIGA 详情页 Item Code 二次校验失败: expected={normalized}, product_id={product_id}"
            )

    return {
        "item_code": normalized,
        "product_id": product_id,
        "search_url": search_url,
        "detail_url": detail_url,
    }


async def _upsert_asset(
    db: AsyncSession,
    *,
    product_id: int,
    parent_asset_id: int | None,
    source_task_run_id: int | None,
    test_session_key: str | None,
    package_type: str | None,
    asset_kind: str,
    path: Path,
    material_dir: Path,
    metadata: dict[str, Any] | None = None,
) -> ProductMaterialAsset:
    content_hash = sha256_file(path)
    result = await db.execute(
        select(ProductMaterialAsset).where(
            ProductMaterialAsset.product_id == product_id,
            ProductMaterialAsset.content_hash == content_hash,
            ProductMaterialAsset.asset_kind == asset_kind,
        )
    )
    asset = result.scalar_one_or_none()
    now = datetime.now()
    previous_metadata = _json_loads(asset.metadata_json, {}) if asset else {}
    sources = previous_metadata.get("sources") if isinstance(previous_metadata.get("sources"), list) else []
    source_entry = {
        "path": str(path),
        "package_type": package_type,
        "parent_asset_id": parent_asset_id,
        "source_task_run_id": source_task_run_id,
    }
    if source_entry not in sources:
        sources.append(source_entry)
    combined_metadata = {**previous_metadata, **(metadata or {}), "sources": sources}

    if asset is None:
        asset = ProductMaterialAsset(
            product_id=product_id,
            content_hash=content_hash,
            asset_kind=asset_kind,
            original_filename=path.name,
            path=str(path.resolve()),
            created_at=now,
        )
        db.add(asset)
    asset.parent_asset_id = asset.parent_asset_id or parent_asset_id
    asset.source_task_run_id = source_task_run_id
    asset.test_session_key = test_session_key
    if not asset.package_type or asset.package_type == "unknown":
        asset.package_type = package_type
    asset.path = str(path.resolve())
    asset.relative_path = _relative_path(path, material_dir)
    asset.file_size = path.stat().st_size
    asset.mime_type = mimetypes.guess_type(path.name)[0]
    asset.processing_status = "ready"
    asset.rejection_reason = None
    asset.metadata_json = _json_dumps(combined_metadata)
    asset.updated_at = now

    if asset_kind == "image":
        width, height, error = _image_dimensions(path)
        asset.width = width
        asset.height = height
        if error:
            asset.processing_status = "rejected"
            asset.rejection_reason = error
    await db.flush()
    return asset


async def write_material_manifest(db: AsyncSession, product_id: int, material_dir: Path) -> Path:
    result = await db.execute(
        select(ProductMaterialAsset)
        .where(ProductMaterialAsset.product_id == product_id)
        .order_by(ProductMaterialAsset.parent_asset_id.is_(None).desc(), ProductMaterialAsset.id.asc())
    )
    assets = result.scalars().all()
    manifest = {
        "schema": "fbm-product-material-manifest.v1",
        "product_id": product_id,
        "generated_at": datetime.now().isoformat(),
        "asset_count": len(assets),
        "assets": [
            {
                "id": asset.id,
                "parent_asset_id": asset.parent_asset_id,
                "package_type": asset.package_type,
                "asset_kind": asset.asset_kind,
                "original_filename": asset.original_filename,
                "path": asset.path,
                "relative_path": asset.relative_path,
                "content_hash": asset.content_hash,
                "file_size": asset.file_size,
                "mime_type": asset.mime_type,
                "width": asset.width,
                "height": asset.height,
                "processing_status": asset.processing_status,
                "rejection_reason": asset.rejection_reason,
                "downstream_usage": _json_loads(asset.downstream_usage_json, []),
                "contact_sheet_path": asset.contact_sheet_path,
                "contact_sheet_page": asset.contact_sheet_page,
                "contact_sheet_label": asset.contact_sheet_label,
            }
            for asset in assets
        ],
    }
    manifest_path = material_dir / "material_manifest.json"
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    return manifest_path


async def prepare_product_materials(
    db: AsyncSession,
    *,
    product_id: int,
    source_task_run_id: int | None = None,
    test_session_key: str | None = None,
) -> dict[str, Any]:
    result = await db.execute(
        select(Product)
        .where(Product.id == product_id)
        .options(selectinload(Product.data), selectinload(Product.material_assets))
    )
    product = result.scalar_one_or_none()
    if not product or not product.data:
        raise ProductMaterialPrepareError(f"商品或商品资料不存在: {product_id}")
    item_code = str(product.data.item_code or "").strip()
    if not item_code:
        raise ProductMaterialPrepareError("商品缺少 item_code，不能解析 GIGA 商品页")

    page = await resolve_gigab2b_product_page(item_code)
    product.gigab2b_product_id = page["product_id"]
    product.gigab2b_url = page["detail_url"]
    product.pipeline_test_session_key = test_session_key or product.pipeline_test_session_key
    product.pipeline_origin_task_run_id = source_task_run_id or product.pipeline_origin_task_run_id

    material_dir = Path(product.data.material_dir).expanduser() if product.data.material_dir else (
        settings.PRODUCT_BASE_DIR / "GIGA" / (product.source_site or "US").upper() / item_code
    )
    material_dir = material_dir.resolve()
    material_dir.mkdir(parents=True, exist_ok=True)
    raw_dir = material_dir / RAW_ASSETS_DIR
    raw_dir.mkdir(parents=True, exist_ok=True)
    product.data.material_dir = str(material_dir)

    download_results = await _download_material_zips(
        raw_dir,
        {"itemCode": item_code, "_gigab2bProductId": page["product_id"]},
        None,
        page["detail_url"],
        required_options={"To B素材包", "Information"},
        download_all=True,
    )
    package_types = {str(item.get("type") or "unknown") for item in download_results}
    missing_types = [key for key in ("to_b", "information") if key not in package_types]
    if missing_types:
        raise ProductMaterialPrepareError(f"缺少必需素材包: {', '.join(missing_types)}")
    empty_packages = [
        str(item.get("path") or item.get("type") or "unknown")
        for item in download_results
        if int(item.get("extracted_count") or 0) <= 0
    ]
    if empty_packages:
        raise ProductMaterialPrepareError(
            "素材包没有解压出文件: " + ", ".join(empty_packages)
        )

    registered_package_ids: list[int] = []
    touched_asset_ids: set[int] = set()
    for download in download_results:
        zip_path = Path(str(download.get("path") or "")).expanduser().resolve()
        extracted_dir = Path(str(download.get("extracted_dir") or "")).expanduser().resolve()
        package_type = str(download.get("type") or "unknown")
        if not zip_path.is_file() or not extracted_dir.is_dir():
            raise ProductMaterialPrepareError(f"素材包未正确保存或解压: {zip_path}")
        package_asset = await _upsert_asset(
            db,
            product_id=product.id,
            parent_asset_id=None,
            source_task_run_id=source_task_run_id,
            test_session_key=test_session_key,
            package_type=package_type,
            asset_kind="zip",
            path=zip_path,
            material_dir=material_dir,
            metadata={"download": download},
        )
        registered_package_ids.append(package_asset.id)
        touched_asset_ids.add(package_asset.id)
        _append_usage(package_asset, "source_archive", "product_detail_preview")
        for member in sorted(extracted_dir.rglob("*")):
            if not member.is_file():
                continue
            member_asset = await _upsert_asset(
                db,
                product_id=product.id,
                parent_asset_id=package_asset.id,
                source_task_run_id=source_task_run_id,
                test_session_key=test_session_key,
                package_type=package_type,
                asset_kind=classify_material_asset(member),
                path=member,
                material_dir=material_dir,
                metadata={"extracted_from": zip_path.name},
            )
            touched_asset_ids.add(member_asset.id)
            if member_asset.package_type == "to_b" and member_asset.asset_kind == "image":
                _append_usage(member_asset, "product_detail_preview", "auto_image_selection_candidate")
            elif member_asset.asset_kind == "video":
                _append_usage(member_asset, "product_detail_preview", "preview_only")
            else:
                _append_usage(member_asset, "product_detail_preview")

    existing_asset_result = await db.execute(
        select(ProductMaterialAsset).where(ProductMaterialAsset.product_id == product.id)
    )
    existing_assets = existing_asset_result.scalars().all()
    now = datetime.now()
    for asset in existing_assets:
        if asset.id in touched_asset_ids:
            continue
        asset.processing_status = "stale"
        asset.rejection_reason = "本轮素材准备未再次发现此文件"
        asset.updated_at = now

    image_result = await db.execute(
        select(ProductMaterialAsset).where(
            ProductMaterialAsset.product_id == product.id,
            ProductMaterialAsset.asset_kind == "image",
            ProductMaterialAsset.processing_status == "ready",
        )
    )
    image_assets = image_result.scalars().all()
    if not image_assets:
        raise ProductMaterialPrepareError("To B 素材包未解压出可读取图片")
    to_b_image_assets = [asset for asset in image_assets if asset.package_type == "to_b"]
    if not to_b_image_assets:
        raise ProductMaterialPrepareError("To B 素材包已下载，但没有登记到可读取图片")
    product.data.image_count = len(to_b_image_assets)
    material_facts_path = material_dir / "image analysis" / "material_facts.json"
    material_facts = write_information_material_facts(
        existing_assets,
        material_facts_path,
    )
    snapshot = _json_loads(product.data.gigab2b_raw_snapshot, {})
    if not isinstance(snapshot, dict):
        snapshot = {}
    snapshot["gigab2b_page_mapping"] = page
    snapshot["material_prepare"] = {
        "source_task_run_id": source_task_run_id,
        "test_session_key": test_session_key,
        "package_types": sorted(package_types),
        "package_asset_ids": registered_package_ids,
        "image_asset_ids": [asset.id for asset in image_assets],
        "to_b_image_asset_ids": [asset.id for asset in to_b_image_assets],
        "material_facts_path": str(material_facts_path),
    }
    snapshot["material_facts"] = material_facts
    product.data.gigab2b_raw_snapshot = _json_dumps(snapshot)
    await refresh_upc_binding(db, product)
    manifest_path = await write_material_manifest(db, product.id, material_dir)
    await register_pipeline_artifact(
        db,
        product_id=product.id,
        file_type="material_facts_json",
        path=material_facts_path,
        metadata={"source_package": "information", "source_count": material_facts.get("source_count", 0)},
    )
    await register_pipeline_artifact(
        db,
        product_id=product.id,
        file_type="material_manifest_json",
        path=manifest_path,
        metadata={"package_types": sorted(package_types), "asset_count": len(existing_assets)},
    )
    await db.commit()

    return {
        "product_id": product.id,
        "item_code": item_code,
        "gigab2b_product_id": page["product_id"],
        "gigab2b_url": page["detail_url"],
        "package_count": len(download_results),
        "package_types": sorted(package_types),
        "image_count": len(image_assets),
        "to_b_image_count": len(to_b_image_assets),
        "material_facts_path": str(material_facts_path),
        "material_fact_source_count": material_facts.get("source_count", 0),
        "material_dir": str(material_dir),
        "manifest_path": str(manifest_path),
        "test_session_key": test_session_key,
    }
