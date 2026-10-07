#!/usr/bin/env python3
"""Read-only audit of supplier rules against exported rows. Never updates DB/ZIP."""
import argparse
import json
import sqlite3
import sys
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from zipfile import ZipFile

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))
from app.pipeline.amazon_export.attribute_rules import resolve_attributes
from app.pipeline.step10_amazon_template import _load_template_mapping, _template_product_type_for_semantic_fields, _flatten_mapping_values

FIELDS = {
    "BICYCLE": ["material", "included_components", "age_range_description", "tire_type", "import_designation"],
    "BED_FRAME": ["finish_type"],
    "CABINET": ["construction_type", "door_style", "is_fragile", "special_features", "mounting_type"],
    "DRESSER": ["finish_type", "is_fragile"],
    "STORAGE_BOX": ["special_features"],
    "STEP_STOOL": ["maximum_height", "maximum_height_unit", "required_product_compliance_certificate"],
    "MAKEUP_VANITY": ["number_of_drawers", "furniture_finish"],
    "RIDE_ON_TOY": ["material", "theme", "target_audience", "target_gender", "search_terms"],
}


def audit(database, zip_paths):
    db = sqlite3.connect(f"file:{database.resolve()}?mode=ro", uri=True)
    db.row_factory = sqlite3.Row
    records = []
    for zip_path in zip_paths:
        with ZipFile(zip_path) as archive:
            manifest = json.loads(archive.read("export-manifest.json"))
            for row in manifest["rows"]:
                pd = SimpleNamespace(**dict(db.execute("select * from product_data where product_id=?", (row["product_id"],)).fetchone()))
                p = SimpleNamespace(**dict(db.execute("select * from products where id=?", (row["product_id"],)).fetchone()))
                source = db.execute("select payload_json from product_source_payload where product_id=?", (p.id,)).fetchone()
                if source:
                    for key, value in json.loads(source[0]).items():
                        setattr(pd, key, json.dumps(value) if isinstance(value, (dict, list)) else value)
                snapshot = db.execute("select payload_json from product_source_snapshot where product_id=?", (p.id,)).fetchone()
                if snapshot: pd.gigab2b_raw_snapshot = snapshot[0]
                m = _load_template_mapping(p, pd)
                typ = _template_product_type_for_semantic_fields(m, pd)
                decisions = resolve_attributes(pd, m, Path(m["template_path"]), typ)
                w = load_workbook(BytesIO(archive.read(row["workbook"])), read_only=True)
                ws = w["Template"]
                columns = {str(cell.value): cell.column for cell in ws[5] if cell.value}
                old_type = ws.cell(row["row_number"], columns["product_type#1.value"]).value
                record = {"sku": row["seller_sku"], "product_id": p.id, "old_type": old_type,
                          "new_type": typ, "needs_rerouting": old_type != typ, "fields": {}}
                for key in FIELDS.get(typ, []):
                    attrs = _flatten_mapping_values(m["dynamic_fields"].get(key))
                    if key == "theme": attrs += _flatten_mapping_values(m["dynamic_fields"].get("theme_1"))
                    if key == "target_audience": attrs += [v for k,v in m["dynamic_fields"].items() if k.startswith("target_audience_")]
                    before = [ws.cell(row["row_number"], columns[attr]).value for attr in attrs if attr in columns]
                    record["fields"][key] = {"before": [v for v in before if v is not None],
                                              "after": decisions.get(key, {}).get("values", []),
                                              "source": decisions.get(key, {}).get("source")}
                w.close()
                records.append(record)
    db.close()
    fields = [f for r in records if not r["needs_rerouting"] for f in r["fields"].values()]
    return {"products": len(records), "rerouting": sum(r["needs_rerouting"] for r in records),
            "same_category_products": sum(not r["needs_rerouting"] for r in records),
            "missing_before": sum(not f["before"] for f in fields),
            "missing_after": sum(not f["after"] for f in fields),
            "llm_calls": 0, "rows": records}


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--database", type=Path, default=ROOT / "data/fbm-pipeline.db")
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("zip_paths", nargs="+", type=Path)
    args = parser.parse_args()
    report = audit(args.database, args.zip_paths)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(report, ensure_ascii=False, indent=2))
    print(json.dumps({k: v for k, v in report.items() if k != "rows"}, ensure_ascii=False))
