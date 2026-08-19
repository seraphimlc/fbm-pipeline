#!/usr/bin/env python3
"""Focused regression checks for the Vindhvisk BED_FRAME export mapping."""

from __future__ import annotations

import copy
import json
import sys
import tempfile
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))

from app.models import Product, ProductData
from app.api.products import _clear_template_data_rows, _copy_import_data_row
from app.pipeline.amazon_export.strategies.bed_frame import apply_bed_frame_strategy
from app.pipeline.amazon_export.writer import build_amazon_template_file
from app.pipeline.step10_amazon_template import _index_template_columns, _load_template_mapping


MAPPING_PATH = ROOT / "backend" / "app" / "pipeline" / "template_mappings" / "vindhvisk_bed_frame.json"


def _flatten(value):
    if isinstance(value, str):
        return [value]
    if isinstance(value, list):
        return [field for item in value for field in _flatten(item)]
    if isinstance(value, dict):
        return [field for item in value.values() for field in _flatten(item)]
    return []


def _mapping() -> dict:
    mapping = json.loads(MAPPING_PATH.read_text(encoding="utf-8"))
    mapping["template_path"] = str(ROOT / "backend" / "app" / "pipeline" / mapping["template_path"])
    return mapping


def _product_data(output_dir: Path) -> SimpleNamespace:
    return SimpleNamespace(
        item_code="W5819P522963",
        title="Queen Size Bed Frame Upholstered Platform Bed with Adjustable Headboard, Linen Fabric Padded, Modern Platform Bed No Box Spring Needed, Easy Assembly, Beige",
        product_type="Bed Frames",
        color="Beige",
        material="Fabric,Linen,Plywood,Wood",
        filler=None,
        description=(
            "Product Information\nColor\nBeige\nMaterial\nPoplar wood+Plywood+Linen fabric+foam+Plastic bed feet\n"
            "Bed Type\nPlatform Bed\nSize\nQueen\nSpring Box\nNo need\nNumber of Bed Slats\n12pcs\n"
            "Bed Weight Capability\n1000 lbs\nAssembly Required\nYes\nAdjustable headboard\n32.87-37 two heights adjustable\nOrigin\nChina\n"
            "Detailed instructions are provided."
        ),
        features=json.dumps(["Modern upholstered platform bed with 12 slats and no box spring needed."]),
        variants=json.dumps([{"attributes": {"Bed Size": "Queen", "Bed Type": "Bed Frame"}}]),
        gigab2b_raw_snapshot=None,
        origin=None,
        dimension_length=81.89,
        dimension_width=62.20,
        dimension_height=47.64,
        weight=72.97,
        packages=json.dumps([{"length": 63.58, "width": 24.61, "height": 5.91, "weight_value": 80.47, "qty": "1"}]),
        stock=111,
        suggested_price=230.45,
        profit=32.0,
        profit_rate=15.0,
        material_dir=str(output_dir),
        amazon_template_path=None,
        leaf_category="Bed Frames",
        categories=json.dumps(["Home & Kitchen", "Furniture", "Bed Frames"]),
        listing_title="Vindhvisk Queen Upholstered Platform Bed Frame with Adjustable Headboard",
        listing_description="Queen platform bed frame with linen upholstered headboard and wood slats.",
        listing_bullets=json.dumps([
            "Adjustable upholstered headboard supports comfortable mattress placement.",
            "Platform bed uses wood slats and does not need a box spring.",
            "Queen bed frame has a 1000 lb maximum weight recommendation.",
            "Linen fabric and wood construction create a modern bedroom look.",
            "Installation manual is included for assembly.",
        ]),
        listing_product_highlights=json.dumps([]),
        listing_search_terms="queen platform bed frame upholstered adjustable headboard linen wood slats",
        listing_check=json.dumps({}),
        listing_primary_keyword="queen platform bed frame",
    )


def _strategy_context(product, pd, mapping, package=None):
    return SimpleNamespace(
        product=product,
        product_data=pd,
        fields=mapping["dynamic_fields"],
        package=package or {"length": 63.58, "width": 24.61, "height": 5.91, "weight": 80.47},
        fill={mapping["dynamic_fields"]["country_of_origin"]: "China"},
        warnings=[],
    )


def _expect_strategy_error(product, pd, mapping, expected: str) -> None:
    try:
        apply_bed_frame_strategy(_strategy_context(product, pd, mapping))
    except ValueError as exc:
        assert expected in str(exc), exc
    else:
        raise AssertionError(f"expected error containing {expected!r}")


def _fake_upload(_, __, mapping, ___):
    return ({mapping["image_fields"]["main"]: "https://example.test/bed-frame-main.jpg"}, [], [])


def main() -> None:
    mapping = _mapping()
    workbook = load_workbook(mapping["template_path"], keep_vba=True, read_only=True)
    assert mapping["data_row"] == 7
    assert len(_index_template_columns(workbook["Template"])) == 331
    template_fields = set(_index_template_columns(workbook["Template"]))
    mapped_fields = set(mapping["fixed_values"]) | set(_flatten(mapping["dynamic_fields"]))
    mapped_fields.discard("::record_action")
    mapped_fields |= set(_flatten(mapping["bullet_fields"])) | set(_flatten(mapping["image_fields"]))
    mapped_fields |= set(_flatten(mapping["package_fields"])) | set(_flatten(mapping["required_fields"]))
    assert not (mapped_fields - template_fields), sorted(mapped_fields - template_fields)

    with tempfile.TemporaryDirectory() as tmp:
        output_dir = Path(tmp)
        product = SimpleNamespace(id=4, upc="787471049980", brand="Vindhvisk", images=None, aplus=None)
        pd = _product_data(output_dir)
        with patch("app.pipeline.step10_amazon_template._upload_listing_images", _fake_upload):
            result = build_amazon_template_file(product, pd, mapping)
        ws = load_workbook(result["path"], keep_vba=True, read_only=True)["Template"]
        columns = _index_template_columns(ws)
        fields = mapping["dynamic_fields"]
        row = mapping["data_row"]
        value = lambda attr: ws[f"{columns[attr]}{row}"].value
        assert value("contribution_sku#1.value") == "W5819P522963"
        assert value("product_type#1.value") == "BED_FRAME"
        assert value(fields["product_id_value"]) == "787471049980"
        assert value(fields["size"]) == "Queen (U.S. Standard)"
        assert value(fields["form_factor"]) == "Platform Bed"
        assert value(fields["maximum_weight"]) == 1000
        assert value(fields["country_of_origin"]) == "China"
        assert value(fields["item_length"]) == 81.89 and value(fields["item_width"]) == 62.2
        assert value(fields["item_height"]) == 47.64 and value(fields["item_weight"]) == 72.97
        assert value(mapping["package_fields"]["length_value"]) == 63.58
        assert value(mapping["package_fields"]["weight_value"]) == 80.47
        assert [value(field) for field in fields["materials"][:3]] == ["Linen", "Engineered Wood", "Wood"]
        assert [value(field) for field in fields["included_components"][:3]] == ["Headboard", "Installation Manual", "Slat"]
        assert [value(field) for field in fields["special_features"][:4]] == ["Adjustable", "No Box Spring Needed", "Slatted", "Upholstered"]
        for attr in (
            "target_audience[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value",
            "item_shape[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value",
            "furniture_finish[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value",
            "finish_type[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value",
            "headboard_height[marketplace_id=ATVPDKIKX0DER]#1.value",
        ):
            assert value(attr) is None, attr

        combined_path = output_dir / "combined.xlsm"
        combined_path.write_bytes(Path(mapping["template_path"]).read_bytes())
        combined_wb = load_workbook(combined_path, keep_vba=True)
        combined_ws = combined_wb["Template"]
        _clear_template_data_rows(combined_ws, 1, data_row=7)
        _copy_import_data_row(Path(result["path"]), combined_ws, 7, data_row=7)
        combined_wb.save(combined_path)
        combined_check = load_workbook(combined_path, keep_vba=True, read_only=True)["Template"]
        assert combined_check["A7"].value == "W5819P522963"
        assert combined_check["A8"].value is None

        invalid_mapping = copy.deepcopy(mapping)
        invalid_mapping["output_filename"] = "invalid_dropdown.xlsm"
        invalid_mapping["fixed_values"]["package_level[marketplace_id=ATVPDKIKX0DER]#1.value"] = "Parcel"
        with patch("app.pipeline.step10_amazon_template._upload_listing_images", _fake_upload):
            invalid_result = build_amazon_template_file(product, _product_data(output_dir), invalid_mapping)
        assert any("不在模板可选项中" in warning for warning in invalid_result["warnings"]), invalid_result["warnings"]

        no_origin = _product_data(output_dir)
        no_origin.description = no_origin.description.replace("Origin\nChina\n", "")
        _expect_strategy_error(product, no_origin, mapping, "原产地 Origin")
        no_size = _product_data(output_dir)
        no_size.title = no_size.title.replace("Queen", "")
        no_size.description = no_size.description.replace("Size\nQueen\n", "")
        no_size.variants = "[]"
        _expect_strategy_error(product, no_size, mapping, "床型 Size")
        no_stock = _product_data(output_dir)
        no_stock.stock = None
        _expect_strategy_error(product, no_stock, mapping, "可售库存")
        negative_stock = _product_data(output_dir)
        negative_stock.stock = -1
        _expect_strategy_error(product, negative_stock, mapping, "可售库存不能为负数")
        no_upc = SimpleNamespace(id=4, upc=None, brand="Vindhvisk", images=None, aplus=None)
        _expect_strategy_error(no_upc, _product_data(output_dir), mapping, "UPC")

    for leaf, title in (("Children Bed Frames", "Kids Bed Frame"), ("Adjustable Bed Bases", "Adjustable Bed Base")):
        try:
            _load_template_mapping(Product(brand="Vindhvisk"), ProductData(leaf_category=leaf, title=title))
        except ValueError:
            pass
        else:
            raise AssertionError(f"non-standard bed category unexpectedly mapped: {leaf}")
    try:
        _load_template_mapping(Product(brand="Other"), ProductData(leaf_category="Bed Frames", title="Queen Bed Frame"))
    except ValueError:
        pass
    else:
        raise AssertionError("non-Vindhvisk Bed Frames unexpectedly mapped")

    print("OK: BED_FRAME mapping, strategy, dropdown, and output checks passed")


if __name__ == "__main__":
    main()
