#!/usr/bin/env python3
"""R1 Batch 2 catalog-export outcome integration checks on isolated MySQL."""

from __future__ import annotations

import asyncio
import ipaddress
import json
import socket
import sys
from contextlib import contextmanager
from datetime import datetime
from pathlib import Path


ROOT = Path(__file__).resolve().parents[1]
BACKEND = ROOT / "backend"
if str(BACKEND) not in sys.path:
    sys.path.insert(0, str(BACKEND))


@contextmanager
def external_socket_guard():
    """Allow isolated MySQL loopback traffic and fail on any external socket."""

    original_connect = socket.socket.connect
    external_attempts: list[object] = []

    def guarded_connect(sock, address):
        if isinstance(address, tuple) and address:
            host = str(address[0])
            try:
                is_loopback = ipaddress.ip_address(host).is_loopback
            except ValueError:
                is_loopback = host.lower() == "localhost"
            if not is_loopback:
                external_attempts.append(address)
                raise AssertionError(f"unexpected external socket during R1 catalog export test: {address}")
        return original_connect(sock, address)

    socket.socket.connect = guarded_connect
    try:
        yield
    finally:
        socket.socket.connect = original_connect
    assert not external_attempts, external_attempts


def test_catalog_export_response_normalizer_contract() -> None:
    from app.api.schemas import CatalogExportFileResponse, CatalogExportResultResponse
    from app.task_runtime.catalog_export_status import normalize_catalog_export_response

    raw_payload = {
        "status": {"nested": "done"},
        "artifact_available": {"nested": True},
        "requested_count": {"nested": 6},
        "success_count": [1],
        "skipped_count": {"nested": 1},
        "failed_count": "4",
        "report_count": [6],
        "filename": {"nested": "unsafe.zip"},
        "reason": ["unsafe reason"],
        "categories": [{"nested": "unsafe"}, "Safe Category"],
        "rows": [
            None,
            7,
            {
                "catalog_id": 71001,
                "item_code": {"nested": "unsafe"},
                "seller_sku": ["unsafe"],
                "category": {"nested": "unsafe"},
                "status": {"nested": "exported"},
                "reason": ["unsafe"],
            },
            {"catalog_id": 71002, "status": "exported", "reason": "duplicate first"},
            {"catalog_id": 71002, "status": "skipped", "reason": "duplicate second"},
            {"catalog_id": None, "product_id": None, "status": None, "reason": None},
        ],
    }

    normalized = normalize_catalog_export_response(raw_payload)
    validated = CatalogExportResultResponse.model_validate(normalized)
    assert validated.status == "failed", validated
    assert validated.artifact_available is False, validated
    assert validated.requested_count == 6, validated
    assert validated.success_count == 1, validated
    assert validated.skipped_count == 1, validated
    assert validated.failed_count == 4, validated
    assert validated.report_count == 6, validated
    assert validated.filename is None and validated.reason == "导出结果格式异常", validated
    assert validated.categories == ["Safe Category"], validated
    assert [row.row_ordinal for row in validated.rows] == [1, 2, 3, 4, 5, 6], validated
    assert validated.rows[0].reason == "导出结果行格式异常", validated.rows[0]
    assert validated.rows[1].reason == "导出结果行格式异常", validated.rows[1]
    assert validated.rows[2].item_code is None, validated.rows[2]
    assert validated.rows[2].status == "failed", validated.rows[2]
    assert validated.rows[2].reason == "导出结果行格式异常", validated.rows[2]
    assert validated.rows[3].catalog_id == validated.rows[4].catalog_id == 71002, validated.rows
    assert validated.rows[5].reason == "导出结果行格式异常", validated.rows[5]

    file_response = CatalogExportFileResponse(
        task_id=1,
        task_status="failed",
        rows=normalized["rows"],
    )
    assert len(file_response.rows) == 6, file_response


def test_catalog_export_response_provenance_contract() -> None:
    from app.task_runtime.catalog_export_status import (
        ARTIFACT_MODE_UNAVAILABLE,
        MAX_SAFE_JSON_INTEGER,
        _parse_catalog_export_payload_result,
        normalize_catalog_export_response,
        project_catalog_effective_terminal_record,
        resolve_catalog_export_artifact,
        select_catalog_export_payload,
    )

    older_material = {
        "status": "done",
        "filename": "older-valid.zip",
        "oss_url": "https://downloads.invalid/older-valid.zip",
        "category": "R1 Provenance Older Valid",
        "categories": ["R1 Provenance Older Valid"],
    }
    invalid = _parse_catalog_export_payload_result("{not-json")
    assert invalid.payload == {}, invalid
    assert invalid.json_valid is False and invalid.material is False and invalid.malformed is True, invalid
    valid_material = _parse_catalog_export_payload_result(json.dumps(older_material))
    assert valid_material.payload == older_material, valid_material
    assert valid_material.json_valid is True and valid_material.material is True and valid_material.malformed is False

    assert select_catalog_export_payload("{not-json", [json.dumps(older_material)]) == older_material
    assert select_catalog_export_payload(
        '{"status":"done","file_path":NaN}',
        ["{still-not-json", json.dumps(older_material)],
    ) == older_material
    assert select_catalog_export_payload("{not-json", ["{still-not-json", '{"status":"done"}']) == {}

    fallback_record = project_catalog_effective_terminal_record(
        owner_kind="task_run",
        task_type="catalog_export",
        status="succeeded",
        summary_json_or_dict="{not-json",
        step_result_json_or_dict=json.dumps(older_material),
    )
    assert fallback_record.authoritative is True, fallback_record
    assert fallback_record.effective_status == "succeeded", fallback_record
    assert fallback_record.selected_payload == older_material, fallback_record

    invalid_only_record = project_catalog_effective_terminal_record(
        owner_kind="task_run",
        task_type="catalog_export",
        status="succeeded",
        summary_json_or_dict="{not-json",
        step_result_json_or_dict='{ "status": NaN }',
    )
    assert invalid_only_record.authoritative is False, invalid_only_record
    assert invalid_only_record.effective_status == "succeeded", invalid_only_record
    assert invalid_only_record.selected_payload == {}, invalid_only_record

    raw_payload = {
        "status": "done",
        "artifact_available": True,
        "filename": "unsafe-provenance.zip",
        "oss_url": "\ud800",
        "rows": [
            {
                "catalog_id": MAX_SAFE_JSON_INTEGER + 1,
                "item_code": "TOO-LARGE-ID",
                "status": "exported",
                "reason": "ok",
            },
            {"catalog_id": 73002, "item_code": "\ud800", "status": "exported", "reason": "ok"},
            {"catalog_id": 73003, "item_code": "SAFE", "status": "exported", "reason": "\ud800"},
            {
                "catalog_id": 73004,
                "item_code": {"nested": "unsafe"},
                "status": "exported",
                "reason": "ok",
            },
            {"catalog_id": 73005, "item_code": "INVALID-STATUS", "status": "done", "reason": "ok"},
        ],
    }
    selected = select_catalog_export_payload(json.dumps(raw_payload), [])
    assert selected["rows"][0]["catalog_id"] == MAX_SAFE_JSON_INTEGER + 1, selected
    assert selected["rows"][1]["item_code"] == "\ud800", selected
    assert selected["rows"][2]["reason"] == "\ud800", selected
    assert selected["rows"][3]["item_code"] == {"nested": "unsafe"}, selected
    assert selected["rows"][4]["status"] == "done", selected

    normalized = normalize_catalog_export_response(selected)
    assert normalized["success_count"] == 0, normalized
    assert normalized["failed_count"] == 5, normalized
    for row in normalized["rows"]:
        assert row["status"] == "failed", row
        assert row["reason"] == "导出结果行格式异常", row
    assert normalized["rows"][0]["catalog_id"] is None, normalized
    assert normalized["rows"][1]["item_code"] is None, normalized
    assert normalized["rows"][2]["reason"] == "导出结果行格式异常", normalized
    assert normalized["rows"][3]["item_code"] is None, normalized
    assert resolve_catalog_export_artifact(
        selected,
        allowed_export_root=ROOT / "tmp" / "r1-provenance-exports",
        object_cache_subdir="provenance",
    ).mode == ARTIFACT_MODE_UNAVAILABLE

    unsafe_record = project_catalog_effective_terminal_record(
        owner_kind="task_run",
        task_type="catalog_export",
        status="succeeded",
        summary_json_or_dict=json.dumps(raw_payload),
    )
    assert unsafe_record.authoritative is True, unsafe_record
    assert unsafe_record.effective_status == "failed", unsafe_record


def test_single_validated_catalog_export_outcome_contract() -> None:
    from tempfile import TemporaryDirectory

    from app.task_runtime.catalog_export_status import (
        ARTIFACT_MODE_LOCAL,
        ARTIFACT_MODE_UNAVAILABLE,
        catalog_export_resolution_is_ready,
        normalize_catalog_export_response,
        resolve_catalog_export_artifact,
    )

    with TemporaryDirectory(prefix="r1-validated-outcome-", dir=ROOT / "tmp") as temporary_dir:
        export_root = Path(temporary_dir)
        artifact = export_root / "validated-outcome.zip"
        artifact.write_bytes(b"validated-outcome")

        unsafe_row_raw = {
            "status": "done",
            "artifact_available": True,
            "requested_count": 1,
            "success_count": 1,
            "exported_count": 1,
            "skipped_count": 0,
            "failed_count": 0,
            "report_count": 1,
            "filename": artifact.name,
            "file_path": str(artifact),
            "rows": [{"catalog_id": 2 ** 53, "status": "exported", "reason": "raw ok"}],
        }
        unsafe_outcome = normalize_catalog_export_response(unsafe_row_raw)
        assert unsafe_outcome["status"] == "failed", unsafe_outcome
        assert unsafe_outcome["requested_count"] == 1, unsafe_outcome
        assert unsafe_outcome["success_count"] == 0, unsafe_outcome
        assert unsafe_outcome["exported_count"] == 0, unsafe_outcome
        assert unsafe_outcome["skipped_count"] == 0, unsafe_outcome
        assert unsafe_outcome["failed_count"] == 1, unsafe_outcome
        assert unsafe_outcome["report_count"] == 1, unsafe_outcome
        assert unsafe_outcome["rows"][0]["status"] == "failed", unsafe_outcome
        unsafe_resolution = resolve_catalog_export_artifact(
            unsafe_outcome,
            allowed_export_root=export_root,
            object_cache_subdir="unsafe",
        )
        assert unsafe_resolution.mode == ARTIFACT_MODE_LOCAL, unsafe_resolution
        assert catalog_export_resolution_is_ready(unsafe_resolution) is False, unsafe_resolution

        legacy_raw = {
            "status": "done",
            "filename": artifact.name,
            "file_path": str(artifact),
        }
        legacy_outcome = normalize_catalog_export_response(legacy_raw)
        assert "artifact_available" not in legacy_outcome, legacy_outcome
        assert legacy_outcome["status"] == "done", legacy_outcome
        assert legacy_outcome["requested_count"] == 0, legacy_outcome
        assert legacy_outcome["requested_count"] == (
            legacy_outcome["success_count"]
            + legacy_outcome["skipped_count"]
            + legacy_outcome["failed_count"]
        ), legacy_outcome
        legacy_resolution = resolve_catalog_export_artifact(
            legacy_outcome,
            allowed_export_root=export_root,
            object_cache_subdir="legacy",
        )
        assert legacy_resolution.mode == ARTIFACT_MODE_LOCAL, legacy_resolution
        assert catalog_export_resolution_is_ready(legacy_resolution) is True, legacy_resolution

        inconsistent_outcome = normalize_catalog_export_response({
            "status": "done",
            "requested_count": 2,
            "success_count": 1,
            "skipped_count": 0,
            "failed_count": 1,
            "filename": artifact.name,
            "file_path": str(artifact),
        })
        assert inconsistent_outcome["status"] == "failed", inconsistent_outcome
        assert inconsistent_outcome["requested_count"] == 2, inconsistent_outcome
        assert inconsistent_outcome["reason"] == "导出结果格式异常", inconsistent_outcome

        explicit_unavailable = normalize_catalog_export_response({
            "status": "done",
            "artifact_available": False,
            "filename": artifact.name,
            "file_path": str(artifact),
            "rows": [{"catalog_id": 74001, "status": "exported"}],
        })
        assert explicit_unavailable["status"] == "failed", explicit_unavailable
        assert explicit_unavailable["success_count"] == 1, explicit_unavailable
        assert resolve_catalog_export_artifact(
            explicit_unavailable,
            allowed_export_root=export_root,
            object_cache_subdir="explicit-unavailable",
        ).mode == ARTIFACT_MODE_UNAVAILABLE


def _malformed_catalog_response_payload(category: str) -> dict:
    return {
        "status": {"nested": "done"},
        "artifact_available": {"nested": True},
        "requested_count": {"nested": 10},
        "success_count": [1],
        "skipped_count": {"nested": 1},
        "failed_count": "8",
        "report_count": [10],
        "filename": {"nested": "unsafe.zip"},
        "reason": ["unsafe reason"],
        "categories": [{"nested": "unsafe"}, category],
        "rows": [
            None,
            7,
            {
                "catalog_id": 72001,
                "product_id": {"nested": 1},
                "item_code": {"nested": "unsafe"},
                "seller_sku": ["unsafe"],
                "category": {"nested": category},
                "status": {"nested": "exported"},
                "reason": ["unsafe"],
                "template_file": {"nested": "unsafe.xlsm"},
                "output_file": ["unsafe.xlsx"],
            },
            {"catalog_id": 72002, "status": "exported", "reason": "duplicate first"},
            {"catalog_id": 72002, "status": "skipped", "reason": "duplicate second"},
            {"catalog_id": None, "product_id": None, "status": None, "reason": None},
            {"catalog_id": None, "product_id": None, "status": None, "reason": None},
            {"item_code": "A/B", "status": "failed", "reason": "slash code"},
            {"item_code": "A B", "status": "failed", "reason": "space code"},
            {"item_code": "商品-甲", "status": "failed", "reason": "Unicode 行"},
            {"catalog_id": (2 ** 53), "status": "exported", "reason": "ok unsafe id"},
            {"catalog_id": 72003, "item_code": "\ud800", "status": "exported", "reason": "ok surrogate code"},
            {"catalog_id": 72004, "item_code": "SAFE", "status": "exported", "reason": "\ud800"},
            {
                "catalog_id": 72005,
                "item_code": {"nested": "unsafe"},
                "status": "exported",
                "reason": "ok nested field",
            },
            {"catalog_id": 72006, "item_code": "BAD-STATUS", "status": "done", "reason": "ok invalid status"},
        ],
    }


async def test_catalog_export_response_normalization_across_apis(environment) -> None:
    import httpx

    from app.main import app
    from app.models import OfflineTask, OfflineTaskStep, TaskGroup, TaskRun, TaskStep
    from app.task_runtime.catalog_export_status import (
        ARTIFACT_MODE_UNAVAILABLE,
        normalize_catalog_export_response,
        resolve_catalog_export_artifact,
    )

    category = "R1 Malformed Response Category"
    raw_payload = _malformed_catalog_response_payload(category)
    expected = normalize_catalog_export_response(raw_payload)
    raw_resolution = resolve_catalog_export_artifact(
        raw_payload,
        allowed_export_root=environment.data_dir / "exports",
        object_cache_subdir="malformed-response",
    )
    assert raw_resolution.mode == ARTIFACT_MODE_UNAVAILABLE, raw_resolution

    async with environment.session_factory() as db:
        run = TaskRun(
            task_type="catalog_export",
            title="R1 malformed response API",
            status="failed",
            summary_json=json.dumps(raw_payload, ensure_ascii=True),
        )
        db.add(run)
        await db.flush()
        group = TaskGroup(
            task_run_id=run.id,
            group_key="export_file",
            title="R1 malformed response API",
            status="failed",
            summary_json=json.dumps(raw_payload, ensure_ascii=True),
        )
        db.add(group)
        await db.flush()
        db.add(TaskStep(
            task_run_id=run.id,
            task_group_id=group.id,
            step_key="catalog_export_template_malformed_response",
            step_type="catalog_export_template",
            status="failed",
            result_json=json.dumps(raw_payload, ensure_ascii=True),
        ))
        offline_task = OfflineTask(
            task_type="catalog_export",
            title="R1 malformed response OfflineTask",
            status="failed",
            total_steps=1,
            failed_steps=1,
            result_json=json.dumps({"status": "failed"}),
        )
        db.add(offline_task)
        await db.flush()
        db.add(OfflineTaskStep(
            task_id=offline_task.id,
            step_type="catalog_export_template",
            title="R1 malformed response OfflineTask step",
            status="failed",
            progress_current=15,
            progress_total=15,
            result_json=json.dumps(raw_payload, ensure_ascii=True),
        ))
        action_task = OfflineTask(
            task_type="catalog_export",
            title="R1 malformed response OfflineTask action",
            status="pending",
            total_steps=1,
            result_json=json.dumps({"status": "failed"}),
        )
        db.add(action_task)
        await db.flush()
        db.add(OfflineTaskStep(
            task_id=action_task.id,
            step_type="catalog_export_template",
            title="R1 malformed response OfflineTask action step",
            status="pending",
            progress_current=0,
            progress_total=15,
            result_json=json.dumps(raw_payload, ensure_ascii=True),
        ))
        await db.commit()
        run_id = run.id
        offline_task_id = offline_task.id
        action_task_id = action_task.id

    transport = httpx.ASGITransport(app=app, client=("127.0.0.1", 41331))
    async with httpx.AsyncClient(transport=transport, base_url="http://testserver") as client:
        with external_socket_guard():
            task_list = await client.get(
                "/api/task-runs",
                params={"view": "all", "q": "R1 malformed response API", "page_size": 20},
            )
            assert task_list.status_code == 200, task_list.text
            task_item = next(item for item in task_list.json()["items"] if item["id"] == run_id)
            assert task_item["catalog_export_result"] == expected, task_item
            assert "download_result" not in task_item["available_actions"], task_item

            detail = await client.get(f"/api/task-runs/{run_id}")
            assert detail.status_code == 200, detail.text
            detail_payload = detail.json()["catalog_export_result"]
            assert detail_payload == expected, detail_payload
            assert [row["row_ordinal"] for row in detail_payload["rows"]] == list(range(1, 16)), detail_payload
            assert detail_payload["rows"][0]["reason"] == "导出结果行格式异常", detail_payload
            assert detail_payload["rows"][1]["reason"] == "导出结果行格式异常", detail_payload
            assert detail_payload["rows"][2]["reason"] == "导出结果行格式异常", detail_payload
            assert detail_payload["rows"][3]["catalog_id"] == detail_payload["rows"][4]["catalog_id"] == 72002
            assert detail_payload["rows"][7]["item_code"] == "A/B", detail_payload
            assert detail_payload["rows"][8]["item_code"] == "A B", detail_payload
            assert detail_payload["rows"][9]["item_code"] == "商品-甲", detail_payload
            for row in detail_payload["rows"][10:15]:
                assert row["status"] == "failed", row
                assert row["reason"] == "导出结果行格式异常", row
            assert detail_payload["rows"][10]["catalog_id"] is None, detail_payload
            assert detail_payload["rows"][11]["item_code"] is None, detail_payload
            assert detail_payload["rows"][12]["item_code"] == "SAFE", detail_payload
            assert detail_payload["rows"][13]["item_code"] is None, detail_payload
            assert detail_payload["rows"][14]["item_code"] == "BAD-STATUS", detail_payload

            offline_list = await client.get(
                "/api/offline-tasks",
                params={"task_type": "catalog_export", "page_size": 100},
            )
            assert offline_list.status_code == 200, offline_list.text
            offline_item = next(item for item in offline_list.json()["items"] if item["id"] == offline_task_id)
            assert offline_item["catalog_export_result"] == expected, offline_item
            assert json.loads(offline_item["result_json"]) == expected, offline_item
            assert offline_item["can_download"] is False, offline_item

            offline_detail = await client.get(f"/api/offline-tasks/{offline_task_id}")
            assert offline_detail.status_code == 200, offline_detail.text
            offline_detail_payload = offline_detail.json()
            assert offline_detail_payload["catalog_export_result"] == expected, offline_detail_payload
            assert json.loads(offline_detail_payload["result_json"]) == expected, offline_detail_payload
            assert all(step["result_json"] is None for step in offline_detail_payload["steps"]), offline_detail_payload

            paused = await client.post(f"/api/offline-tasks/{action_task_id}/pause")
            assert paused.status_code == 200, paused.text
            paused_payload = paused.json()
            assert paused_payload["status"] == "paused", paused_payload
            assert paused_payload["catalog_export_result"] == expected, paused_payload
            assert json.loads(paused_payload["result_json"]) == expected, paused_payload
            assert all(step["result_json"] is None for step in paused_payload["steps"]), paused_payload
            assert paused_payload["can_download"] is False, paused_payload

            export_files = await client.get("/api/products/catalog/export-files", params={"page_size": 100})
            assert export_files.status_code == 200, export_files.text
            export_item = next(
                item for item in export_files.json()["items"]
                if item["task_source"] == "task_run" and item["task_id"] == run_id
            )
            assert export_item["rows"] == expected["rows"], export_item
            assert export_item["can_download"] is False, export_item
            offline_export_item = next(
                item for item in export_files.json()["items"]
                if item["task_source"] == "offline_task" and item["task_id"] == offline_task_id
            )
            assert offline_export_item["rows"] == expected["rows"], offline_export_item
            assert offline_export_item["can_download"] is False, offline_export_item

            categories = await client.get("/api/products/catalog/export-categories")
            assert categories.status_code == 200, categories.text
            json.dumps(categories.json(), ensure_ascii=False, allow_nan=False).encode("utf-8")

            download = await client.get(f"/api/task-runs/{run_id}/download")
            assert download.status_code == 400, download.text
            offline_download = await client.get(f"/api/offline-tasks/{offline_task_id}/download")
            assert offline_download.status_code == 400, offline_download.text


async def test_unsafe_valid_artifact_fails_all_catalog_consumers(environment) -> None:
    import httpx
    from sqlalchemy import select
    from sqlalchemy.orm import selectinload

    from app.main import app
    from app.models import OfflineTask, OfflineTaskStep, TaskGroup, TaskRun, TaskStep
    from app.services import offline_tasks
    from app.task_runtime import catalog_export_workers
    from app.task_runtime.catalog_export_status import normalize_catalog_export_response
    from app.task_runtime.catalog_export_workers import register_catalog_export_workers
    from app.task_runtime.scheduler import drain_ready_steps

    category = "R1 Unsafe Valid Artifact Category"
    unsafe_catalog_id = 2 ** 53
    artifact = environment.data_dir / "exports" / "r1-unsafe-valid-artifact.zip"
    artifact.parent.mkdir(parents=True, exist_ok=True)
    artifact.write_bytes(b"r1-unsafe-valid-artifact")
    raw_payload = {
        "status": "done",
        "artifact_available": True,
        "requested_count": 1,
        "success_count": 1,
        "exported_count": 1,
        "skipped_count": 0,
        "failed_count": 0,
        "report_count": 1,
        "filename": artifact.name,
        "file_path": str(artifact),
        "file_size": artifact.stat().st_size,
        "category": category,
        "categories": [category],
        "catalog_product_ids": [unsafe_catalog_id],
        "rows": [{
            "catalog_id": unsafe_catalog_id,
            "product_id": 99001,
            "item_code": "R1-UNSAFE-VALID-ARTIFACT",
            "category": category,
            "status": "exported",
            "reason": "raw payload claims success",
        }],
    }
    expected = normalize_catalog_export_response(raw_payload)
    assert expected["status"] == "failed", expected
    assert expected["requested_count"] == 1, expected
    assert expected["success_count"] == 0, expected
    assert expected["skipped_count"] == 0, expected
    assert expected["failed_count"] == 1, expected
    assert expected["rows"][0]["catalog_id"] is None, expected

    async with environment.session_factory() as db:
        api_run = TaskRun(
            task_type="catalog_export",
            title="R1 unsafe outer projection API unsafe",
            status="succeeded",
            summary_json=json.dumps(raw_payload),
        )
        db.add(api_run)
        await db.flush()
        api_group = TaskGroup(
            task_run_id=api_run.id,
            group_key="export_file",
            title="R1 unsafe outer projection API unsafe",
            status="succeeded",
            summary_json=json.dumps(raw_payload),
        )
        db.add(api_group)
        await db.flush()
        db.add(TaskStep(
            task_run_id=api_run.id,
            task_group_id=api_group.id,
            step_key="catalog_export_template_unsafe_valid_artifact_api",
            step_type="catalog_export_template",
            status="succeeded",
            result_json=json.dumps(raw_payload),
        ))

        raw_succeeded_run = TaskRun(
            task_type="catalog_export",
            title="R1 unsafe outer projection API raw succeeded no material",
            status="succeeded",
            summary_json=json.dumps({"status": "done"}),
        )
        raw_failed_run = TaskRun(
            task_type="catalog_export",
            title="R1 unsafe outer projection API raw failed",
            status="failed",
            summary_json=json.dumps({"status": "failed"}),
        )
        db.add_all([raw_succeeded_run, raw_failed_run])

        raw_done_offline_task = OfflineTask(
            task_type="catalog_export",
            title="R1 effective projection OfflineTask raw done no material",
            status="done",
            total_steps=0,
            result_json=json.dumps({"status": "done"}),
        )
        api_offline_task = OfflineTask(
            task_type="catalog_export",
            title="R1 effective projection OfflineTask unsafe",
            status="done",
            total_steps=1,
            success_steps=1,
            result_json=json.dumps(raw_payload),
        )
        raw_failed_offline_task = OfflineTask(
            task_type="catalog_export",
            title="R1 effective projection OfflineTask raw failed",
            status="failed",
            total_steps=0,
            result_json=json.dumps({"status": "failed"}),
        )
        db.add_all([raw_done_offline_task, api_offline_task, raw_failed_offline_task])
        await db.flush()
        db.add(OfflineTaskStep(
            task_id=api_offline_task.id,
            step_type="catalog_export_template",
            title="R1 unsafe valid artifact OfflineTask API step",
            status="done",
            progress_current=1,
            progress_total=1,
            result_json=json.dumps(raw_payload),
        ))

        reuse_run = TaskRun(
            task_type="catalog_export",
            title="R1 unsafe valid artifact TaskRun reuse",
            status="pending",
            payload_json=json.dumps({
                "catalog_product_ids": [unsafe_catalog_id],
                "category": category,
                "categories": [category],
            }),
            summary_json=json.dumps(raw_payload),
        )
        db.add(reuse_run)
        await db.flush()
        reuse_group = TaskGroup(
            task_run_id=reuse_run.id,
            group_key="export_file",
            title="R1 unsafe valid artifact TaskRun reuse",
            status="pending",
        )
        db.add(reuse_group)
        await db.flush()
        reuse_step = TaskStep(
            task_run_id=reuse_run.id,
            task_group_id=reuse_group.id,
            step_key="catalog_export_template_unsafe_valid_artifact_reuse",
            step_type="catalog_export_template",
            status="ready",
            payload_json=json.dumps({
                "catalog_product_ids": [unsafe_catalog_id],
                "category": category,
                "categories": [category],
            }),
            progress_current=0,
            progress_total=1,
            max_attempts=1,
        )
        db.add(reuse_step)

        reuse_offline_task = OfflineTask(
            task_type="catalog_export",
            title="R1 unsafe valid artifact OfflineTask reuse",
            status="pending",
            total_steps=1,
            result_json=json.dumps(raw_payload),
        )
        db.add(reuse_offline_task)
        await db.flush()
        reuse_offline_step = OfflineTaskStep(
            task_id=reuse_offline_task.id,
            step_type="catalog_export_template",
            title="R1 unsafe valid artifact OfflineTask reuse step",
            status="pending",
            progress_current=0,
            progress_total=1,
            payload_json=json.dumps({
                "catalog_product_ids": [unsafe_catalog_id],
                "category": category,
                "categories": [category],
            }),
            result_json=json.dumps(raw_payload),
        )
        db.add(reuse_offline_step)
        await db.commit()
        api_run_id = api_run.id
        raw_succeeded_run_id = raw_succeeded_run.id
        raw_failed_run_id = raw_failed_run.id
        raw_done_offline_task_id = raw_done_offline_task.id
        api_offline_task_id = api_offline_task.id
        raw_failed_offline_task_id = raw_failed_offline_task.id
        reuse_run_id = reuse_run.id
        reuse_offline_task_id = reuse_offline_task.id
        reuse_offline_step_id = reuse_offline_step.id

    transport = httpx.ASGITransport(app=app, client=("127.0.0.1", 41332))
    async with httpx.AsyncClient(transport=transport, base_url="http://testserver") as client:
        with external_socket_guard():
            task_list = await client.get(
                "/api/task-runs",
                params={"view": "all", "q": "R1 unsafe outer projection API", "page_size": 20},
            )
            assert task_list.status_code == 200, task_list.text
            task_item = next(item for item in task_list.json()["items"] if item["id"] == api_run_id)
            assert task_item["status"] == "failed", task_item
            assert task_item["display_status"] == "failed", task_item
            assert task_item["display_status_label"] == "失败", task_item
            assert "导出结果格式异常" in str(task_item["display_reason"]), task_item
            assert task_item["catalog_export_result"] == expected, task_item
            assert "download_result" not in task_item["available_actions"], task_item

            current = await client.get(
                "/api/task-runs",
                params={"view": "current", "q": "R1 unsafe outer projection API", "page_size": 20},
            )
            assert current.status_code == 200, current.text
            current_payload = current.json()
            assert {item["id"] for item in current_payload["items"]} == {
                api_run_id,
                raw_failed_run_id,
            }, current_payload
            assert current_payload["total"] == 2, current_payload
            assert current_payload["base_total"] == 2, current_payload
            assert current_payload["filtered_total"] == 2, current_payload

            history = await client.get(
                "/api/task-runs",
                params={"view": "history", "q": "R1 unsafe outer projection API", "page_size": 20},
            )
            assert history.status_code == 200, history.text
            history_payload = history.json()
            assert [item["id"] for item in history_payload["items"]] == [raw_succeeded_run_id], history_payload
            assert history_payload["total"] == 1, history_payload
            assert history_payload["base_total"] == 1, history_payload
            assert history_payload["filtered_total"] == 1, history_payload

            failed_pages = []
            for page_number in (1, 2):
                failed_page = await client.get(
                    "/api/task-runs",
                    params={
                        "view": "current",
                        "display_status": "failed",
                        "q": "R1 unsafe outer projection API",
                        "page": page_number,
                        "page_size": 1,
                    },
                )
                assert failed_page.status_code == 200, failed_page.text
                failed_payload = failed_page.json()
                assert failed_payload["total"] == 2, failed_payload
                assert failed_payload["base_total"] == 2, failed_payload
                assert failed_payload["filtered_total"] == 2, failed_payload
                assert len(failed_payload["items"]) == 1, failed_payload
                assert failed_payload["items"][0]["status"] == "failed", failed_payload
                assert failed_payload["items"][0]["display_status"] == "failed", failed_payload
                failed_pages.append(failed_payload["items"][0]["id"])
            assert set(failed_pages) == {api_run_id, raw_failed_run_id}, failed_pages

            succeeded_history = await client.get(
                "/api/task-runs",
                params={
                    "view": "history",
                    "display_status": "succeeded",
                    "q": "R1 unsafe outer projection API",
                    "page_size": 20,
                },
            )
            assert succeeded_history.status_code == 200, succeeded_history.text
            succeeded_history_payload = succeeded_history.json()
            assert [item["id"] for item in succeeded_history_payload["items"]] == [raw_succeeded_run_id], succeeded_history_payload
            assert succeeded_history_payload["total"] == 1, succeeded_history_payload
            assert succeeded_history_payload["base_total"] == 1, succeeded_history_payload
            assert succeeded_history_payload["filtered_total"] == 1, succeeded_history_payload

            task_detail = await client.get(f"/api/task-runs/{api_run_id}")
            assert task_detail.status_code == 200, task_detail.text
            task_detail_payload = task_detail.json()
            assert task_detail_payload["status"] == "failed", task_detail_payload
            assert task_detail_payload["display_status"] == "failed", task_detail_payload
            assert task_detail_payload["display_status_label"] == "失败", task_detail_payload
            assert "导出结果格式异常" in str(task_detail_payload["display_reason"]), task_detail_payload
            assert task_detail_payload["catalog_export_result"] == expected, task_detail_payload
            assert "download_result" not in task_detail_payload["available_actions"], task_detail_payload

            offline_list = await client.get(
                "/api/offline-tasks",
                params={"task_type": "catalog_export", "page_size": 100},
            )
            assert offline_list.status_code == 200, offline_list.text
            offline_item = next(
                item for item in offline_list.json()["items"] if item["id"] == api_offline_task_id
            )
            assert offline_item["status"] == "failed", offline_item
            assert offline_item["catalog_export_result"] == expected, offline_item
            assert json.loads(offline_item["result_json"]) == expected, offline_item
            assert offline_item["can_download"] is False, offline_item

            offline_detail = await client.get(f"/api/offline-tasks/{api_offline_task_id}")
            assert offline_detail.status_code == 200, offline_detail.text
            offline_detail_payload = offline_detail.json()
            assert offline_detail_payload["status"] == "failed", offline_detail_payload
            assert offline_detail_payload["catalog_export_result"] == expected, offline_detail_payload
            assert offline_detail_payload["can_download"] is False, offline_detail_payload

            offline_failed_pages = []
            for page_number in (1, 2):
                offline_failed_page = await client.get(
                    "/api/offline-tasks",
                    params={
                        "task_type": "catalog_export",
                        "status": "failed",
                        "page": page_number,
                        "page_size": 1,
                    },
                )
                assert offline_failed_page.status_code == 200, offline_failed_page.text
                offline_failed_payload = offline_failed_page.json()
                assert offline_failed_payload["total"] >= 2, offline_failed_payload
                assert len(offline_failed_payload["items"]) == 1, offline_failed_payload
                assert offline_failed_payload["items"][0]["status"] == "failed", offline_failed_payload
                offline_failed_pages.append(offline_failed_payload["items"][0]["id"])
            assert offline_failed_pages == [raw_failed_offline_task_id, api_offline_task_id], offline_failed_pages

            offline_done = await client.get(
                "/api/offline-tasks",
                params={"task_type": "catalog_export", "status": "done", "page_size": 100},
            )
            assert offline_done.status_code == 200, offline_done.text
            offline_done_payload = offline_done.json()
            offline_done_ids = {item["id"] for item in offline_done_payload["items"]}
            assert raw_done_offline_task_id in offline_done_ids, offline_done_payload
            assert api_offline_task_id not in offline_done_ids, offline_done_payload
            assert all(item["status"] == "done" for item in offline_done_payload["items"]), offline_done_payload

            export_files = await client.get("/api/products/catalog/export-files", params={"page_size": 100})
            assert export_files.status_code == 200, export_files.text
            export_items = {
                (item["task_source"], item["task_id"]): item
                for item in export_files.json()["items"]
            }
            for key in (("task_run", api_run_id), ("offline_task", api_offline_task_id)):
                export_item = export_items[key]
                assert export_item["task_status"] == "failed", export_item
                assert export_item["success_count"] == 0, export_item
                assert export_item["skipped_count"] == 0, export_item
                assert export_item["failed_count"] == 1, export_item
                assert export_item["can_download"] is False, export_item

            categories = await client.get("/api/products/catalog/export-categories")
            assert categories.status_code == 200, categories.text
            exported_categories = {item["category"] for item in categories.json()["exported"]}
            assert category not in exported_categories, exported_categories

            task_download = await client.get(f"/api/task-runs/{api_run_id}/download")
            assert task_download.status_code == 400, task_download.text
            offline_download = await client.get(f"/api/offline-tasks/{api_offline_task_id}/download")
            assert offline_download.status_code == 400, offline_download.text

    task_run_upload_calls: list[tuple[Path, str]] = []
    original_task_run_upload = catalog_export_workers.upload_private_file

    def fake_task_run_upload(path: Path, object_key: str) -> dict:
        task_run_upload_calls.append((Path(path), object_key))
        return {"object_key": object_key, "url": f"https://oss.invalid/{object_key}"}

    register_catalog_export_workers()
    catalog_export_workers.upload_private_file = fake_task_run_upload
    try:
        with external_socket_guard():
            await drain_ready_steps()
    finally:
        catalog_export_workers.upload_private_file = original_task_run_upload
    assert task_run_upload_calls == [], task_run_upload_calls

    offline_upload_calls: list[tuple[Path, str]] = []
    original_offline_upload = offline_tasks.upload_private_file

    def fake_offline_upload(path: Path, object_key: str) -> dict:
        offline_upload_calls.append((Path(path), object_key))
        return {"object_key": object_key, "url": f"https://oss.invalid/{object_key}"}

    offline_tasks.upload_private_file = fake_offline_upload
    try:
        with external_socket_guard():
            await offline_tasks._execute_offline_task(reuse_offline_task_id, [reuse_offline_step_id])
    finally:
        offline_tasks.upload_private_file = original_offline_upload
    assert offline_upload_calls == [], offline_upload_calls

    async with environment.session_factory() as db:
        processed_run = (
            await db.execute(
                select(TaskRun)
                .where(TaskRun.id == reuse_run_id)
                .options(selectinload(TaskRun.groups), selectinload(TaskRun.steps))
            )
        ).scalar_one()
        assert processed_run.status == "failed", processed_run.status
        assert processed_run.groups[0].status == "failed", processed_run.groups[0].status
        assert processed_run.steps[0].status == "failed", processed_run.steps[0].status
        processed_payload = json.loads(processed_run.summary_json or "{}")
        assert processed_payload["status"] == "failed", processed_payload
        assert processed_payload["success_count"] == 0, processed_payload
        assert processed_payload["failed_count"] == 1, processed_payload

        processed_offline_task = await db.get(OfflineTask, reuse_offline_task_id)
        processed_offline_step = await db.get(OfflineTaskStep, reuse_offline_step_id)
        assert processed_offline_task is not None and processed_offline_task.status == "failed", processed_offline_task
        assert processed_offline_step is not None and processed_offline_step.status == "failed", processed_offline_step
        assert "导出商品不存在" in str(processed_offline_step.error_message or ""), processed_offline_step


async def test_canonical_payload_contract(environment) -> None:
    from app.services.offline_tasks import _catalog_export_result_payload

    artifact = environment.data_dir / "exports" / "canonical.zip"
    artifact.parent.mkdir(parents=True, exist_ok=True)
    artifact.write_bytes(b"r1-catalog-export")
    payload = _catalog_export_result_payload(
        category="Sofas & Couches",
        categories=["Sofas & Couches"],
        template_name="CHAIR_SOFA.xlsm",
        template_path="/fixture/CHAIR_SOFA.xlsm",
        catalog_ids=[10, 11],
        report_rows=[{
            "商品资料ID": 10,
            "商品ID": 110,
            "商品Code": "R1-DONE",
            "Seller SKU": "R1-DONE",
            "类目": "Sofas & Couches",
            "状态": "已导出",
            "原因": "fixture success",
            "模板文件": "CHAIR_SOFA.xlsm",
            "导出文件": "r1.xlsx",
        }],
        created_at=datetime.now(),
        filename=artifact.name,
        file_path=str(artifact),
        file_size=artifact.stat().st_size,
    )

    assert payload["status"] == "partial_failed", payload
    assert payload["artifact_available"] is True, payload
    assert payload["requested_count"] == 2, payload
    assert payload["success_count"] == 1, payload
    assert payload["skipped_count"] == 0, payload
    assert payload["failed_count"] == 1, payload
    assert len(payload["rows"]) == 2, payload
    assert payload["requested_count"] == (
        payload["success_count"] + payload["skipped_count"] + payload["failed_count"]
    ), payload
    missing = next(row for row in payload["rows"] if row["catalog_id"] == 11)
    assert missing["status"] == "failed", missing
    assert missing["reason"] == "导出结果缺少行级结果", missing

    no_artifact = _catalog_export_result_payload(
        category="Sofas & Couches",
        categories=["Sofas & Couches"],
        template_name="CHAIR_SOFA.xlsm",
        template_path="/fixture/CHAIR_SOFA.xlsm",
        catalog_ids=[12],
        report_rows=[{
            "商品资料ID": 12,
            "商品ID": 112,
            "商品Code": "R1-NO-ARTIFACT",
            "状态": "已导出",
            "原因": "fixture success without artifact",
        }],
        created_at=datetime.now(),
    )
    assert no_artifact["status"] == "failed", no_artifact
    assert no_artifact["artifact_available"] is False, no_artifact


async def test_pre_r1_artifact_availability_contract(environment) -> None:
    from app.services.offline_tasks import _catalog_export_result_ready
    from app.task_runtime.catalog_export_status import (
        ARTIFACT_MODE_LOCAL,
        ARTIFACT_MODE_OBJECT_KEY,
        ARTIFACT_MODE_REDIRECT,
        ARTIFACT_MODE_UNAVAILABLE,
        resolve_catalog_export_artifact,
        select_catalog_export_payload,
    )

    artifact = environment.data_dir / "exports" / "pre-r1-availability.zip"
    artifact.parent.mkdir(parents=True, exist_ok=True)
    artifact.write_bytes(b"pre-r1-availability")
    stale_path = artifact.with_name("missing-pre-r1-availability.zip")

    legacy_local = {
        "status": "partial_failed",
        "filename": artifact.name,
        "file_path": str(artifact),
    }
    legacy_object_key = {
        "status": "done",
        "filename": "pre-r1-object-key.zip",
        "oss_object_key": "exports/pre-r1-object-key.zip",
    }
    legacy_oss_url = {
        "status": "done",
        "filename": "pre-r1-oss-url.zip",
        "oss_url": "https://oss.invalid/exports/pre-r1-oss-url.zip",
    }
    legacy_url_file_path = {
        "status": "done",
        "filename": "pre-r1-url-path.zip",
        "file_path": "https://oss.invalid/exports/pre-r1-url-path.zip",
    }
    stale_local = {
        "status": "done",
        "filename": stale_path.name,
        "file_path": str(stale_path),
    }

    def resolve(payload: object):
        return resolve_catalog_export_artifact(
            payload,
            allowed_export_root=environment.data_dir / "exports",
            object_cache_subdir="pre-r1-test-cache",
        )

    expected_modes = {
        "local": (legacy_local, ARTIFACT_MODE_LOCAL),
        "object_key": (legacy_object_key, ARTIFACT_MODE_OBJECT_KEY),
        "oss_url": (legacy_oss_url, ARTIFACT_MODE_REDIRECT),
        "url_file_path": (legacy_url_file_path, ARTIFACT_MODE_REDIRECT),
    }
    for payload, expected_mode in expected_modes.values():
        assert "artifact_available" not in payload, payload
        assert resolve(payload).mode == expected_mode, payload
        assert _catalog_export_result_ready(payload) is True, payload

    assert resolve(legacy_oss_url).redirect_url == legacy_oss_url["oss_url"], legacy_oss_url
    assert resolve(legacy_url_file_path).redirect_url == legacy_url_file_path["file_path"], legacy_url_file_path

    assert resolve(stale_local).mode == ARTIFACT_MODE_UNAVAILABLE, stale_local
    assert _catalog_export_result_ready(stale_local) is False, stale_local
    for explicit_value in (False, None):
        explicit_unavailable = {
            **legacy_local,
            "artifact_available": explicit_value,
            "oss_object_key": "exports/must-not-infer.zip",
            "oss_url": "https://oss.invalid/exports/must-not-infer.zip",
        }
        assert resolve(explicit_unavailable).mode == ARTIFACT_MODE_UNAVAILABLE, explicit_unavailable
        assert _catalog_export_result_ready(explicit_unavailable) is False, explicit_unavailable

    authoritative_false = {"status": "done", "artifact_available": False}
    authoritative_rows = {"status": "done", "rows": [], "requested_count": 0}
    assert select_catalog_export_payload(authoritative_false, [legacy_local]) == authoritative_false
    assert select_catalog_export_payload(json.dumps(authoritative_rows), [json.dumps(legacy_local)]) == authoritative_rows
    assert select_catalog_export_payload({"status": "done"}, [legacy_local]) == legacy_local
    assert resolve({"status": "done", "artifact_available": True}).mode == ARTIFACT_MODE_UNAVAILABLE


async def test_material_predicate_python_mysql_type_equivalence(environment) -> None:
    from sqlalchemy import event

    from app.database import engine
    from app.models import TaskGroup, TaskRun, TaskStep
    from app.task_runtime.catalog_export_status import (
        ARTIFACT_MODE_LOCAL,
        ARTIFACT_MODE_UNAVAILABLE,
        ARTIFACT_REFERENCE_FIELDS,
        CATALOG_STEP_OWNER_TASK_RUN,
        load_newest_material_catalog_step_results,
        resolve_catalog_export_artifact,
        select_catalog_export_payload,
    )

    export_root = environment.data_dir / "exports"
    export_root.mkdir(parents=True, exist_ok=True)
    older_artifact = export_root / "material-matrix-older.zip"
    older_artifact.write_bytes(b"material-matrix-older")
    newer_artifact = export_root / "material-matrix-newer.zip"
    newer_artifact.write_bytes(b"material-matrix-newer")
    older_payload = {
        "status": "done",
        "filename": older_artifact.name,
        "file_path": str(older_artifact),
    }

    mysql_depth_payload = (
        '{"status":"done","file_path":null,"nested":'
        + "[" * 150
        + "0"
        + "]" * 150
        + "}"
    )
    python_recursion_payload = (
        '{"status":"done","file_path":null,"nested":'
        + "[" * 1200
        + "0"
        + "]" * 1200
        + "}"
    )
    huge_integer = "9" * 310
    python_limit_integer = "9" * 5000
    authoritative_cases: list[tuple[str, str | None, str]] = [
        ("valid_string", json.dumps({"status": "done", "file_path": str(newer_artifact)}), ARTIFACT_MODE_LOCAL),
        ("empty_string", '{"status":"done","file_path":""}', ARTIFACT_MODE_UNAVAILABLE),
        ("whitespace_string", '{"status":"done","file_path":" \\t\\n"}', ARTIFACT_MODE_UNAVAILABLE),
        ("positive_number", '{"status":"done","file_path":123}', ARTIFACT_MODE_UNAVAILABLE),
        ("negative_number", '{"status":"done","file_path":-123}', ARTIFACT_MODE_UNAVAILABLE),
        ("zero_number", '{"status":"done","file_path":0}', ARTIFACT_MODE_UNAVAILABLE),
        ("true", '{"status":"done","file_path":true}', ARTIFACT_MODE_UNAVAILABLE),
        ("false", '{"status":"done","file_path":false}', ARTIFACT_MODE_UNAVAILABLE),
        ("object", '{"status":"done","file_path":{"private":"must-not-leak"}}', ARTIFACT_MODE_UNAVAILABLE),
        ("list", '{"status":"done","file_path":["private-must-not-leak"]}', ARTIFACT_MODE_UNAVAILABLE),
        ("null", '{"status":"done","file_path":null}', ARTIFACT_MODE_UNAVAILABLE),
        ("nan", '{"status":"done","file_path":NaN}', ARTIFACT_MODE_UNAVAILABLE),
        ("positive_infinity", '{"status":"done","file_path":Infinity}', ARTIFACT_MODE_UNAVAILABLE),
        ("negative_infinity", '{"status":"done","file_path":-Infinity}', ARTIFACT_MODE_UNAVAILABLE),
        ("overflow_float", '{"status":"done","file_path":1e309}', ARTIFACT_MODE_UNAVAILABLE),
        ("long_integer", f'{{"status":"done","file_path":{huge_integer}}}', ARTIFACT_MODE_UNAVAILABLE),
        ("python_limit_integer", f'{{"status":"done","file_path":{python_limit_integer}}}', ARTIFACT_MODE_UNAVAILABLE),
        ("lone_high_surrogate", '{"status":"done","file_path":"\\ud800"}', ARTIFACT_MODE_UNAVAILABLE),
        ("lone_low_surrogate", '{"status":"done","file_path":"\\udc00"}', ARTIFACT_MODE_UNAVAILABLE),
        ("high_surrogate_text", '{"status":"done","file_path":"\\ud800text"}', ARTIFACT_MODE_UNAVAILABLE),
        ("paired_surrogate", '{"status":"done","file_path":"\\ud83d\\ude00"}', ARTIFACT_MODE_UNAVAILABLE),
        ("mysql_depth_limit", mysql_depth_payload, ARTIFACT_MODE_UNAVAILABLE),
        ("python_recursion_limit", python_recursion_payload, ARTIFACT_MODE_UNAVAILABLE),
        ("ordinary_invalid_json", "{not-json", ARTIFACT_MODE_UNAVAILABLE),
        ("invalid_escape", '{"status":"done","file_path":"\\x"}', ARTIFACT_MODE_UNAVAILABLE),
        ("invalid_control", '{"status":"done","file_path":"line\nbreak"}', ARTIFACT_MODE_UNAVAILABLE),
        ("duplicate_key", '{"status":"done","file_path":"private","file_path":null}', ARTIFACT_MODE_UNAVAILABLE),
        ("canonical_nested_nan", '{"status":"done","rows":[{"reason":NaN}]}', ARTIFACT_MODE_UNAVAILABLE),
        ("canonical_nested_overflow_float", '{"status":"done","rows":[{"score":1e309}]}', ARTIFACT_MODE_UNAVAILABLE),
        ("canonical_nested_long_integer", f'{{"status":"done","rows":[{{"score":{huge_integer}}}]}}', ARTIFACT_MODE_UNAVAILABLE),
        ("canonical_nested_surrogate", '{"status":"done","rows":[{"reason":"\\ud800"}]}', ARTIFACT_MODE_UNAVAILABLE),
        ("artifact_false", '{"status":"done","artifact_available":false}', ARTIFACT_MODE_UNAVAILABLE),
        ("artifact_null", '{"status":"done","artifact_available":null}', ARTIFACT_MODE_UNAVAILABLE),
        ("canonical_rows_null", '{"status":"done","rows":null}', ARTIFACT_MODE_UNAVAILABLE),
        ("canonical_count_zero", '{"status":"done","requested_count":0}', ARTIFACT_MODE_UNAVAILABLE),
    ]
    for field in ARTIFACT_REFERENCE_FIELDS:
        if field != "file_path":
            authoritative_cases.append((
                f"{field}_null_presence",
                json.dumps({"status": "done", field: None}),
                ARTIFACT_MODE_UNAVAILABLE,
            ))

    fallback_cases: list[tuple[str, str | None]] = [
        ("none", None),
        ("empty", ""),
        ("whitespace", " \t\n"),
        ("status_only", '{"status":"done"}'),
        ("payload_scalar_string", '"done"'),
        ("payload_scalar_number", "123"),
        ("payload_scalar_true", "true"),
        ("payload_scalar_null", "null"),
        ("payload_array", '[{"file_path":"private-array-value"}]'),
    ]
    invalid_fallback_names = {
        "nan",
        "positive_infinity",
        "negative_infinity",
        "overflow_float",
        "python_limit_integer",
        "python_recursion_limit",
        "ordinary_invalid_json",
        "invalid_escape",
        "invalid_control",
        "canonical_nested_nan",
        "canonical_nested_overflow_float",
    }
    fallback_cases.extend(
        (name, raw)
        for name, raw, _expected_mode in authoritative_cases
        if name in invalid_fallback_names
    )
    authoritative_cases = [
        item
        for item in authoritative_cases
        if item[0] not in invalid_fallback_names
    ]

    run_ids: dict[str, int] = {}
    expected_modes: dict[str, str] = {}
    older_result_json = json.dumps(older_payload)
    async with environment.session_factory() as db:
        for name, newer_result_json, expected_mode in authoritative_cases:
            expected_modes[name] = expected_mode
            run = TaskRun(task_type="catalog_export", title=f"R1 material matrix {name}", status="succeeded")
            db.add(run)
            await db.flush()
            group = TaskGroup(
                task_run_id=run.id,
                group_key="export_file",
                title=f"R1 material matrix {name}",
                status="succeeded",
            )
            db.add(group)
            await db.flush()
            db.add(TaskStep(
                task_run_id=run.id,
                task_group_id=group.id,
                step_key="catalog_export_template_older",
                step_type="catalog_export_template",
                status="succeeded",
                result_json=older_result_json,
            ))
            await db.flush()
            db.add(TaskStep(
                task_run_id=run.id,
                task_group_id=group.id,
                step_key="catalog_export_template_newer",
                step_type="catalog_export_template",
                status="succeeded",
                result_json=newer_result_json,
            ))
            await db.flush()
            run_ids[name] = run.id

        for name, newer_result_json in fallback_cases:
            expected_modes[name] = ARTIFACT_MODE_LOCAL
            run = TaskRun(task_type="catalog_export", title=f"R1 material matrix {name}", status="succeeded")
            db.add(run)
            await db.flush()
            group = TaskGroup(
                task_run_id=run.id,
                group_key="export_file",
                title=f"R1 material matrix {name}",
                status="succeeded",
            )
            db.add(group)
            await db.flush()
            db.add(TaskStep(
                task_run_id=run.id,
                task_group_id=group.id,
                step_key="catalog_export_template_older",
                step_type="catalog_export_template",
                status="succeeded",
                result_json=older_result_json,
            ))
            await db.flush()
            db.add(TaskStep(
                task_run_id=run.id,
                task_group_id=group.id,
                step_key="catalog_export_template_newer",
                step_type="catalog_export_template",
                status="succeeded",
                result_json=newer_result_json,
            ))
            await db.flush()
            run_ids[name] = run.id
        await db.commit()

    recorded_queries: list[tuple[str, object]] = []

    def record_query(_connection, _cursor, statement, parameters, _context, _executemany):
        normalized = " ".join(statement.lower().split())
        if "task_steps.task_run_id" in normalized and "task_steps.result_json" in normalized:
            recorded_queries.append((normalized, parameters))

    event.listen(engine.sync_engine, "before_cursor_execute", record_query)
    async with environment.session_factory() as db:
        try:
            sql_selected = await load_newest_material_catalog_step_results(
                db,
                owner_kind=CATALOG_STEP_OWNER_TASK_RUN,
                owner_ids=run_ids.values(),
            )
        finally:
            event.remove(engine.sync_engine, "before_cursor_execute", record_query)
    assert len(recorded_queries) == 1, recorded_queries
    statement, parameters = recorded_queries[0]
    assert "json_" not in statement, statement
    assert "row_number()" not in statement, statement
    assert "task_steps.task_run_id" in statement, statement
    assert "task_steps.id" in statement, statement
    assert "task_steps.result_json" in statement, statement
    assert "order by task_steps.task_run_id asc, task_steps.id desc" in statement, statement
    parameter_values = parameters.values() if isinstance(parameters, dict) else parameters
    queried_owner_ids = {
        value
        for value in parameter_values
        if isinstance(value, int) and not isinstance(value, bool)
    }
    assert queried_owner_ids == set(run_ids.values()), (queried_owner_ids, run_ids)
    assert set(sql_selected) == set(run_ids.values()), sql_selected
    authoritative_names = {name for name, _raw, _mode in authoritative_cases}
    raw_by_name = {
        name: raw
        for name, raw, _mode in authoritative_cases
    } | {
        name: raw
        for name, raw in fallback_cases
    }
    for name, run_id in run_ids.items():
        python_selected = select_catalog_export_payload({}, [raw_by_name[name], older_result_json])
        selected_payload = sql_selected[run_id]
        assert isinstance(selected_payload, dict), (name, selected_payload)
        python_signature = json.dumps(python_selected, sort_keys=True, ensure_ascii=True, allow_nan=False)
        sql_signature = json.dumps(selected_payload, sort_keys=True, ensure_ascii=True, allow_nan=False)
        assert sql_signature == python_signature, (name, selected_payload, python_selected)
        sql_signature.encode("utf-8")
        if name in authoritative_names:
            assert selected_payload != older_payload, (name, selected_payload)
        else:
            assert selected_payload == older_payload, (name, selected_payload)
        resolution = resolve_catalog_export_artifact(
            selected_payload,
            allowed_export_root=export_root,
            object_cache_subdir="material-matrix-cache",
        )
        assert resolution.mode == expected_modes[name], (name, resolution, selected_payload)


async def test_invalid_step_payload_falls_back_across_apis(environment) -> None:
    import httpx
    from sqlalchemy import event

    from app.database import engine
    from app.main import app
    from app.models import OfflineTask, OfflineTaskStep, TaskGroup, TaskRun, TaskStep
    from app.task_runtime.catalog_export_status import normalize_catalog_export_response

    export_root = environment.data_dir / "exports"
    export_root.mkdir(parents=True, exist_ok=True)
    older_artifact = export_root / "parser-api-older.zip"
    older_artifact.write_bytes(b"parser-api-older")
    task_category = "R1 Parser Barrier TaskRun Older Category"
    offline_category = "R1 Parser Barrier Offline Older Category"

    def older_payload(category: str) -> dict:
        return {
            "status": "done",
            "category": category,
            "categories": [category],
            "filename": older_artifact.name,
            "file_path": str(older_artifact),
        }

    api_huge_integer = "9" * 310
    canonical_unsafe_raw = (
        '{"status":"done","rows":[{"category":"R1 Parser Unsafe Canonical",'
        f'"reason":"\\ud800","score":{api_huge_integer}}}]}}'
    )
    canonical_overflow_raw = '{"status":"done","rows":[{"score":1e309}]}'
    mysql_depth_raw = (
        '{"status":"done","file_path":null,"nested":'
        + "[" * 1200
        + "0"
        + "]" * 150
        + "}"
    )

    run_ids: dict[str, int] = {}
    async with environment.session_factory() as db:
        for name, newer_raw in (
            ("invalid", "{not-json"),
            ("canonical_unsafe", canonical_unsafe_raw),
            ("canonical_overflow", canonical_overflow_raw),
        ):
            run = TaskRun(
                task_type="catalog_export",
                title=f"R1 parser API TaskRun {name}",
                status="succeeded",
                summary_json=json.dumps({"status": "done"}),
            )
            db.add(run)
            await db.flush()
            run_ids[name] = run.id
            group = TaskGroup(
                task_run_id=run.id,
                group_key="export_file",
                title=f"R1 parser API TaskRun {name}",
                status="succeeded",
            )
            db.add(group)
            await db.flush()
            db.add(TaskStep(
                task_run_id=run.id,
                task_group_id=group.id,
                step_key="catalog_export_template_older",
                step_type="catalog_export_template",
                status="succeeded",
                result_json=json.dumps(older_payload(task_category)),
            ))
            await db.flush()
            db.add(TaskStep(
                task_run_id=run.id,
                task_group_id=group.id,
                step_key="catalog_export_template_newer",
                step_type="catalog_export_template",
                status="succeeded",
                result_json=newer_raw,
            ))
            await db.flush()

        offline_task = OfflineTask(
            task_type="catalog_export",
            title="R1 parser API OfflineTask depth",
            status="done",
            total_steps=2,
            success_steps=2,
            result_json=json.dumps({"status": "done"}),
        )
        db.add(offline_task)
        await db.flush()
        db.add(OfflineTaskStep(
            task_id=offline_task.id,
            step_type="catalog_export_template",
            title="R1 parser API OfflineTask older",
            status="done",
            progress_current=1,
            progress_total=1,
            result_json=json.dumps(older_payload(offline_category)),
        ))
        await db.flush()
        db.add(OfflineTaskStep(
            task_id=offline_task.id,
            step_type="catalog_export_template",
            title="R1 parser API OfflineTask newer depth",
            status="done",
            progress_current=1,
            progress_total=1,
            result_json=mysql_depth_raw,
        ))
        await db.commit()
        offline_task_id = offline_task.id

    def batch_queries(records: list[tuple[str, object]], table: str, owner_column: str):
        return [
            item
            for item in records
            if f"{table}.{owner_column}" in item[0] and f"{table}.result_json" in item[0]
        ]

    def assert_single_batch_query(
        records: list[tuple[str, object]],
        *,
        table: str,
        owner_column: str,
        expected_owner_ids: set[int],
    ) -> None:
        matches = batch_queries(records, table, owner_column)
        assert len(matches) == 1, (table, records)
        statement, parameters = matches[0]
        assert "json_" not in statement, statement
        assert "row_number()" not in statement, statement
        assert f"order by {table}.{owner_column} asc, {table}.id desc" in statement, statement
        assert "events" not in statement and "groups" not in statement, statement
        parameter_values = parameters.values() if isinstance(parameters, dict) else parameters
        queried_owner_ids = {
            value
            for value in parameter_values
            if isinstance(value, int) and not isinstance(value, bool)
        }
        assert expected_owner_ids.issubset(queried_owner_ids), (table, queried_owner_ids, expected_owner_ids)

    transport = httpx.ASGITransport(app=app, client=("127.0.0.1", 41329))
    async with httpx.AsyncClient(
        transport=transport,
        base_url="http://testserver",
        follow_redirects=False,
    ) as client:
        with external_socket_guard():
            task_list_queries: list[tuple[str, object]] = []

            def record_task_list(_connection, _cursor, statement, parameters, _context, _executemany):
                task_list_queries.append((" ".join(statement.lower().split()), parameters))

            event.listen(engine.sync_engine, "before_cursor_execute", record_task_list)
            try:
                task_list = await client.get(
                    "/api/task-runs",
                    params={"view": "all", "q": "R1 parser API TaskRun", "page_size": 20},
                )
            finally:
                event.remove(engine.sync_engine, "before_cursor_execute", record_task_list)
            assert task_list.status_code == 200, task_list.text
            task_items = {item["id"]: item for item in task_list.json()["items"]}
            assert set(task_items) == set(run_ids.values()), task_items
            assert_single_batch_query(
                task_list_queries,
                table="task_steps",
                owner_column="task_run_id",
                expected_owner_ids=set(run_ids.values()),
            )
            for name in ("invalid", "canonical_overflow"):
                item = task_items[run_ids[name]]
                assert item["status"] == "succeeded" and item["display_status"] == "succeeded", item
                assert "download_result" in item["available_actions"], item
                assert item["catalog_export_result"]["status"] == "done", item
                assert item["catalog_export_result"]["category"] == task_category, item
            canonical_payload = task_items[run_ids["canonical_unsafe"]]["catalog_export_result"]
            assert task_items[run_ids["canonical_unsafe"]]["status"] == "failed", task_items
            assert task_items[run_ids["canonical_unsafe"]]["display_status"] == "failed", task_items
            assert "download_result" not in task_items[run_ids["canonical_unsafe"]]["available_actions"]
            assert canonical_payload["rows"][0]["reason"] == "导出结果行格式异常", canonical_payload
            assert "score" not in canonical_payload["rows"][0], canonical_payload

            for name, run_id in run_ids.items():
                detail = await client.get(f"/api/task-runs/{run_id}")
                assert detail.status_code == 200, (name, detail.text)
                detail_payload = detail.json()["catalog_export_result"]
                if name in {"invalid", "canonical_overflow"}:
                    assert detail.json()["status"] == "succeeded", detail.json()
                    assert detail_payload["status"] == "done", detail_payload
                    assert detail_payload["category"] == task_category, detail_payload
                else:
                    assert detail.json()["status"] == "failed", detail.json()
                    assert detail_payload["rows"][0]["reason"] == "导出结果行格式异常", detail_payload
                download = await client.get(f"/api/task-runs/{run_id}/download")
                if name in {"invalid", "canonical_overflow"}:
                    assert download.status_code == 200 and download.content == older_artifact.read_bytes(), name
                else:
                    assert download.status_code == 400, (name, download.status_code, download.text)

            offline_list_queries: list[tuple[str, object]] = []

            def record_offline_list(_connection, _cursor, statement, parameters, _context, _executemany):
                offline_list_queries.append((" ".join(statement.lower().split()), parameters))

            event.listen(engine.sync_engine, "before_cursor_execute", record_offline_list)
            try:
                offline_list = await client.get(
                    "/api/offline-tasks",
                    params={"task_type": "catalog_export", "page_size": 100},
                )
            finally:
                event.remove(engine.sync_engine, "before_cursor_execute", record_offline_list)
            assert offline_list.status_code == 200, offline_list.text
            offline_items = {item["id"]: item for item in offline_list.json()["items"]}
            assert offline_items[offline_task_id]["status"] == "done", offline_items[offline_task_id]
            assert offline_items[offline_task_id]["can_download"] is True, offline_items[offline_task_id]
            assert offline_items[offline_task_id]["catalog_export_result"]["category"] == offline_category
            assert_single_batch_query(
                offline_list_queries,
                table="offline_task_steps",
                owner_column="task_id",
                expected_owner_ids=set(offline_items),
            )
            offline_detail = await client.get(f"/api/offline-tasks/{offline_task_id}")
            assert offline_detail.status_code == 200, offline_detail.text
            assert offline_detail.json()["status"] == "done", offline_detail.json()
            assert offline_detail.json()["can_download"] is True, offline_detail.json()
            assert offline_detail.json()["catalog_export_result"]["category"] == offline_category
            offline_download = await client.get(f"/api/offline-tasks/{offline_task_id}/download")
            assert offline_download.status_code == 200, offline_download.text
            assert offline_download.content == older_artifact.read_bytes()

            export_queries: list[tuple[str, object]] = []

            def record_exports(_connection, _cursor, statement, parameters, _context, _executemany):
                export_queries.append((" ".join(statement.lower().split()), parameters))

            event.listen(engine.sync_engine, "before_cursor_execute", record_exports)
            try:
                export_files = await client.get("/api/products/catalog/export-files", params={"page_size": 100})
            finally:
                event.remove(engine.sync_engine, "before_cursor_execute", record_exports)
            assert export_files.status_code == 200, export_files.text
            export_items = {
                (item["task_source"], item["task_id"]): item
                for item in export_files.json()["items"]
            }
            for name in ("invalid", "canonical_overflow"):
                assert export_items[("task_run", run_ids[name])]["can_download"] is True, export_items
                assert export_items[("task_run", run_ids[name])]["task_status"] == "succeeded", export_items
            assert export_items[("task_run", run_ids["canonical_unsafe"])]["can_download"] is False, export_items
            assert export_items[("task_run", run_ids["canonical_unsafe"])]["task_status"] == "failed", export_items
            assert export_items[("offline_task", offline_task_id)]["can_download"] is True, export_items
            assert export_items[("offline_task", offline_task_id)]["task_status"] == "done", export_items
            export_task_step_queries = batch_queries(export_queries, "task_steps", "task_run_id")
            export_offline_step_queries = batch_queries(export_queries, "offline_task_steps", "task_id")
            assert len(export_task_step_queries) == 1, export_queries
            assert len(export_offline_step_queries) == 1, export_queries
            assert all(
                "json_" not in statement
                for statement, _params in (*export_task_step_queries, *export_offline_step_queries)
            ), export_queries

            category_queries: list[tuple[str, object]] = []

            def record_categories(_connection, _cursor, statement, parameters, _context, _executemany):
                category_queries.append((" ".join(statement.lower().split()), parameters))

            event.listen(engine.sync_engine, "before_cursor_execute", record_categories)
            try:
                categories = await client.get("/api/products/catalog/export-categories")
            finally:
                event.remove(engine.sync_engine, "before_cursor_execute", record_categories)
            assert categories.status_code == 200, categories.text
            exported_categories = {item["category"] for item in categories.json()["exported"]}
            assert task_category in exported_categories, exported_categories
            assert offline_category in exported_categories, exported_categories
            category_task_step_queries = batch_queries(category_queries, "task_steps", "task_run_id")
            category_offline_step_queries = batch_queries(category_queries, "offline_task_steps", "task_id")
            assert len(category_task_step_queries) == 1, category_queries
            assert len(category_offline_step_queries) == 1, category_queries
            assert all(
                "json_" not in statement
                for statement, _params in (*category_task_step_queries, *category_offline_step_queries)
            ), category_queries

            response_payloads = [task_list.json(), offline_list.json(), export_files.json(), categories.json()]
            for payload in response_payloads:
                json.dumps(payload, ensure_ascii=False, allow_nan=False).encode("utf-8")


async def test_export_center_shared_effective_projection_contract(environment) -> None:
    import httpx
    from sqlalchemy import event

    from app.database import engine
    from app.main import app
    from app.models import OfflineTask, TaskRun

    export_root = environment.data_dir / "exports"
    export_root.mkdir(parents=True, exist_ok=True)
    artifact = export_root / "shared-effective-projection.zip"
    artifact.write_bytes(b"shared-effective-projection")

    def safe_payload(category: str) -> dict:
        return {
            "status": "done",
            "category": category,
            "categories": [category],
            "filename": artifact.name,
            "file_path": str(artifact),
        }

    def unsafe_payload(category: str) -> dict:
        return {
            **safe_payload(category),
            "artifact_available": True,
            "rows": [
                {
                    "catalog_id": 2 ** 53,
                    "category": category,
                    "status": "exported",
                }
            ],
        }

    task_categories = {
        "interrupted_material": "R1 Shared Projection TaskRun Interrupted",
        "canceled_material": "R1 Shared Projection TaskRun Canceled",
        "paused_material": "R1 Shared Projection TaskRun Paused",
        "failed_safe": "R1 Shared Projection TaskRun Failed Safe",
        "failed_unsafe": "R1 Shared Projection TaskRun Failed Unsafe",
    }
    offline_categories = {
        "paused_material": "R1 Shared Projection Offline Paused",
        "failed_safe": "R1 Shared Projection Offline Failed Safe",
        "failed_unsafe": "R1 Shared Projection Offline Failed Unsafe",
    }
    run_ids: dict[str, int] = {}
    offline_ids: dict[str, int] = {}
    async with environment.session_factory() as db:
        for name, status in (
            ("interrupted_material", "interrupted"),
            ("canceled_material", "canceled"),
            ("paused_material", "paused"),
            ("failed_safe", "failed"),
            ("failed_empty", "failed"),
            ("succeeded_invalid", "succeeded"),
            ("failed_unsafe", "failed"),
        ):
            if name == "failed_empty":
                summary_json = None
            elif name == "succeeded_invalid":
                summary_json = "{not-json"
            elif name == "failed_unsafe":
                summary_json = json.dumps(unsafe_payload(task_categories[name]))
            else:
                summary_json = json.dumps(safe_payload(task_categories[name]))
            run = TaskRun(
                task_type="catalog_export",
                title=f"R1 shared projection TaskRun {name}",
                status=status,
                summary_json=summary_json,
            )
            db.add(run)
            await db.flush()
            run_ids[name] = run.id

        for name, status in (
            ("paused_material", "paused"),
            ("failed_safe", "failed"),
            ("failed_empty", "failed"),
            ("done_invalid", "done"),
            ("failed_unsafe", "failed"),
        ):
            if name == "failed_empty":
                result_json = None
            elif name == "done_invalid":
                result_json = "{not-json"
            elif name == "failed_unsafe":
                result_json = json.dumps(unsafe_payload(offline_categories[name]))
            else:
                result_json = json.dumps(safe_payload(offline_categories[name]))
            task = OfflineTask(
                task_type="catalog_export",
                title=f"R1 shared projection OfflineTask {name}",
                status=status,
                result_json=result_json,
            )
            db.add(task)
            await db.flush()
            offline_ids[name] = task.id
        await db.commit()

    def projection_scalar_queries(records: list[tuple[str, object]], *, table: str, payload_column: str):
        return [
            item
            for item in records
            if f"select {table}.id, {table}.status, {table}.{payload_column}" in item[0]
            and f"{table}.task_type" in item[0]
        ]

    def assert_shared_projection_queries(records: list[tuple[str, object]]) -> None:
        assert len(projection_scalar_queries(
            records,
            table="task_runs",
            payload_column="summary_json",
        )) == 1, records
        assert len(projection_scalar_queries(
            records,
            table="offline_tasks",
            payload_column="result_json",
        )) == 1, records
        assert sum(
            "task_steps.task_run_id" in statement and "task_steps.result_json" in statement
            for statement, _parameters in records
        ) == 1, records
        assert sum(
            "offline_task_steps.task_id" in statement and "offline_task_steps.result_json" in statement
            for statement, _parameters in records
        ) == 1, records
        assert not any(
            "select task_runs.id, task_runs.summary_json" in statement
            for statement, _parameters in records
        ), records
        assert all("json_" not in statement for statement, _parameters in records), records

    transport = httpx.ASGITransport(app=app, client=("127.0.0.1", 41330))
    async with httpx.AsyncClient(
        transport=transport,
        base_url="http://testserver",
        follow_redirects=False,
    ) as client:
        with external_socket_guard():
            task_list = await client.get(
                "/api/task-runs",
                params={"view": "all", "q": "R1 shared projection TaskRun", "page_size": 20},
            )
            assert task_list.status_code == 200, task_list.text
            task_items = {item["id"]: item for item in task_list.json()["items"]}
            for name, expected_status in (
                ("interrupted_material", "interrupted"),
                ("canceled_material", "canceled"),
                ("paused_material", "paused"),
                ("failed_safe", "succeeded"),
                ("failed_empty", "failed"),
                ("succeeded_invalid", "succeeded"),
                ("failed_unsafe", "failed"),
            ):
                assert task_items[run_ids[name]]["status"] == expected_status, (name, task_items)

            offline_list = await client.get(
                "/api/offline-tasks",
                params={"task_type": "catalog_export", "page_size": 100},
            )
            assert offline_list.status_code == 200, offline_list.text
            offline_items = {item["id"]: item for item in offline_list.json()["items"]}
            for name, expected_status in (
                ("paused_material", "paused"),
                ("failed_safe", "done"),
                ("failed_empty", "failed"),
                ("done_invalid", "done"),
                ("failed_unsafe", "failed"),
            ):
                assert offline_items[offline_ids[name]]["status"] == expected_status, (name, offline_items)

            export_queries: list[tuple[str, object]] = []

            def record_exports(_connection, _cursor, statement, parameters, _context, _executemany):
                export_queries.append((" ".join(statement.lower().split()), parameters))

            event.listen(engine.sync_engine, "before_cursor_execute", record_exports)
            try:
                export_files = await client.get(
                    "/api/products/catalog/export-files",
                    params={"page_size": 100},
                )
            finally:
                event.remove(engine.sync_engine, "before_cursor_execute", record_exports)
            assert export_files.status_code == 200, export_files.text
            assert_shared_projection_queries(export_queries)
            export_items = {
                (item["task_source"], item["task_id"]): item
                for item in export_files.json()["items"]
            }
            for name in ("interrupted_material", "canceled_material", "paused_material"):
                assert ("task_run", run_ids[name]) not in export_items, (name, export_items)
            assert ("offline_task", offline_ids["paused_material"]) not in export_items, export_items
            for source, owner_ids, expected_statuses in (
                (
                    "task_run",
                    run_ids,
                    {
                        "failed_safe": ("succeeded", True),
                        "failed_empty": ("failed", False),
                        "succeeded_invalid": ("succeeded", False),
                        "failed_unsafe": ("failed", False),
                    },
                ),
                (
                    "offline_task",
                    offline_ids,
                    {
                        "failed_safe": ("done", True),
                        "failed_empty": ("failed", False),
                        "done_invalid": ("done", False),
                        "failed_unsafe": ("failed", False),
                    },
                ),
            ):
                for name, (expected_status, expected_downloadable) in expected_statuses.items():
                    item = export_items[(source, owner_ids[name])]
                    assert item["task_status"] == expected_status, (source, name, item)
                    assert item["can_download"] is expected_downloadable, (source, name, item)

            category_queries: list[tuple[str, object]] = []

            def record_categories(_connection, _cursor, statement, parameters, _context, _executemany):
                category_queries.append((" ".join(statement.lower().split()), parameters))

            event.listen(engine.sync_engine, "before_cursor_execute", record_categories)
            try:
                categories = await client.get("/api/products/catalog/export-categories")
            finally:
                event.remove(engine.sync_engine, "before_cursor_execute", record_categories)
            assert categories.status_code == 200, categories.text
            assert_shared_projection_queries(category_queries)
            exported_categories = {item["category"] for item in categories.json()["exported"]}
            assert task_categories["failed_safe"] in exported_categories, exported_categories
            assert offline_categories["failed_safe"] in exported_categories, exported_categories
            assert all(
                category not in exported_categories
                for category in (
                    task_categories["interrupted_material"],
                    task_categories["canceled_material"],
                    task_categories["paused_material"],
                    task_categories["failed_unsafe"],
                    offline_categories["paused_material"],
                    offline_categories["failed_unsafe"],
                )
            ), exported_categories

            for source, owner_id, expected_status in (
                ("task-runs", run_ids["failed_safe"], 200),
                ("task-runs", run_ids["failed_empty"], 400),
                ("task-runs", run_ids["failed_unsafe"], 400),
                ("task-runs", run_ids["interrupted_material"], 400),
                ("offline-tasks", offline_ids["failed_safe"], 200),
                ("offline-tasks", offline_ids["failed_empty"], 400),
                ("offline-tasks", offline_ids["failed_unsafe"], 400),
                ("offline-tasks", offline_ids["paused_material"], 400),
            ):
                response = await client.get(f"/api/{source}/{owner_id}/download")
                assert response.status_code == expected_status, (source, owner_id, response.text)


def test_worker_outcome_contract_types() -> None:
    from app.task_runtime.constants import (
        RETRYABLE_STEP_STATUSES,
        STEP_STATUS_PARTIAL_FAILED,
        TERMINAL_STEP_STATUSES,
    )
    from app.task_runtime.registry import TaskOutcomeContractError, TaskWorkerOutcome

    outcome = TaskWorkerOutcome(
        payload={"status": "partial_failed"},
        terminal_status="partial_failed",
        event_type="warning",
        event_message="部分商品导出失败",
        propagate_single_step_run=True,
    )
    assert outcome.payload["status"] == "partial_failed"
    assert STEP_STATUS_PARTIAL_FAILED in TERMINAL_STEP_STATUSES
    assert STEP_STATUS_PARTIAL_FAILED not in RETRYABLE_STEP_STATUSES
    assert issubclass(TaskOutcomeContractError, RuntimeError)


async def test_progress_update_is_flush_only(environment) -> None:
    from sqlalchemy import func, select

    from app.models import TaskGroup, TaskRun, TaskStep, TaskStepEvent
    from app.task_runtime.events import update_step_progress_in_session

    async with environment.session_factory() as db:
        run = TaskRun(task_type="catalog_export", title="R1 flush-only", status="running")
        db.add(run)
        await db.flush()
        group = TaskGroup(task_run_id=run.id, group_key="export_file", title="导出文件", status="running")
        db.add(group)
        await db.flush()
        step = TaskStep(
            task_run_id=run.id,
            task_group_id=group.id,
            step_key="catalog_export_template",
            step_type="catalog_export_template",
            status="running",
            progress_current=0,
            progress_total=2,
        )
        db.add(step)
        await db.commit()
        step_id = step.id

    async with environment.session_factory() as db:
        step = await db.get(TaskStep, step_id)
        assert step is not None
        await update_step_progress_in_session(
            db,
            step,
            current=1,
            total=2,
            message="flush only",
            data={"fixture": True},
        )
        local_events = await db.scalar(
            select(func.count(TaskStepEvent.id)).where(TaskStepEvent.task_step_id == step_id)
        )
        assert local_events == 1
        async with environment.session_factory() as observer:
            visible_events = await observer.scalar(
                select(func.count(TaskStepEvent.id)).where(TaskStepEvent.task_step_id == step_id)
            )
            visible_step = await observer.get(TaskStep, step_id)
            assert visible_events == 0
            assert visible_step is not None and visible_step.progress_current == 0
        await db.rollback()

    async with environment.session_factory() as db:
        persisted_events = await db.scalar(
            select(func.count(TaskStepEvent.id)).where(TaskStepEvent.task_step_id == step_id)
        )
        persisted_step = await db.get(TaskStep, step_id)
        assert persisted_events == 0
        assert persisted_step is not None and persisted_step.progress_current == 0


async def _seed_catalog_export_products(environment, specs: list[dict]) -> list[int]:
    from app.models import CatalogProduct, GigaInventory, GigaSyncBatch, Product, ProductData, ProductImage
    from app.pipeline.step10_amazon_template import AMAZON_TEMPLATE_LOGIC_VERSION, SEMANTIC_DROPDOWN_FIELD_KEYS

    template_path = ROOT / "backend" / "app" / "pipeline" / "templates" / "CHAIR_SOFA.xlsm"
    assert template_path.is_file(), template_path
    semantic_fields = {
        ("target_audience_keyword" if key == "target_audience" else key): {
            "values": [],
            "reason": "R1 deterministic fixture",
            "source": "fixture",
        }
        for key in (*SEMANTIC_DROPDOWN_FIELD_KEYS, "fabric_type")
    }
    batch_id = f"r1-catalog-{len(specs)}-{datetime.now().strftime('%H%M%S%f')}"
    catalog_ids: list[int] = []
    inventories: list[GigaInventory] = []
    async with environment.session_factory() as db:
        db.add(
            GigaSyncBatch(
                task_id=batch_id,
                batch_id=batch_id,
                site="US",
                status="done",
                inventory_count=len(specs),
                finished_at=datetime.now(),
            )
        )
        for index, spec in enumerate(specs, start=1):
            item_code = str(spec.get("item_code") or f"R1-CATALOG-{index}-{batch_id[-6:]}")
            upc = spec.get("upc") if "upc" in spec else f"{725000000000 + index}"
            product = Product(
                gigab2b_url=f"https://example.invalid/{item_code}",
                gigab2b_product_id=item_code,
                brand="Vindhvisk",
                source_site="US",
                status="completed",
                current_step=6,
                upc=upc,
                amazon_asin=spec.get("amazon_asin"),
                competitor_asin="B0R1COMPETE",
            )
            product.data = ProductData(
                item_code=item_code,
                title="R1 Modular Sectional Sofa",
                product_type="SOFA",
                leaf_category="Sofas & Couches",
                categories=json.dumps(["Sofas & Couches"]),
                material="Polyester",
                color="Grey",
                features=json.dumps(["Modular", "Sectional", "Living room"]),
                description="R1 deterministic catalog export fixture.",
                dimension_length=80,
                dimension_width=35,
                dimension_height=32,
                weight=120,
                listing_title=f"{item_code} Modular Sectional Sofa",
                listing_bullets=json.dumps([
                    "Modular sectional design",
                    "Durable polyester upholstery",
                    "Comfortable living room seating",
                    "Sturdy furniture construction",
                    "Easy assembly with included instructions",
                ]),
                listing_search_terms="modular sofa,sectional couch,living room furniture",
                listing_description="A modular sectional sofa for living room seating.",
                suggested_price=499.99,
                amazon_template_path=str(template_path),
                amazon_template_fill_summary=json.dumps({"logic_version": AMAZON_TEMPLATE_LOGIC_VERSION}),
                listing_check=json.dumps({"amazon_template_fields": semantic_fields}),
                gigab2b_raw_snapshot=json.dumps({"site": "US"}),
            )
            product.images = ProductImage(main_image_path=f"/fixture/{item_code}.jpg")
            product.catalog_item = CatalogProduct(
                gigab2b_url=product.gigab2b_url,
                gigab2b_product_id=item_code,
                brand="Vindhvisk",
                item_code=item_code,
                title=product.data.title,
                leaf_category="Sofas & Couches",
                status="completed",
                confirmed_at=datetime.now(),
                amazon_asin=spec.get("amazon_asin"),
                upc=product.upc,
            )
            db.add(product)
            await db.flush()
            catalog_ids.append(product.catalog_item.id)
            inventories.append(
                GigaInventory(
                    batch_id=batch_id,
                    site="US",
                    sku_code=item_code,
                    stock_qty=5,
                    seller_available_inventory=5,
                    total_buyer_available_inventory=5,
                    availability_status="available",
                )
            )
        db.add_all(inventories)
        await db.commit()
    return catalog_ids


async def _prepare_catalog_run(environment, catalog_ids: list[int]) -> tuple[int, int]:
    from sqlalchemy import select

    from app.api.schemas import OfflineTaskCatalogExportRequest
    from app.models import TaskStep
    from app.task_planners.catalog_export import create_catalog_export_runs

    async with environment.session_factory() as db:
        runs, errors = await create_catalog_export_runs(
            db,
            OfflineTaskCatalogExportRequest(catalog_product_ids=catalog_ids),
            created_by="r1_catalog_test",
            auto_start=False,
        )
        assert not errors, errors
        assert len(runs) == 1, runs
        run_id = runs[0].id
        step = (
            await db.execute(select(TaskStep).where(TaskStep.task_run_id == run_id))
        ).scalar_one()
        step.status = "ready"
        await db.commit()
        return run_id, step.id


async def test_sync_catalog_export_api_commits_builder_mutations(environment) -> None:
    import httpx
    from openpyxl import Workbook
    from sqlalchemy import select

    from app.main import app
    from app.models import CatalogProduct, Product, ProductData, UpcPoolItem
    from app.pipeline.step10_amazon_template import AMAZON_TEMPLATE_LOGIC_VERSION

    catalog_id = (await _seed_catalog_export_products(environment, [
        {"item_code": "R1-SYNC-EXPORT-COMMIT", "upc": None},
    ]))[0]
    async with environment.session_factory() as db:
        pool_item = UpcPoolItem(upc="725999999903", status="available", source="r1_fixture")
        db.add(pool_item)
        await db.commit()
        pool_item_id = pool_item.id

    transport = httpx.ASGITransport(app=app, client=("127.0.0.1", 41322))
    async with httpx.AsyncClient(transport=transport, base_url="http://testserver") as client:
        with external_socket_guard():
            response = await client.post("/api/products/catalog/export", json=[catalog_id])
    assert response.status_code == 200, response.text
    assert response.content.startswith(b"PK"), response.content[:20]

    async with environment.session_factory() as db:
        catalog = await db.get(CatalogProduct, catalog_id)
        product = await db.get(Product, catalog.source_product_id if catalog else 0)
        pool_item = await db.get(UpcPoolItem, pool_item_id)
        assert catalog is not None and catalog.upc == "725999999903", catalog
        assert product is not None and product.upc == "725999999903", product
        assert pool_item is not None and pool_item.status == "bound", pool_item
        assert pool_item.product_id == product.id, pool_item

    failed_catalog_id = (await _seed_catalog_export_products(environment, [
        {"item_code": "R1-SYNC-EXPORT-ROLLBACK", "upc": None},
    ]))[0]
    invalid_source = environment.data_dir / "exports" / "invalid-cached-template.xlsm"
    invalid_source.parent.mkdir(parents=True, exist_ok=True)
    Workbook().save(invalid_source)
    async with environment.session_factory() as db:
        failed_catalog = await db.get(CatalogProduct, failed_catalog_id)
        failed_product = await db.get(Product, failed_catalog.source_product_id if failed_catalog else 0)
        failed_data = await db.scalar(select(ProductData).where(ProductData.product_id == failed_product.id))
        assert failed_catalog is not None and failed_product is not None and failed_data is not None
        failed_data.amazon_template_path = str(invalid_source)
        failed_data.amazon_template_fill_summary = json.dumps({"logic_version": AMAZON_TEMPLATE_LOGIC_VERSION})
        rollback_pool_item = UpcPoolItem(upc="725999999904", status="available", source="r1_fixture")
        db.add(rollback_pool_item)
        await db.commit()
        rollback_pool_item_id = rollback_pool_item.id

    failed_transport = httpx.ASGITransport(app=app, client=("127.0.0.1", 41322))
    async with httpx.AsyncClient(transport=failed_transport, base_url="http://testserver") as client:
        with external_socket_guard():
            failed_response = await client.post("/api/products/catalog/export", json=[failed_catalog_id])
    assert failed_response.status_code == 400, failed_response.text

    async with environment.session_factory() as db:
        failed_catalog = await db.get(CatalogProduct, failed_catalog_id)
        failed_product = await db.get(Product, failed_catalog.source_product_id if failed_catalog else 0)
        rollback_pool_item = await db.get(UpcPoolItem, rollback_pool_item_id)
        assert failed_catalog is not None and failed_catalog.upc is None, failed_catalog
        assert failed_product is not None and failed_product.upc is None, failed_product
        assert rollback_pool_item is not None and rollback_pool_item.status == "available", rollback_pool_item
        assert rollback_pool_item.product_id is None, rollback_pool_item


async def test_real_worker_late_business_failure_rolls_back_row_mutations(environment) -> None:
    from openpyxl import Workbook
    from sqlalchemy import func, select

    from app.models import CatalogProduct, Product, ProductData, ProductFile, TaskGroup, TaskRun, TaskStep, UpcPoolItem
    from app.pipeline.step10_amazon_template import AMAZON_TEMPLATE_LOGIC_VERSION
    from app.task_runtime import catalog_export_workers
    from app.task_runtime.catalog_export_workers import register_catalog_export_workers
    from app.task_runtime.scheduler import drain_ready_steps

    catalog_id = (await _seed_catalog_export_products(environment, [
        {"item_code": "R1-LATE-BUSINESS-ROLLBACK", "upc": None},
    ]))[0]
    invalid_source = environment.data_dir / "exports" / "late-business-invalid-cache.xlsm"
    invalid_source.parent.mkdir(parents=True, exist_ok=True)
    Workbook().save(invalid_source)
    async with environment.session_factory() as db:
        catalog = await db.get(CatalogProduct, catalog_id)
        product = await db.get(Product, catalog.source_product_id if catalog else 0)
        product_data = await db.scalar(select(ProductData).where(ProductData.product_id == product.id))
        assert catalog is not None and product is not None and product_data is not None
        product_data.amazon_template_path = str(invalid_source)
        product_data.amazon_template_fill_summary = json.dumps({"logic_version": AMAZON_TEMPLATE_LOGIC_VERSION})
        pool_item = UpcPoolItem(upc="725999999905", status="available", source="r1_fixture")
        db.add(pool_item)
        await db.commit()
        product_id = product.id
        pool_item_id = pool_item.id
        baseline_listing_check = product_data.listing_check
        baseline_fill_summary = product_data.amazon_template_fill_summary

    run_id, step_id = await _prepare_catalog_run(environment, [catalog_id])
    upload_calls: list[Path] = []
    original_upload = catalog_export_workers.upload_private_file

    def fake_upload(path: Path, object_key: str) -> dict:
        upload_calls.append(Path(path))
        return {"object_key": object_key, "url": f"https://oss.invalid/{object_key}"}

    register_catalog_export_workers()
    catalog_export_workers.upload_private_file = fake_upload
    try:
        with external_socket_guard():
            await drain_ready_steps()
    finally:
        catalog_export_workers.upload_private_file = original_upload

    assert upload_calls == [], upload_calls
    async with environment.session_factory() as db:
        run = await db.get(TaskRun, run_id)
        step = await db.get(TaskStep, step_id)
        group = await db.scalar(select(TaskGroup).where(TaskGroup.task_run_id == run_id))
        catalog = await db.get(CatalogProduct, catalog_id)
        product = await db.get(Product, product_id)
        product_data = await db.scalar(select(ProductData).where(ProductData.product_id == product_id))
        pool_item = await db.get(UpcPoolItem, pool_item_id)
        product_file_count = await db.scalar(
            select(func.count(ProductFile.id)).where(ProductFile.product_id == product_id)
        )
        assert run is not None and run.status == "failed", run
        assert group is not None and group.status == "failed", group
        assert step is not None and step.status == "failed", step
        payload = json.loads(run.summary_json or "{}")
        assert payload.get("status") == "failed" and payload.get("failed_count") == 1, payload
        assert len(payload.get("rows") or []) == 1, payload
        assert "导入表格缺少 Template 工作表" in str(payload["rows"][0].get("reason") or ""), payload
        assert catalog is not None and catalog.upc is None, catalog
        assert product is not None and product.upc is None, product
        assert pool_item is not None and pool_item.status == "available" and pool_item.product_id is None, pool_item
        assert product_file_count == 0, product_file_count
        assert product_data is not None and product_data.amazon_template_path == str(invalid_source), product_data
        assert product_data.amazon_template_fill_summary == baseline_fill_summary, product_data
        assert product_data.listing_check == baseline_listing_check, product_data
        assert product_data.amazon_template_generated_at is None, product_data
        assert product_data.amazon_template_warnings is None, product_data


async def test_real_worker_upload_system_error_rolls_back_builder_state(environment) -> None:
    from sqlalchemy import func, select

    from app.database import engine
    from app.models import CatalogProduct, Product, TaskRun, TaskStep, TaskStepEvent, UpcPoolItem
    from app.task_runtime import catalog_export_workers
    from app.task_runtime.catalog_export_workers import register_catalog_export_workers
    from app.task_runtime.scheduler import drain_ready_steps

    catalog_id = (await _seed_catalog_export_products(environment, [
        {"item_code": "R1-UPLOAD-SYSTEM-ERROR", "upc": None},
    ]))[0]
    async with environment.session_factory() as db:
        pool_item = UpcPoolItem(upc="725999999901", status="available", source="r1_fixture")
        db.add(pool_item)
        await db.commit()
        pool_item_id = pool_item.id
    run_id, step_id = await _prepare_catalog_run(environment, [catalog_id])

    upload_calls: list[Path] = []
    original_upload = catalog_export_workers.upload_private_file

    def failing_upload(path: Path, _object_key: str) -> dict:
        upload_calls.append(Path(path))
        raise OSError("R1 injected upload outage")

    register_catalog_export_workers()
    catalog_export_workers.upload_private_file = failing_upload
    try:
        with external_socket_guard():
            await drain_ready_steps()
    finally:
        catalog_export_workers.upload_private_file = original_upload

    assert len(upload_calls) == 1, upload_calls
    assert upload_calls[0].is_file(), upload_calls
    async with environment.session_factory() as db:
        run = await db.get(TaskRun, run_id)
        step = await db.get(TaskStep, step_id)
        catalog = await db.get(CatalogProduct, catalog_id)
        product = await db.get(Product, catalog.source_product_id if catalog else 0)
        pool_item = await db.get(UpcPoolItem, pool_item_id)
        progress_events = await db.scalar(
            select(func.count(TaskStepEvent.id)).where(
                TaskStepEvent.task_run_id == run_id,
                TaskStepEvent.event_type == "progress",
            )
        )
        assert run is not None and run.status == "failed" and run.summary_json is None, run
        assert step is not None and step.status == "failed" and step.result_json is None, step
        assert step.error_message and step.error_message.startswith("OSError: R1 injected upload outage"), step
        assert catalog is not None and catalog.upc is None and catalog.exported_at is None, catalog
        assert product is not None and product.upc is None, product
        assert pool_item is not None and pool_item.status == "available" and pool_item.product_id is None, pool_item
        assert progress_events == 0, progress_events


async def test_real_builder_db_error_is_not_business_partial(environment) -> None:
    from sqlalchemy import event, func, select
    from sqlalchemy.exc import OperationalError

    from app.database import engine
    from app.models import TaskRun, TaskStep, TaskStepEvent, UpcPoolItem
    from app.task_runtime import catalog_export_workers
    from app.task_runtime.catalog_export_workers import register_catalog_export_workers
    from app.task_runtime.scheduler import drain_ready_steps

    catalog_ids = await _seed_catalog_export_products(environment, [
        {"item_code": "R1-DB-ERROR-FIRST"},
        {"item_code": "R1-DB-ERROR-SECOND", "upc": None},
    ])
    async with environment.session_factory() as db:
        db.add(UpcPoolItem(upc="725999999902", status="available", source="r1_fixture"))
        await db.commit()
    run_id, step_id = await _prepare_catalog_run(environment, catalog_ids)

    upload_calls: list[Path] = []
    original_upload = catalog_export_workers.upload_private_file

    def fake_upload(path: Path, object_key: str) -> dict:
        upload_calls.append(Path(path))
        return {"object_key": object_key, "url": f"https://oss.invalid/{object_key}"}

    def fail_upc_query(_connection, _cursor, statement, parameters, _context, _executemany):
        if "upc_pool_items" in statement.lower():
            raise OperationalError(statement, parameters, RuntimeError("R1 injected database outage"))

    register_catalog_export_workers()
    catalog_export_workers.upload_private_file = fake_upload
    event.listen(engine.sync_engine, "before_cursor_execute", fail_upc_query)
    try:
        with external_socket_guard():
            await drain_ready_steps()
    finally:
        event.remove(engine.sync_engine, "before_cursor_execute", fail_upc_query)
        catalog_export_workers.upload_private_file = original_upload

    assert upload_calls == [], upload_calls
    async with environment.session_factory() as db:
        run = await db.get(TaskRun, run_id)
        step = await db.get(TaskStep, step_id)
        progress_events = await db.scalar(
            select(func.count(TaskStepEvent.id)).where(
                TaskStepEvent.task_run_id == run_id,
                TaskStepEvent.event_type == "progress",
            )
        )
        assert run is not None and run.status == "failed" and run.summary_json is None, run
        assert step is not None and step.status == "failed" and step.result_json is None, step
        assert step.error_message and step.error_message.startswith("OperationalError:"), step
        assert "R1 injected database outage" in step.error_message, step
        assert progress_events == 0, progress_events


async def test_real_mixed_worker_scheduler_outcome(environment) -> None:
    from openpyxl import Workbook
    from sqlalchemy import select
    from sqlalchemy.orm import selectinload

    from app.api.schemas import OfflineTaskCatalogExportRequest
    from app.models import CatalogProduct, Product, ProductData, TaskGroup, TaskRun, TaskStep, TaskStepEvent, UpcPoolItem
    from app.pipeline.step10_amazon_template import AMAZON_TEMPLATE_LOGIC_VERSION
    from app.task_planners.catalog_export import create_catalog_export_runs
    from app.task_runtime import catalog_export_workers
    from app.task_runtime.catalog_export_workers import register_catalog_export_workers
    from app.task_runtime.scheduler import drain_ready_steps

    catalog_ids = await _seed_catalog_export_products(environment, [
        {"item_code": "R1-EXPORT-DONE"},
        {"item_code": "R1-EXPORT-LATE-FAIL", "upc": None},
    ])
    invalid_source = environment.data_dir / "exports" / "mixed-late-business-invalid-cache.xlsm"
    invalid_source.parent.mkdir(parents=True, exist_ok=True)
    Workbook().save(invalid_source)
    async with environment.session_factory() as db:
        failed_catalog = await db.get(CatalogProduct, catalog_ids[1])
        failed_product = await db.get(Product, failed_catalog.source_product_id if failed_catalog else 0)
        failed_data = await db.scalar(select(ProductData).where(ProductData.product_id == failed_product.id))
        assert failed_catalog is not None and failed_product is not None and failed_data is not None
        failed_data.amazon_template_path = str(invalid_source)
        failed_data.amazon_template_fill_summary = json.dumps({"logic_version": AMAZON_TEMPLATE_LOGIC_VERSION})
        pool_item = UpcPoolItem(upc="725999999906", status="available", source="r1_fixture")
        db.add(pool_item)
        await db.commit()
        failed_product_id = failed_product.id
        pool_item_id = pool_item.id

        runs, errors = await create_catalog_export_runs(
            db,
            OfflineTaskCatalogExportRequest(catalog_product_ids=catalog_ids),
            created_by="r1_catalog_test",
            auto_start=False,
        )
        assert not errors, errors
        assert len(runs) == 1, runs
        run_id = runs[0].id
        step = (
            await db.execute(select(TaskStep).where(TaskStep.task_run_id == run_id))
        ).scalar_one()
        step.status = "ready"
        await db.commit()

    upload_calls: list[tuple[Path, str]] = []
    original_upload = catalog_export_workers.upload_private_file

    def fake_upload(path: Path, object_key: str) -> dict:
        upload_calls.append((Path(path), object_key))
        return {"object_key": object_key, "url": f"https://oss.invalid/{object_key}"}

    register_catalog_export_workers()
    catalog_export_workers.upload_private_file = fake_upload
    try:
        with external_socket_guard():
            await drain_ready_steps()
    finally:
        catalog_export_workers.upload_private_file = original_upload

    assert len(upload_calls) == 1, upload_calls
    assert upload_calls[0][0].is_file(), upload_calls
    async with environment.session_factory() as db:
        run = (
            await db.execute(
                select(TaskRun)
                .where(TaskRun.id == run_id)
                .options(selectinload(TaskRun.groups), selectinload(TaskRun.steps))
            )
        ).scalar_one()
        assert len(run.groups) == 1
        assert len(run.steps) == 1
        assert run.status == "partial_failed", run.status
        assert run.groups[0].status == "partial_failed", run.groups[0].status
        assert run.steps[0].status == "partial_failed", run.steps[0].status
        summary = json.loads(run.summary_json or "{}")
        result = json.loads(run.steps[0].result_json or "{}")
        assert summary == result
        assert summary["status"] == "partial_failed", summary
        assert summary["success_count"] == 1, summary
        assert summary["skipped_count"] == 0, summary
        assert summary["failed_count"] == 1, summary
        assert summary["artifact_available"] is True, summary
        failed_row = next(row for row in summary["rows"] if row["catalog_id"] == catalog_ids[1])
        assert failed_row["status"] == "failed", failed_row
        assert "导入表格缺少 Template 工作表" in str(failed_row.get("reason") or ""), failed_row
        events = (
            await db.execute(
                select(TaskStepEvent)
                .where(TaskStepEvent.task_run_id == run_id)
                .order_by(TaskStepEvent.id.asc())
            )
        ).scalars().all()
        assert [event.event_type for event in events].count("warning") == 1, events
        assert not any(event.message == "step 执行成功" for event in events), events
        exported = await db.get(CatalogProduct, catalog_ids[0])
        failed = await db.get(CatalogProduct, catalog_ids[1])
        failed_product = await db.get(Product, failed_product_id)
        pool_item = await db.get(UpcPoolItem, pool_item_id)
        assert exported is not None and exported.export_task_id == run_id, exported
        assert failed is not None and failed.export_task_id is None and failed.upc is None, failed
        assert failed_product is not None and failed_product.upc is None, failed_product
        assert pool_item is not None and pool_item.status == "available" and pool_item.product_id is None, pool_item


async def test_real_done_and_business_failed_outcomes(environment) -> None:
    from sqlalchemy import select
    from sqlalchemy.orm import selectinload

    from app.api.schemas import OfflineTaskCatalogExportRequest
    from app.models import TaskRun, TaskStep, TaskStepEvent
    from app.task_planners.catalog_export import create_catalog_export_runs
    from app.task_runtime import catalog_export_workers
    from app.task_runtime.catalog_export_workers import register_catalog_export_workers
    from app.task_runtime.scheduler import drain_ready_steps

    async def execute(specs: list[dict]) -> tuple[int, list[tuple[Path, str]]]:
        catalog_ids = await _seed_catalog_export_products(environment, specs)
        async with environment.session_factory() as db:
            runs, errors = await create_catalog_export_runs(
                db,
                OfflineTaskCatalogExportRequest(catalog_product_ids=catalog_ids),
                created_by="r1_catalog_test",
                auto_start=False,
            )
            assert not errors, errors
            assert len(runs) == 1, runs
            run_id = runs[0].id
            step = (
                await db.execute(select(TaskStep).where(TaskStep.task_run_id == run_id))
            ).scalar_one()
            step.status = "ready"
            await db.commit()

        upload_calls: list[tuple[Path, str]] = []
        original_upload = catalog_export_workers.upload_private_file

        def fake_upload(path: Path, object_key: str) -> dict:
            upload_calls.append((Path(path), object_key))
            return {"object_key": object_key, "url": f"https://oss.invalid/{object_key}"}

        register_catalog_export_workers()
        catalog_export_workers.upload_private_file = fake_upload
        try:
            with external_socket_guard():
                await drain_ready_steps()
        finally:
            catalog_export_workers.upload_private_file = original_upload
        return run_id, upload_calls

    done_run_id, done_uploads = await execute([{"item_code": "R1-EXPORT-ALL-DONE"}])
    failed_run_id, failed_uploads = await execute([
        {"item_code": "R1-EXPORT-ALL-FAILED", "amazon_asin": "B0R1FAIL01"},
    ])
    assert len(done_uploads) == 1, done_uploads
    assert failed_uploads == [], failed_uploads

    async with environment.session_factory() as db:
        for run_id, expected_status, event_type in [
            (done_run_id, "succeeded", "status"),
            (failed_run_id, "failed", "error"),
        ]:
            run = (
                await db.execute(
                    select(TaskRun)
                    .where(TaskRun.id == run_id)
                    .options(selectinload(TaskRun.groups), selectinload(TaskRun.steps))
                )
            ).scalar_one()
            assert run.status == expected_status, run.status
            assert run.groups[0].status == expected_status, run.groups[0].status
            assert run.steps[0].status == expected_status, run.steps[0].status
            payload = json.loads(run.summary_json or "{}")
            if expected_status == "succeeded":
                assert payload["status"] == "done", payload
                assert payload["success_count"] == 1, payload
                assert payload["artifact_available"] is True, payload
            else:
                assert payload["status"] == "failed", payload
                assert payload["success_count"] == 0, payload
                assert payload["skipped_count"] == 1, payload
                assert payload["artifact_available"] is False, payload
                assert payload["rows"][0]["reason"], payload
            events = (
                await db.execute(
                    select(TaskStepEvent)
                    .where(TaskStepEvent.task_run_id == run_id)
                    .order_by(TaskStepEvent.id.asc())
                )
            ).scalars().all()
            assert [event.event_type for event in events].count(event_type) >= 1, events
            assert not any(event.message == "step 执行成功" for event in events), events


async def test_effective_legacy_partial_list_api(environment) -> None:
    import httpx

    from app.main import app
    from app.models import TaskRun
    from app.services.offline_tasks import _catalog_export_result_payload

    artifact = environment.data_dir / "exports" / "legacy-partial.zip"
    artifact.parent.mkdir(parents=True, exist_ok=True)
    artifact.write_bytes(b"legacy-partial")
    legacy_category = "R1 Legacy Effective Category"
    native_category = "R1 Native Effective Category"
    raw_failed_category = "R1 Raw Failed Partial Artifact Category"
    legacy_partial_payload = {
        "status": "partial_failed",
        "category": legacy_category,
        "categories": [legacy_category],
        "template_name": "CHAIR_SOFA.xlsm",
        "template_path": "/fixture/CHAIR_SOFA.xlsm",
        "catalog_product_ids": [501, 502],
        "requested_count": 2,
        "success_count": 1,
        "skipped_count": 1,
        "failed_count": 0,
        "filename": artifact.name,
        "file_path": str(artifact),
        "file_size": artifact.stat().st_size,
        "created_at": datetime.now().isoformat(),
        "rows": [
            {
                "catalog_id": 501,
                "product_id": 1501,
                "item_code": "LEGACY-DONE",
                "category": legacy_category,
                "status": "exported",
                "reason": None,
            },
            {
                "catalog_id": 502,
                "product_id": 1502,
                "item_code": "LEGACY-SKIP",
                "category": legacy_category,
                "status": "skipped",
                "reason": "保护门",
            },
        ],
    }
    native_partial_payload = _catalog_export_result_payload(
        category=native_category,
        categories=[native_category],
        template_name="CHAIR_SOFA.xlsm",
        template_path="/fixture/CHAIR_SOFA.xlsm",
        catalog_ids=[511, 512],
        report_rows=[
            {"商品资料ID": 511, "商品ID": 1511, "商品Code": "NATIVE-DONE", "状态": "已导出"},
            {"商品资料ID": 512, "商品ID": 1512, "商品Code": "NATIVE-SKIP", "状态": "跳过", "原因": "保护门"},
        ],
        created_at=datetime.now(),
        filename=artifact.name,
        file_path=str(artifact),
        file_size=artifact.stat().st_size,
    )
    raw_failed_partial_payload = json.loads(json.dumps(legacy_partial_payload))
    raw_failed_partial_payload["category"] = raw_failed_category
    raw_failed_partial_payload["categories"] = [raw_failed_category]
    raw_failed_partial_payload["catalog_product_ids"] = [503, 504]
    raw_failed_partial_payload["rows"][0]["catalog_id"] = 503
    raw_failed_partial_payload["rows"][1]["catalog_id"] = 504
    cases = [
        ("legacy_partial", "catalog_export", "succeeded", json.dumps(legacy_partial_payload)),
        ("native_partial", "catalog_export", "partial_failed", json.dumps(native_partial_payload)),
        ("done", "catalog_export", "succeeded", json.dumps({"status": "done"})),
        ("null_summary", "catalog_export", "succeeded", None),
        ("empty_summary", "catalog_export", "succeeded", ""),
        ("invalid_summary", "catalog_export", "succeeded", "{invalid-json"),
        ("empty_object", "catalog_export", "succeeded", "{}"),
        ("missing_status", "catalog_export", "succeeded", json.dumps({"other": "value"})),
        ("null_status", "catalog_export", "succeeded", json.dumps({"status": None})),
        ("non_catalog_partial", "giga_pull", "succeeded", json.dumps({"status": "partial_failed"})),
        ("raw_failed_partial", "catalog_export", "failed", json.dumps(raw_failed_partial_payload)),
    ]
    run_ids: dict[str, int] = {}
    async with environment.session_factory() as db:
        for name, task_type, status, summary_json in cases:
            run = TaskRun(
                task_type=task_type,
                title=f"R1 effective {name}",
                status=status,
                summary_json=summary_json,
            )
            db.add(run)
            await db.flush()
            run_ids[name] = run.id
        await db.commit()

    transport = httpx.ASGITransport(app=app, client=("127.0.0.1", 41322))
    async with httpx.AsyncClient(transport=transport, base_url="http://testserver") as client:
        partial = await client.get(
            "/api/task-runs",
            params={"view": "all", "display_status": "partial_failed", "q": "R1 effective", "page_size": 20},
        )
        assert partial.status_code == 200, partial.text
        partial_body = partial.json()
        assert partial_body["total"] == 3, partial_body
        assert partial_body["base_total"] == 11, partial_body
        assert partial_body["filtered_total"] == 3, partial_body
        assert {item["id"] for item in partial_body["items"]} == {
            run_ids["legacy_partial"],
            run_ids["native_partial"],
            run_ids["raw_failed_partial"],
        }, partial_body
        legacy = next(item for item in partial_body["items"] if item["id"] == run_ids["legacy_partial"])
        assert legacy["status"] == "partial_failed", legacy
        assert legacy["display_status"] == "partial_failed", legacy
        assert legacy["catalog_export_result"]["requested_count"] == 2, legacy
        assert legacy["catalog_export_result"]["artifact_available"] is None, legacy
        assert "download_result" in legacy["available_actions"], legacy
        assert "retry_failed_steps" not in legacy["available_actions"], legacy

        legacy_detail = await client.get(f"/api/task-runs/{run_ids['legacy_partial']}")
        assert legacy_detail.status_code == 200, legacy_detail.text
        legacy_detail_body = legacy_detail.json()
        assert legacy_detail_body["status"] == "partial_failed", legacy_detail_body
        assert legacy_detail_body["display_status"] == "partial_failed", legacy_detail_body
        assert legacy_detail_body["catalog_export_result"]["rows"][1]["reason"] == "保护门", legacy_detail_body
        assert "download_result" in legacy_detail_body["available_actions"], legacy_detail_body
        assert "retry_failed_steps" not in legacy_detail_body["available_actions"], legacy_detail_body

        raw_failed_detail = await client.get(f"/api/task-runs/{run_ids['raw_failed_partial']}")
        assert raw_failed_detail.status_code == 200, raw_failed_detail.text
        raw_failed_detail_body = raw_failed_detail.json()
        assert raw_failed_detail_body["status"] == "partial_failed", raw_failed_detail_body
        assert raw_failed_detail_body["display_status"] == "partial_failed", raw_failed_detail_body
        assert raw_failed_detail_body["catalog_export_result"]["artifact_available"] is None, raw_failed_detail_body
        assert "download_result" in raw_failed_detail_body["available_actions"], raw_failed_detail_body

        history_pages: list[dict] = []
        for page_number in range(1, 5):
            response = await client.get(
                "/api/task-runs",
                params={
                    "view": "history",
                    "display_status": "succeeded",
                    "q": "R1 effective",
                    "page": page_number,
                    "page_size": 2,
                },
            )
            assert response.status_code == 200, response.text
            history_pages.append(response.json())
        for body in history_pages:
            assert body["base_total"] == 8, body
            assert body["filtered_total"] == 8, body
            assert body["total"] == 8, body
        history_items = [item for body in history_pages for item in body["items"]]
        history_ids = {item["id"] for item in history_items}
        assert len(history_items) == len(history_ids) == 8, history_pages
        assert history_ids == {
            run_ids["done"],
            run_ids["null_summary"],
            run_ids["empty_summary"],
            run_ids["invalid_summary"],
            run_ids["empty_object"],
            run_ids["missing_status"],
            run_ids["null_status"],
            run_ids["non_catalog_partial"],
        }, history_pages

        succeeded_all = await client.get(
            "/api/task-runs",
            params={"view": "all", "display_status": "succeeded", "q": "R1 effective", "page_size": 20},
        )
        assert succeeded_all.status_code == 200, succeeded_all.text
        succeeded_body = succeeded_all.json()
        assert succeeded_body["base_total"] == 11, succeeded_body
        assert succeeded_body["filtered_total"] == 8, succeeded_body
        assert succeeded_body["total"] == 8, succeeded_body
        assert {item["id"] for item in succeeded_body["items"]} == history_ids, succeeded_body

        current = await client.get(
            "/api/task-runs",
            params={"view": "current", "q": "R1 effective", "page_size": 20},
        )
        assert current.status_code == 200, current.text
        current_body = current.json()
        assert current_body["base_total"] == 3, current_body
        assert current_body["filtered_total"] == 3, current_body
        assert current_body["total"] == 3, current_body
        current_ids = {item["id"] for item in current_body["items"]}
        assert run_ids["legacy_partial"] in current_ids, current_body
        assert run_ids["native_partial"] in current_ids, current_body
        assert run_ids["raw_failed_partial"] in current_ids, current_body
        raw_failed_list_item = next(
            item for item in current_body["items"] if item["id"] == run_ids["raw_failed_partial"]
        )
        assert raw_failed_list_item["status"] == "partial_failed", raw_failed_list_item
        assert raw_failed_list_item["display_status"] == "partial_failed", raw_failed_list_item
        assert "download_result" in raw_failed_list_item["available_actions"], raw_failed_list_item
        assert not current_ids.intersection({
            run_ids["done"],
            run_ids["null_summary"],
            run_ids["empty_summary"],
            run_ids["invalid_summary"],
            run_ids["empty_object"],
            run_ids["missing_status"],
            run_ids["null_status"],
            run_ids["non_catalog_partial"],
        }), current_body

        categories = await client.get("/api/products/catalog/export-categories")
        assert categories.status_code == 200, categories.text
        exported_categories = {item["category"]: item for item in categories.json()["exported"]}
        assert exported_categories[legacy_category]["count"] == 1, exported_categories

        raw_failed_files = await client.get(
            "/api/products/catalog/export-files",
            params={"category": raw_failed_category, "page_size": 20},
        )
        assert raw_failed_files.status_code == 200, raw_failed_files.text
        raw_failed_files_body = raw_failed_files.json()
        assert raw_failed_files_body["total"] == 1, raw_failed_files_body
        raw_failed_file = raw_failed_files_body["items"][0]
        assert raw_failed_file["task_id"] == run_ids["raw_failed_partial"], raw_failed_file
        assert raw_failed_file["task_status"] == "partial_failed", raw_failed_file
        assert raw_failed_file["can_download"] is True, raw_failed_file

        legacy_files = await client.get(
            "/api/products/catalog/export-files",
            params={"category": legacy_category, "page_size": 20},
        )
        assert legacy_files.status_code == 200, legacy_files.text
        legacy_files_body = legacy_files.json()
        assert legacy_files_body["total"] == 1, legacy_files_body
        legacy_file = legacy_files_body["items"][0]
        assert legacy_file["task_id"] == run_ids["legacy_partial"], legacy_file
        assert legacy_file["task_status"] == "partial_failed", legacy_file
        assert legacy_file["can_download"] is True, legacy_file

        raw_failed_download = await client.get(f"/api/task-runs/{run_ids['raw_failed_partial']}/download")
        assert raw_failed_download.status_code == 200, raw_failed_download.text
        assert raw_failed_download.content == artifact.read_bytes(), raw_failed_download.content
        legacy_download = await client.get(f"/api/task-runs/{run_ids['legacy_partial']}/download")
        assert legacy_download.status_code == 200, legacy_download.text
        assert legacy_download.content == artifact.read_bytes(), legacy_download.content
        assert exported_categories[raw_failed_category]["count"] == 1, exported_categories


async def test_summary_total_parser_and_legacy_projection_contract(environment) -> None:
    import httpx
    from sqlalchemy import event

    from app.database import engine
    from app.main import app
    from app.models import TaskGroup, TaskRun, TaskStep
    from app.task_runtime import catalog_export_status

    artifact = environment.data_dir / "exports" / "summary-total-parser.zip"
    artifact.parent.mkdir(parents=True, exist_ok=True)
    artifact.write_bytes(b"summary-total-parser")
    legacy_category = "R1 Summary Total Legacy Partial"
    escaped_category = "R1 Summary Total Escaped Partial"

    def downloadable_payload(category: str) -> dict:
        return {
            "status": "partial_failed",
            "category": category,
            "categories": [category],
            "filename": artifact.name,
            "file_path": str(artifact),
            "rows": [
                {"catalog_id": 96001, "category": category, "status": "exported"},
                {"catalog_id": 96002, "category": category, "status": "skipped", "reason": "legacy partial"},
            ],
            "requested_count": 2,
            "success_count": 1,
            "skipped_count": 1,
            "failed_count": 0,
            "report_count": 2,
        }

    mysql_depth_summary = (
        '{"status":"done","file_path":null,"nested":'
        + "[" * 150
        + "0"
        + "]" * 150
        + "}"
    )
    python_depth_summary = (
        '{"status":"done","file_path":null,"nested":'
        + "[" * 10000
        + "0"
        + "]" * 10000
        + "}"
    )
    huge_integer = "9" * 310
    legacy_payload = downloadable_payload(legacy_category)
    escaped_payload = downloadable_payload(escaped_category)
    escaped_summary = json.dumps(escaped_payload).replace(
        '"partial_failed"',
        '"partial\\u005ffailed"',
        1,
    )
    cases: list[tuple[str, str, str | None]] = [
        ("legacy_partial", "succeeded", json.dumps(legacy_payload)),
        ("escaped_legacy_partial", "succeeded", escaped_summary),
        ("native_partial", "partial_failed", json.dumps(downloadable_payload("R1 Summary Total Native Partial"))),
        ("done", "succeeded", json.dumps({"status": "done"})),
        ("mysql_depth", "succeeded", mysql_depth_summary),
        ("python_depth", "succeeded", python_depth_summary),
        ("invalid", "succeeded", "{not-json"),
        ("nan", "succeeded", '{"status":"done","file_path":NaN}'),
        ("infinity", "succeeded", '{"status":"done","file_path":Infinity}'),
        ("overflow_float", "succeeded", '{"status":"done","file_path":1e309}'),
        ("long_integer", "succeeded", f'{{"status":"done","file_path":{huge_integer}}}'),
        ("surrogate", "succeeded", '{"status":"done","file_path":"\\ud800"}'),
    ]
    anomaly_names = {
        "mysql_depth",
        "python_depth",
        "invalid",
        "nan",
        "infinity",
        "overflow_float",
        "long_integer",
        "surrogate",
    }
    invalid_parse_names = {
        "python_depth",
        "invalid",
        "nan",
        "infinity",
        "overflow_float",
    }
    valid_unsafe_names = anomaly_names - invalid_parse_names
    run_ids: dict[str, int] = {}
    old_categories: dict[str, str] = {}
    async with environment.session_factory() as db:
        for name, status, summary_json in cases:
            run = TaskRun(
                task_type="catalog_export",
                title=f"R1 summary total {name}",
                status=status,
                summary_json=summary_json,
            )
            db.add(run)
            await db.flush()
            run_ids[name] = run.id
            if name not in anomaly_names:
                continue
            old_category = f"R1 Summary Total Old Step {name}"
            old_categories[name] = old_category
            group = TaskGroup(
                task_run_id=run.id,
                group_key="export_file",
                title=f"R1 summary total {name}",
                status="succeeded",
            )
            db.add(group)
            await db.flush()
            old_payload = downloadable_payload(old_category)
            old_payload["status"] = "done"
            db.add(TaskStep(
                task_run_id=run.id,
                task_group_id=group.id,
                step_key="catalog_export_template_older",
                step_type="catalog_export_template",
                status="succeeded",
                result_json=json.dumps(old_payload),
            ))
        await db.commit()

    original_json_loads = catalog_export_status.json.loads

    def fail_with_memory_error(*_args, **_kwargs):
        raise MemoryError("R1 injected summary parser memory pressure")

    catalog_export_status.json.loads = fail_with_memory_error
    try:
        assert catalog_export_status.effective_task_run_status(
            task_type="catalog_export",
            status="succeeded",
            summary_json='{"status":"partial_failed"}',
        ) == "succeeded"
        assert catalog_export_status.select_catalog_export_payload('{"status":"done"}', []) == {}
    finally:
        catalog_export_status.json.loads = original_json_loads

    def legacy_projection_queries(records: list[tuple[str, object]]) -> list[tuple[str, object]]:
        return [
            item
            for item in records
            if (
                "select task_runs.id, task_runs.status, task_runs.summary_json" in item[0]
                and "task_runs.task_type =" in item[0]
            ) or (
                "select task_runs.id, task_runs.summary_json" in item[0]
                and "task_runs.summary_json is not null" in item[0]
            )
        ]

    def assert_summary_query_safety(records: list[tuple[str, object]], *, projection_count: int = 1) -> None:
        assert len(legacy_projection_queries(records)) == projection_count, records
        for statement, _parameters in records:
            assert "json_valid" not in statement, statement
            assert "json_extract" not in statement, statement
            assert "json_unquote" not in statement, statement
            assert "json_contains" not in statement, statement

    transport = httpx.ASGITransport(app=app, client=("127.0.0.1", 41330))
    async with httpx.AsyncClient(
        transport=transport,
        base_url="http://testserver",
        follow_redirects=False,
    ) as client:
        with external_socket_guard():
            async def get_task_runs(params: dict) -> dict:
                records: list[tuple[str, object]] = []

                def record(_connection, _cursor, statement, parameters, _context, _executemany):
                    records.append((" ".join(statement.lower().split()), parameters))

                event.listen(engine.sync_engine, "before_cursor_execute", record)
                try:
                    response = await client.get("/api/task-runs", params=params)
                finally:
                    event.remove(engine.sync_engine, "before_cursor_execute", record)
                assert response.status_code == 200, response.text
                assert_summary_query_safety(records)
                assert any("count(task_runs.id)" in statement for statement, _params in records), records
                assert any(" limit " in statement for statement, _params in records), records
                return response.json()

            partial_pages = [
                await get_task_runs({
                    "view": "current",
                    "display_status": "partial_failed",
                    "q": "R1 summary total",
                    "page": page_number,
                    "page_size": 2,
                })
                for page_number in range(1, 6)
            ]
            for body in partial_pages:
                assert body["base_total"] == 8, body
                assert body["filtered_total"] == 8, body
                assert body["total"] == 8, body
            partial_items = [item for body in partial_pages for item in body["items"]]
            assert {item["id"] for item in partial_items} == {
                run_ids["legacy_partial"],
                run_ids["escaped_legacy_partial"],
                run_ids["native_partial"],
                *[run_ids[name] for name in invalid_parse_names],
            }, partial_items

            history_pages = []
            for page_number in range(1, 4):
                history_pages.append(await get_task_runs({
                    "view": "history",
                    "display_status": "succeeded",
                    "q": "R1 summary total",
                    "page": page_number,
                    "page_size": 3,
                }))
            for body in history_pages:
                assert body["base_total"] == 4, body
                assert body["filtered_total"] == 4, body
                assert body["total"] == 4, body
            history_items = [item for body in history_pages for item in body["items"]]
            history_ids = {item["id"] for item in history_items}
            assert len(history_items) == len(history_ids) == 4, history_pages
            assert history_ids == {
                run_ids["done"],
                *[run_ids[name] for name in valid_unsafe_names],
            }, history_pages

            all_succeeded = await get_task_runs({
                "view": "all",
                "display_status": "succeeded",
                "q": "R1 summary total",
                "page_size": 20,
            })
            assert all_succeeded["base_total"] == 12, all_succeeded
            assert all_succeeded["filtered_total"] == 4, all_succeeded
            assert {item["id"] for item in all_succeeded["items"]} == history_ids, all_succeeded

            all_partial = await get_task_runs({
                "view": "all",
                "display_status": "partial_failed",
                "q": "R1 summary total",
                "page_size": 20,
            })
            assert all_partial["base_total"] == 12, all_partial
            assert all_partial["filtered_total"] == 8, all_partial

            for name in invalid_parse_names:
                detail = await client.get(f"/api/task-runs/{run_ids[name]}")
                assert detail.status_code == 200, (name, detail.text)
                detail_body = detail.json()
                assert detail_body["status"] == "partial_failed", (name, detail_body)
                assert "download_result" in detail_body["available_actions"], (name, detail_body)
                assert old_categories[name] in detail_body["catalog_export_result"]["categories"], (name, detail_body)
                download = await client.get(f"/api/task-runs/{run_ids[name]}/download")
                assert download.status_code == 200, (name, download.status_code, download.text)
                assert download.content == artifact.read_bytes(), (name, download.content)

            for name in valid_unsafe_names:
                detail = await client.get(f"/api/task-runs/{run_ids[name]}")
                assert detail.status_code == 200, (name, detail.text)
                detail_body = detail.json()
                assert detail_body["status"] == "succeeded", (name, detail_body)
                assert "download_result" not in detail_body["available_actions"], (name, detail_body)
                assert old_categories[name] not in detail_body["catalog_export_result"]["categories"], (name, detail_body)
                download = await client.get(f"/api/task-runs/{run_ids[name]}/download")
                assert download.status_code == 400, (name, download.text)

            export_records: list[tuple[str, object]] = []

            def record_exports(_connection, _cursor, statement, parameters, _context, _executemany):
                export_records.append((" ".join(statement.lower().split()), parameters))

            event.listen(engine.sync_engine, "before_cursor_execute", record_exports)
            try:
                export_files = await client.get("/api/products/catalog/export-files", params={"page_size": 100})
            finally:
                event.remove(engine.sync_engine, "before_cursor_execute", record_exports)
            assert export_files.status_code == 200, export_files.text
            assert_summary_query_safety(export_records)
            export_items = {
                (item["task_source"], item["task_id"]): item
                for item in export_files.json()["items"]
            }
            for name in invalid_parse_names:
                item = export_items[("task_run", run_ids[name])]
                assert item["task_status"] == "partial_failed", (name, item)
                assert item["can_download"] is True, (name, item)
            for name in valid_unsafe_names:
                item = export_items[("task_run", run_ids[name])]
                assert item["task_status"] == "succeeded", (name, item)
                assert item["can_download"] is False, (name, item)

            category_records: list[tuple[str, object]] = []

            def record_categories(_connection, _cursor, statement, parameters, _context, _executemany):
                category_records.append((" ".join(statement.lower().split()), parameters))

            event.listen(engine.sync_engine, "before_cursor_execute", record_categories)
            try:
                categories = await client.get("/api/products/catalog/export-categories")
            finally:
                event.remove(engine.sync_engine, "before_cursor_execute", record_categories)
            assert categories.status_code == 200, categories.text
            assert_summary_query_safety(category_records)
            exported_categories = {item["category"] for item in categories.json()["exported"]}
            assert {
                old_categories[name]
                for name in invalid_parse_names
            }.issubset(exported_categories), exported_categories
            assert {
                old_categories[name]
                for name in valid_unsafe_names
            }.isdisjoint(exported_categories), exported_categories
            assert legacy_category in exported_categories, exported_categories
            assert escaped_category in exported_categories, exported_categories


async def test_pre_r1_task_run_remote_download_and_action_contract(environment) -> None:
    import httpx

    from app.main import app
    from app.models import TaskGroup, TaskRun, TaskStep

    local_artifact = environment.data_dir / "exports" / "fix5-local.zip"
    local_artifact.parent.mkdir(parents=True, exist_ok=True)
    local_artifact.write_bytes(b"fix5-local")
    stale_path = local_artifact.with_name("fix5-stale.zip")
    url_file_path = "https://downloads.invalid/legacy-url-file-path.zip"
    oss_url = "https://downloads.invalid/legacy-oss-url.zip"
    fallback_url = "https://downloads.invalid/legacy-step-fallback.zip"

    payloads = {
        "url_file_path": {
            "status": "done",
            "category": "R1 Fix5 URL File Path",
            "categories": ["R1 Fix5 URL File Path"],
            "filename": "legacy-url-file-path.zip",
            "file_path": url_file_path,
        },
        "oss_url": {
            "status": "partial_failed",
            "category": "R1 Fix5 OSS URL",
            "categories": ["R1 Fix5 OSS URL"],
            "filename": "legacy-oss-url.zip",
            "oss_url": oss_url,
        },
        "explicit_false": {
            "status": "done",
            "category": "R1 Fix5 Explicit False",
            "categories": ["R1 Fix5 Explicit False"],
            "artifact_available": False,
            "filename": local_artifact.name,
            "file_path": str(local_artifact),
            "oss_url": "https://downloads.invalid/must-not-redirect-false.zip",
        },
        "explicit_null": {
            "status": "partial_failed",
            "category": "R1 Fix5 Explicit Null",
            "categories": ["R1 Fix5 Explicit Null"],
            "artifact_available": None,
            "filename": local_artifact.name,
            "file_path": str(local_artifact),
            "oss_url": "https://downloads.invalid/must-not-redirect-null.zip",
        },
        "stale_local": {
            "status": "done",
            "category": "R1 Fix5 Stale Local",
            "categories": ["R1 Fix5 Stale Local"],
            "filename": stale_path.name,
            "file_path": str(stale_path),
        },
        "local": {
            "status": "done",
            "category": "R1 Fix5 Local",
            "categories": ["R1 Fix5 Local"],
            "filename": local_artifact.name,
            "file_path": str(local_artifact),
        },
    }
    run_ids: dict[str, int] = {}
    async with environment.session_factory() as db:
        for name, payload in payloads.items():
            run = TaskRun(
                task_type="catalog_export",
                title=f"R1 Fix5 {name}",
                status="succeeded",
                summary_json=json.dumps(payload),
            )
            db.add(run)
            await db.flush()
            run_ids[name] = run.id

        fallback_run = TaskRun(
            task_type="catalog_export",
            title="R1 Fix5 step fallback",
            status="succeeded",
            summary_json=json.dumps({"status": "done"}),
        )
        db.add(fallback_run)
        await db.flush()
        fallback_group = TaskGroup(
            task_run_id=fallback_run.id,
            group_key="export_file",
            title="R1 Fix5 fallback group",
            status="succeeded",
        )
        db.add(fallback_group)
        await db.flush()
        fallback_payload = {
            "status": "done",
            "category": "R1 Fix5 Step Fallback",
            "categories": ["R1 Fix5 Step Fallback"],
            "filename": "legacy-step-fallback.zip",
            "file_path": fallback_url,
        }
        db.add(TaskStep(
            task_run_id=fallback_run.id,
            task_group_id=fallback_group.id,
            step_key="catalog_export_template",
            step_type="catalog_export_template",
            status="succeeded",
            result_json=json.dumps(fallback_payload),
        ))
        await db.commit()
        run_ids["step_fallback"] = fallback_run.id

    transport = httpx.ASGITransport(app=app, client=("127.0.0.1", 41325))
    async with httpx.AsyncClient(
        transport=transport,
        base_url="http://testserver",
        follow_redirects=False,
    ) as client:
        with external_socket_guard():
            list_response = await client.get(
                "/api/task-runs",
                params={"view": "all", "q": "R1 Fix5", "page_size": 20},
            )
            assert list_response.status_code == 200, list_response.text
            list_items = {item["id"]: item for item in list_response.json()["items"]}
            for name in ("url_file_path", "oss_url", "local", "step_fallback"):
                assert "download_result" in list_items[run_ids[name]]["available_actions"], list_items[run_ids[name]]
            for name in ("explicit_false", "explicit_null", "stale_local"):
                assert "download_result" not in list_items[run_ids[name]]["available_actions"], list_items[run_ids[name]]

            for name in payloads:
                detail = await client.get(f"/api/task-runs/{run_ids[name]}")
                assert detail.status_code == 200, (name, detail.text)
                actions = detail.json()["available_actions"]
                if name in {"url_file_path", "oss_url", "local"}:
                    assert "download_result" in actions, (name, actions)
                else:
                    assert "download_result" not in actions, (name, actions)

            fallback_detail = await client.get(f"/api/task-runs/{run_ids['step_fallback']}")
            assert fallback_detail.status_code == 200, fallback_detail.text
            assert "download_result" in fallback_detail.json()["available_actions"], fallback_detail.json()

            for name, expected_url in (("url_file_path", url_file_path), ("oss_url", oss_url), ("step_fallback", fallback_url)):
                response = await client.get(f"/api/task-runs/{run_ids[name]}/download")
                assert response.status_code == 307, (name, response.status_code, response.text)
                assert response.headers["location"] == expected_url, (name, response.headers)

            local_download = await client.get(f"/api/task-runs/{run_ids['local']}/download")
            assert local_download.status_code == 200, local_download.text
            assert local_download.content == local_artifact.read_bytes(), local_download.content
            for name in ("explicit_false", "explicit_null", "stale_local"):
                response = await client.get(f"/api/task-runs/{run_ids[name]}/download")
                assert response.status_code == 400, (name, response.status_code, response.text)

            export_files = await client.get("/api/products/catalog/export-files", params={"page_size": 100})
            assert export_files.status_code == 200, export_files.text
            export_items = {item["task_id"]: item for item in export_files.json()["items"]}
            for name in ("url_file_path", "oss_url", "local", "step_fallback"):
                assert export_items[run_ids[name]]["can_download"] is True, export_items[run_ids[name]]
            for name in ("explicit_false", "explicit_null", "stale_local"):
                assert export_items[run_ids[name]]["can_download"] is False, export_items[run_ids[name]]


async def test_catalog_export_single_selector_and_consumption_contract(environment) -> None:
    import httpx
    from sqlalchemy import event

    from app.api import task_runs as task_runs_api
    from app.database import engine
    from app.main import app
    from app.models import TaskGroup, TaskRun, TaskStep

    export_root = environment.data_dir / "exports"
    export_root.mkdir(parents=True, exist_ok=True)
    valid_step_artifact = export_root / "fix6-step-artifact.zip"
    valid_step_artifact.write_bytes(b"fix6-step-artifact")
    outside_root = environment.data_dir / "outside-exports"
    outside_root.mkdir(parents=True, exist_ok=True)
    outside_artifact = outside_root / "fix6-outside.zip"
    outside_artifact.write_bytes(b"must-not-be-read-directly")
    redirect_url = "https://downloads.invalid/fix6-outside-redirect.zip"
    newest_material_url = "https://downloads.invalid/fix6-newest-material.zip"

    step_payload = {
        "status": "done",
        "filename": valid_step_artifact.name,
        "file_path": str(valid_step_artifact),
    }
    cases = {
        "summary_false_step": ({"status": "done", "artifact_available": False}, step_payload),
        "summary_null_step": ({"status": "done", "artifact_available": None}, step_payload),
        "summary_rows_step": ({
            "status": "done",
            "requested_count": 1,
            "success_count": 1,
            "skipped_count": 0,
            "failed_count": 0,
            "rows": [{"catalog_id": 9101, "status": "exported"}],
        }, step_payload),
        "status_only_step": ({"status": "done"}, step_payload),
        "outside_local": ({
            "status": "done",
            "filename": outside_artifact.name,
            "file_path": str(outside_artifact),
        }, None),
        "outside_object_key": ({
            "status": "done",
            "filename": "fix6-object-key.zip",
            "file_path": str(outside_artifact),
            "oss_object_key": "exports/fix6-object-key.zip",
        }, None),
        "outside_redirect": ({
            "status": "done",
            "filename": "fix6-outside-redirect.zip",
            "file_path": str(outside_artifact),
            "oss_url": redirect_url,
        }, None),
        "explicit_true_no_source": ({
            "status": "done",
            "artifact_available": True,
        }, None),
    }
    run_ids: dict[str, int] = {}
    async with environment.session_factory() as db:
        for name, (summary_payload, catalog_step_payload) in cases.items():
            run = TaskRun(
                task_type="catalog_export",
                title=f"R1 Fix6 {name}",
                status="succeeded",
                summary_json=json.dumps(summary_payload),
            )
            db.add(run)
            await db.flush()
            run_ids[name] = run.id
            if catalog_step_payload is None:
                continue
            group = TaskGroup(
                task_run_id=run.id,
                group_key="export_file",
                title=f"R1 Fix6 {name} group",
                status="succeeded",
            )
            db.add(group)
            await db.flush()
            material_step = TaskStep(
                task_run_id=run.id,
                task_group_id=group.id,
                step_key="catalog_export_template",
                step_type="catalog_export_template",
                status="succeeded",
                result_json=json.dumps(catalog_step_payload),
            )
            db.add(material_step)
            await db.flush()
            if name == "status_only_step":
                db.add(TaskStep(
                    task_run_id=run.id,
                    task_group_id=group.id,
                    step_key="catalog_export_template_newest_material",
                    step_type="catalog_export_template",
                    status="succeeded",
                    result_json=json.dumps({
                        "status": "done",
                        "filename": "fix6-newest-material.zip",
                        "oss_url": newest_material_url,
                    }),
                ))
                await db.flush()
                db.add(TaskStep(
                    task_run_id=run.id,
                    task_group_id=group.id,
                    step_key="catalog_export_template_newest_nonmaterial",
                    step_type="catalog_export_template",
                    status="succeeded",
                    result_json=json.dumps({"status": "done"}),
                ))
                await db.flush()
        await db.commit()

    transport = httpx.ASGITransport(app=app, client=("127.0.0.1", 41326))
    async with httpx.AsyncClient(
        transport=transport,
        base_url="http://testserver",
        follow_redirects=False,
    ) as client:
        with external_socket_guard():
            list_sql: list[str] = []

            def record_list_sql(_connection, _cursor, statement, _parameters, _context, _executemany):
                list_sql.append(" ".join(statement.lower().split()))

            event.listen(engine.sync_engine, "before_cursor_execute", record_list_sql)
            try:
                list_response = await client.get(
                    "/api/task-runs",
                    params={"view": "all", "q": "R1 Fix6", "page_size": 20},
                )
            finally:
                event.remove(engine.sync_engine, "before_cursor_execute", record_list_sql)
            assert list_response.status_code == 200, list_response.text
            material_step_queries = [
                statement
                for statement in list_sql
                if "task_steps.task_run_id" in statement and "task_steps.result_json" in statement
            ]
            assert len(material_step_queries) == 1, material_step_queries
            assert "json_" not in material_step_queries[0], material_step_queries[0]
            assert "row_number()" not in material_step_queries[0], material_step_queries[0]
            assert "order by task_steps.task_run_id asc" in material_step_queries[0], material_step_queries[0]
            assert "task_steps.id desc" in material_step_queries[0], material_step_queries[0]
            list_items = {item["id"]: item for item in list_response.json()["items"]}
            assert len(list_items) >= 2, list_items
            for name in ("summary_false_step", "summary_null_step", "summary_rows_step"):
                assert "download_result" not in list_items[run_ids[name]]["available_actions"], (name, list_items[run_ids[name]])
            for name in ("status_only_step", "outside_object_key", "outside_redirect"):
                assert "download_result" in list_items[run_ids[name]]["available_actions"], (name, list_items[run_ids[name]])

            for name in cases:
                detail = await client.get(f"/api/task-runs/{run_ids[name]}")
                assert detail.status_code == 200, (name, detail.text)
                actions = detail.json()["available_actions"]
                if name in {"status_only_step", "outside_object_key", "outside_redirect"}:
                    assert "download_result" in actions, (name, actions, detail.json())
                else:
                    assert "download_result" not in actions, (name, actions, detail.json())

            for name in ("outside_local", "explicit_true_no_source"):
                assert "download_result" not in list_items[run_ids[name]]["available_actions"], (name, list_items[run_ids[name]])

            export_files = await client.get("/api/products/catalog/export-files", params={"page_size": 100})
            assert export_files.status_code == 200, export_files.text
            export_items = {item["task_id"]: item for item in export_files.json()["items"]}
            for name in ("summary_false_step", "summary_null_step", "summary_rows_step", "outside_local", "explicit_true_no_source"):
                assert export_items[run_ids[name]]["can_download"] is False, (name, export_items[run_ids[name]])
            for name in ("status_only_step", "outside_object_key", "outside_redirect"):
                assert export_items[run_ids[name]]["can_download"] is True, (name, export_items[run_ids[name]])
            assert export_items[run_ids["status_only_step"]]["oss_url"] == newest_material_url, export_items[run_ids["status_only_step"]]

            for name in ("summary_false_step", "summary_null_step", "summary_rows_step", "outside_local", "explicit_true_no_source"):
                response = await client.get(f"/api/task-runs/{run_ids[name]}/download")
                assert response.status_code == 400, (name, response.status_code, response.text)

            status_only = await client.get(f"/api/task-runs/{run_ids['status_only_step']}/download")
            assert status_only.status_code == 307, status_only.text
            assert status_only.headers["location"] == newest_material_url, status_only.headers

            outside_redirect = await client.get(f"/api/task-runs/{run_ids['outside_redirect']}/download")
            assert outside_redirect.status_code == 307, outside_redirect.text
            assert outside_redirect.headers["location"] == redirect_url, outside_redirect.headers

    download_calls: list[tuple[str, Path]] = []
    original_download = task_runs_api.download_private_file

    def fake_download(object_key: str, target_path: Path) -> dict:
        target = Path(target_path)
        download_calls.append((object_key, target))
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_bytes(b"fix6-object-key-download")
        return {"object_key": object_key, "path": str(target)}

    task_runs_api.download_private_file = fake_download
    try:
        async with httpx.AsyncClient(
            transport=transport,
            base_url="http://testserver",
            follow_redirects=False,
        ) as client:
            with external_socket_guard():
                object_download = await client.get(f"/api/task-runs/{run_ids['outside_object_key']}/download")
    finally:
        task_runs_api.download_private_file = original_download
    assert object_download.status_code == 200, object_download.text
    assert object_download.content == b"fix6-object-key-download", object_download.content
    assert len(download_calls) == 1, download_calls
    assert download_calls[0][0] == "exports/fix6-object-key.zip", download_calls
    downloaded_path = download_calls[0][1].resolve()
    assert export_root.resolve() in downloaded_path.parents, downloaded_path
    assert downloaded_path != outside_artifact.resolve(), downloaded_path


async def test_catalog_export_download_and_export_center_api(environment) -> None:
    import httpx

    from app.main import app
    from app.models import TaskRun
    from app.services.offline_tasks import _catalog_export_result_payload

    export_dir = environment.data_dir / "exports" / "api-outcomes"
    export_dir.mkdir(parents=True, exist_ok=True)
    artifact = export_dir / "r1-api-outcomes.zip"
    artifact.write_bytes(b"r1-api-outcomes")
    category = "R1 API Outcome Category"

    def payload_for(catalog_id: int, row_status: str) -> dict:
        return _catalog_export_result_payload(
            category=category,
            categories=[category],
            template_name="CHAIR_SOFA.xlsm",
            template_path="/fixture/CHAIR_SOFA.xlsm",
            catalog_ids=[catalog_id],
            report_rows=[{
                "商品资料ID": catalog_id,
                "商品ID": catalog_id + 1000,
                "商品Code": f"R1-API-{catalog_id}",
                "状态": row_status,
                "原因": "R1 API outcome fixture",
            }],
            created_at=datetime.now(),
            filename=artifact.name,
            file_path=str(artifact),
            file_size=artifact.stat().st_size,
        )

    done_payload = payload_for(601, "已导出")
    partial_payload = _catalog_export_result_payload(
        category=category,
        categories=[category],
        template_name="CHAIR_SOFA.xlsm",
        template_path="/fixture/CHAIR_SOFA.xlsm",
        catalog_ids=[602, 603],
        report_rows=[
            {"商品资料ID": 602, "商品ID": 1602, "商品Code": "R1-API-PARTIAL", "状态": "已导出"},
            {"商品资料ID": 603, "商品ID": 1603, "商品Code": "R1-API-SKIP", "状态": "跳过", "原因": "保护门"},
        ],
        created_at=datetime.now(),
        filename=artifact.name,
        file_path=str(artifact),
        file_size=artifact.stat().st_size,
    )
    failed_payload = payload_for(604, "跳过")
    assert done_payload["status"] == "done", done_payload
    assert partial_payload["status"] == "partial_failed", partial_payload
    assert failed_payload["status"] == "failed", failed_payload
    assert failed_payload["artifact_available"] is True, failed_payload

    run_ids: dict[str, int] = {}
    async with environment.session_factory() as db:
        for name, raw_status, payload in [
            ("done", "succeeded", done_payload),
            ("partial", "partial_failed", partial_payload),
            ("failed", "failed", failed_payload),
        ]:
            run = TaskRun(
                task_type="catalog_export",
                title=f"R1 API outcome {name}",
                status=raw_status,
                summary_json=json.dumps(payload),
            )
            db.add(run)
            await db.flush()
            run_ids[name] = run.id
        await db.commit()

    transport = httpx.ASGITransport(app=app, client=("127.0.0.1", 41323))
    async with httpx.AsyncClient(transport=transport, base_url="http://testserver") as client:
        for name in ("done", "partial"):
            response = await client.get(f"/api/task-runs/{run_ids[name]}/download")
            assert response.status_code == 200, (name, response.text)
            assert response.content == artifact.read_bytes(), name
        failed_download = await client.get(f"/api/task-runs/{run_ids['failed']}/download")
        assert failed_download.status_code == 400, failed_download.text

        export_files = await client.get(
            "/api/products/catalog/export-files",
            params={"category": category, "page_size": 20},
        )
        assert export_files.status_code == 200, export_files.text
        body = export_files.json()
        assert body["total"] == 3, body
        items = {item["task_id"]: item for item in body["items"]}
        assert set(items) == set(run_ids.values()), body
        assert items[run_ids["done"]]["can_download"] is True, items[run_ids["done"]]
        assert items[run_ids["partial"]]["can_download"] is True, items[run_ids["partial"]]
        assert items[run_ids["failed"]]["can_download"] is False, items[run_ids["failed"]]
        assert len(items[run_ids["partial"]]["rows"]) == 2, items[run_ids["partial"]]
        assert items[run_ids["partial"]]["rows"][1]["reason"] == "保护门", items[run_ids["partial"]]


async def test_pre_r1_offline_task_export_and_executor_reuse(environment) -> None:
    import httpx

    from app.api import offline_tasks as offline_tasks_api
    from app.main import app
    from app.models import OfflineTask, OfflineTaskStep
    from app.services import offline_tasks

    artifact = environment.data_dir / "exports" / "pre-r1-offline.zip"
    artifact.parent.mkdir(parents=True, exist_ok=True)
    artifact.write_bytes(b"pre-r1-offline")
    category = "R1 Pre-R1 Offline Category"
    unavailable_only_category = "R1 Unavailable Only Offline Category"
    legacy_remote_url = "https://downloads.invalid/pre-r1-offline-url.zip"
    object_key = "exports/pre-r1-offline-object-key.zip"
    outside_root = environment.data_dir / "outside-offline-exports"
    outside_root.mkdir(parents=True, exist_ok=True)
    outside_artifact = outside_root / "pre-r1-offline-outside.zip"
    outside_artifact.write_bytes(b"outside-root-must-not-be-served")
    legacy_payload = {
        "status": "partial_failed",
        "category": category,
        "categories": [category],
        "catalog_product_ids": [801, 802],
        "requested_count": 2,
        "success_count": 1,
        "skipped_count": 1,
        "failed_count": 0,
        "filename": "pre-r1-offline-url.zip",
        "file_path": legacy_remote_url,
        "file_size": artifact.stat().st_size,
        "rows": [
            {"catalog_id": 801, "item_code": "PRE-R1-OFFLINE-DONE", "category": category, "status": "exported"},
            {
                "catalog_id": 802,
                "item_code": "PRE-R1-OFFLINE-SKIP",
                "category": category,
                "status": "skipped",
                "reason": "保护门",
            },
        ],
    }
    assert "artifact_available" not in legacy_payload, legacy_payload

    async with environment.session_factory() as db:
        api_task = OfflineTask(
            task_type="catalog_export",
            title="R1 pre-R1 offline API",
            status="partial_failed",
            total_steps=1,
            success_steps=1,
            result_json=json.dumps(legacy_payload),
        )
        db.add(api_task)
        await db.flush()
        db.add(OfflineTaskStep(
            task_id=api_task.id,
            step_type="catalog_export_template",
            title="R1 pre-R1 offline API step",
            status="done",
            progress_current=2,
            progress_total=2,
            result_json=json.dumps(legacy_payload),
        ))

        local_payload = {
            **legacy_payload,
            "status": "done",
            "category": "R1 Pre-R1 Offline Local",
            "categories": ["R1 Pre-R1 Offline Local"],
            "filename": artifact.name,
            "file_path": str(artifact),
        }
        local_task = OfflineTask(
            task_type="catalog_export",
            title="R1 pre-R1 offline local",
            status="done",
            total_steps=1,
            success_steps=1,
            result_json=json.dumps(local_payload),
        )
        db.add(local_task)
        await db.flush()

        object_key_payload = {
            **legacy_payload,
            "status": "done",
            "category": "R1 Pre-R1 Offline Object Key",
            "categories": ["R1 Pre-R1 Offline Object Key"],
            "filename": "pre-r1-offline-object-key.zip",
            "file_path": None,
            "oss_object_key": object_key,
        }
        object_key_task = OfflineTask(
            task_type="catalog_export",
            title="R1 pre-R1 offline object key",
            status="done",
            total_steps=1,
            success_steps=1,
            result_json=json.dumps(object_key_payload),
        )
        db.add(object_key_task)
        await db.flush()

        status_only_task = OfflineTask(
            task_type="catalog_export",
            title="R1 pre-R1 offline status-only step fallback",
            status="done",
            total_steps=1,
            success_steps=1,
            result_json=json.dumps({"status": "done"}),
        )
        db.add(status_only_task)
        await db.flush()
        db.add(OfflineTaskStep(
            task_id=status_only_task.id,
            step_type="catalog_export_template",
            title="R1 pre-R1 offline status-only step fallback",
            status="done",
            progress_current=1,
            progress_total=1,
            result_json=json.dumps(local_payload),
        ))

        unavailable_task_ids: dict[str, int] = {}
        unavailable_payloads = {
            "explicit_false": {
                **local_payload,
                "category": unavailable_only_category,
                "categories": [unavailable_only_category],
                "artifact_available": False,
                "oss_url": legacy_remote_url,
            },
            "explicit_null": {**local_payload, "artifact_available": None, "oss_url": legacy_remote_url},
            "stale_local": {
                **local_payload,
                "filename": "pre-r1-offline-stale.zip",
                "file_path": str(artifact.with_name("pre-r1-offline-stale.zip")),
            },
            "outside_root": {
                **local_payload,
                "filename": outside_artifact.name,
                "file_path": str(outside_artifact),
            },
        }
        for name, unavailable_payload in unavailable_payloads.items():
            unavailable_task = OfflineTask(
                task_type="catalog_export",
                title=f"R1 pre-R1 offline {name}",
                status="done",
                total_steps=1,
                success_steps=1,
                result_json=json.dumps(unavailable_payload),
            )
            db.add(unavailable_task)
            await db.flush()
            unavailable_task_ids[name] = unavailable_task.id

        rerun_payload = {
            **legacy_payload,
            "status": "done",
            "category": "R1 Pre-R1 Offline Rerun",
            "categories": ["R1 Pre-R1 Offline Rerun"],
            "catalog_product_ids": [999999999],
            "requested_count": 1,
            "success_count": 1,
            "skipped_count": 0,
            "rows": [{
                "catalog_id": 999999999,
                "item_code": "PRE-R1-RERUN",
                "category": "R1 Pre-R1 Offline Rerun",
                "status": "exported",
            }],
        }
        rerun_task = OfflineTask(
            task_type="catalog_export",
            title="R1 pre-R1 offline executor rerun",
            status="pending",
            total_steps=1,
        )
        db.add(rerun_task)
        await db.flush()
        rerun_step = OfflineTaskStep(
            task_id=rerun_task.id,
            step_type="catalog_export_template",
            title="R1 pre-R1 offline executor rerun step",
            status="pending",
            progress_current=0,
            progress_total=1,
            payload_json=json.dumps({
                "catalog_product_ids": [999999999],
                "category": "R1 Pre-R1 Offline Rerun",
                "categories": ["R1 Pre-R1 Offline Rerun"],
            }),
            result_json=json.dumps(rerun_payload),
        )
        db.add(rerun_step)
        await db.commit()
        api_task_id = api_task.id
        local_task_id = local_task.id
        object_key_task_id = object_key_task.id
        status_only_task_id = status_only_task.id
        rerun_task_id = rerun_task.id
        rerun_step_id = rerun_step.id

    transport = httpx.ASGITransport(app=app, client=("127.0.0.1", 41324))
    async with httpx.AsyncClient(
        transport=transport,
        base_url="http://testserver",
        follow_redirects=False,
    ) as client:
        with external_socket_guard():
            offline_list = await client.get(
                "/api/offline-tasks",
                params={"task_type": "catalog_export", "page_size": 100},
            )
            assert offline_list.status_code == 200, offline_list.text
            offline_items = {item["id"]: item for item in offline_list.json()["items"]}
            for task_id in (api_task_id, local_task_id, object_key_task_id, status_only_task_id):
                assert offline_items[task_id]["can_download"] is True, offline_items[task_id]
            for name, task_id in unavailable_task_ids.items():
                assert offline_items[task_id]["can_download"] is False, (name, offline_items[task_id])

            for task_id in (api_task_id, local_task_id, object_key_task_id, status_only_task_id):
                detail = await client.get(f"/api/offline-tasks/{task_id}")
                assert detail.status_code == 200, detail.text
                assert detail.json()["can_download"] is True, detail.json()
            for name, task_id in unavailable_task_ids.items():
                detail = await client.get(f"/api/offline-tasks/{task_id}")
                assert detail.status_code == 200, (name, detail.text)
                assert detail.json()["can_download"] is False, (name, detail.json())

            export_files = await client.get(
                "/api/products/catalog/export-files",
                params={"category": category, "page_size": 20},
            )
            assert export_files.status_code == 200, export_files.text
            items = {item["task_id"]: item for item in export_files.json()["items"]}
            assert items[api_task_id]["can_download"] is True, items[api_task_id]

            categories = await client.get("/api/products/catalog/export-categories")
            assert categories.status_code == 200, categories.text
            exported_categories = {item["category"]: item for item in categories.json()["exported"]}
            assert exported_categories[category]["count"] == 1, exported_categories
            assert unavailable_only_category not in exported_categories, exported_categories

            remote_download = await client.get(f"/api/offline-tasks/{api_task_id}/download")
            assert remote_download.status_code == 307, remote_download.text
            assert remote_download.headers["location"] == legacy_remote_url, remote_download.headers
            local_download = await client.get(f"/api/offline-tasks/{local_task_id}/download")
            assert local_download.status_code == 200, local_download.text
            assert local_download.content == artifact.read_bytes(), local_download.content
            status_only_download = await client.get(f"/api/offline-tasks/{status_only_task_id}/download")
            assert status_only_download.status_code == 200, status_only_download.text
            assert status_only_download.content == artifact.read_bytes(), status_only_download.content
            for name, task_id in unavailable_task_ids.items():
                response = await client.get(f"/api/offline-tasks/{task_id}/download")
                assert response.status_code == 400, (name, response.status_code, response.text)

    download_calls: list[tuple[str, Path]] = []
    original_download = offline_tasks_api.download_private_file

    def fake_download(object_key_value: str, target_path: Path) -> dict:
        target = Path(target_path)
        download_calls.append((object_key_value, target))
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_bytes(b"pre-r1-offline-object-key")
        return {"object_key": object_key_value, "path": str(target)}

    offline_tasks_api.download_private_file = fake_download
    try:
        async with httpx.AsyncClient(
            transport=transport,
            base_url="http://testserver",
            follow_redirects=False,
        ) as client:
            with external_socket_guard():
                object_download = await client.get(f"/api/offline-tasks/{object_key_task_id}/download")
    finally:
        offline_tasks_api.download_private_file = original_download
    assert object_download.status_code == 200, object_download.text
    assert object_download.content == b"pre-r1-offline-object-key", object_download.content
    assert len(download_calls) == 1, download_calls
    assert download_calls[0][0] == object_key, download_calls
    assert (environment.data_dir / "exports").resolve() in download_calls[0][1].resolve().parents, download_calls

    upload_calls: list[tuple[Path, str]] = []
    original_upload = offline_tasks.upload_private_file

    def fake_upload(path: Path, object_key: str) -> dict:
        upload_calls.append((Path(path), object_key))
        return {"object_key": object_key, "url": f"https://oss.invalid/{object_key}"}

    offline_tasks.upload_private_file = fake_upload
    try:
        with external_socket_guard():
            await offline_tasks._execute_offline_task(rerun_task_id, [rerun_step_id])
    finally:
        offline_tasks.upload_private_file = original_upload
    assert upload_calls == [], upload_calls

    async with environment.session_factory() as db:
        reused_task = await db.get(OfflineTask, rerun_task_id)
        reused_step = await db.get(OfflineTaskStep, rerun_step_id)
        assert reused_task is not None and reused_task.status == "done", reused_task
        assert reused_step is not None and reused_step.status == "done", reused_step
        assert json.loads(reused_step.result_json or "{}") == offline_tasks.normalize_catalog_export_response(
            rerun_payload
        ), reused_step.result_json


async def test_catalog_export_zip_recovery_preserves_outcome(environment) -> None:
    from io import BytesIO
    from zipfile import ZipFile

    from openpyxl import Workbook

    from app.api.products import _summary_workbook
    from app.services.offline_tasks import (
        _catalog_export_result_ready,
        _recover_catalog_export_result_from_file,
    )

    export_dir = environment.data_dir / "exports" / "recovery-mixed"
    export_dir.mkdir(parents=True, exist_ok=True)
    archive_path = export_dir / "mixed.zip"
    dummy = Workbook()
    dummy.active.append(["not", "the", "report"])
    dummy_stream = BytesIO()
    dummy.save(dummy_stream)
    report_rows = [
        {"商品资料ID": 701, "商品ID": 1701, "商品Code": "R1-RECOVER-DONE", "状态": "已导出"},
        {"商品资料ID": 702, "商品ID": 1702, "商品Code": "R1-RECOVER-SKIP", "状态": "跳过", "原因": "保护门"},
    ]
    with ZipFile(archive_path, "w") as archive:
        archive.writestr("Sofas.xlsx", dummy_stream.getvalue())
        archive.writestr("导出报告.xlsx", _summary_workbook(report_rows))

    recovered = _recover_catalog_export_result_from_file(
        export_dir=export_dir,
        category="Sofas & Couches",
        categories=["Sofas & Couches"],
        template_name="CHAIR_SOFA.xlsm",
        template_path="/fixture/CHAIR_SOFA.xlsm",
        catalog_ids=[701, 702],
    )
    assert recovered is not None, recovered
    assert recovered["status"] == "partial_failed", recovered
    assert recovered["success_count"] == 1, recovered
    assert recovered["skipped_count"] == 1, recovered
    assert recovered["artifact_available"] is True, recovered
    assert _catalog_export_result_ready(recovered) is True, recovered

    failed_dir = environment.data_dir / "exports" / "recovery-failed"
    failed_dir.mkdir(parents=True, exist_ok=True)
    failed_path = failed_dir / "failed.zip"
    failed_rows = [
        {"商品资料ID": 703, "商品ID": 1703, "商品Code": "R1-RECOVER-FAILED", "状态": "跳过", "原因": "保护门"},
    ]
    with ZipFile(failed_path, "w") as archive:
        archive.writestr("导出报告.xlsx", _summary_workbook(failed_rows))
    failed_recovery = _recover_catalog_export_result_from_file(
        export_dir=failed_dir,
        category="Sofas & Couches",
        categories=["Sofas & Couches"],
        template_name="CHAIR_SOFA.xlsm",
        template_path="/fixture/CHAIR_SOFA.xlsm",
        catalog_ids=[703],
    )
    assert failed_recovery is None, failed_recovery
    assert _catalog_export_result_ready({
        "status": "failed",
        "artifact_available": True,
        "filename": failed_path.name,
        "file_path": str(failed_path),
    }) is False


async def test_catalog_export_outcome_rejects_invalid_topology(environment) -> None:
    from sqlalchemy import func, select

    from app.models import TaskGroup, TaskRun, TaskStep, TaskStepEvent
    from app.services.offline_tasks import _catalog_export_result_payload
    from app.task_runtime.registry import TaskOutcomeContractError, TaskWorkerOutcome
    from app.task_runtime.scheduler import _project_single_step_worker_outcome

    for topology in ("extra_group", "extra_step"):
        catalog_id = (await _seed_catalog_export_products(environment, [
            {"item_code": f"R1-TOPOLOGY-{topology}"},
        ]))[0]
        run_id, step_id = await _prepare_catalog_run(environment, [catalog_id])
        async with environment.session_factory() as db:
            step = await db.get(TaskStep, step_id)
            assert step is not None
            if topology == "extra_group":
                db.add(TaskGroup(
                    task_run_id=run_id,
                    group_key="unexpected_group",
                    title="unexpected group",
                    status="pending",
                ))
            else:
                db.add(TaskStep(
                    task_run_id=run_id,
                    task_group_id=step.task_group_id,
                    step_key="unexpected_step",
                    step_type="catalog_export_template",
                    status="pending",
                ))
            await db.commit()

        artifact = environment.data_dir / "exports" / "topology" / f"{topology}.zip"
        artifact.parent.mkdir(parents=True, exist_ok=True)
        artifact.write_bytes(topology.encode())
        payload = _catalog_export_result_payload(
            category="Sofas & Couches",
            categories=["Sofas & Couches"],
            template_name="CHAIR_SOFA.xlsm",
            template_path="/fixture/CHAIR_SOFA.xlsm",
            catalog_ids=[catalog_id],
            report_rows=[{"商品资料ID": catalog_id, "状态": "已导出"}],
            created_at=datetime.now(),
            filename=artifact.name,
            file_path=str(artifact),
            file_size=artifact.stat().st_size,
        )
        outcome = TaskWorkerOutcome(
            payload=payload,
            terminal_status="succeeded",
            event_type="status",
            event_message="导出文件生成完成",
            propagate_single_step_run=True,
        )
        async with environment.session_factory() as db:
            run = await db.get(TaskRun, run_id)
            step = await db.get(TaskStep, step_id)
            group = await db.get(TaskGroup, step.task_group_id if step else 0)
            assert run is not None and group is not None and step is not None
            try:
                await _project_single_step_worker_outcome(
                    db,
                    run=run,
                    group=group,
                    step=step,
                    outcome=outcome,
                )
            except TaskOutcomeContractError:
                await db.rollback()
            else:
                raise AssertionError(f"invalid topology propagated terminal outcome: {topology}")

        async with environment.session_factory() as observer:
            run = await observer.get(TaskRun, run_id)
            step = await observer.get(TaskStep, step_id)
            event_count = await observer.scalar(
                select(func.count(TaskStepEvent.id)).where(TaskStepEvent.task_run_id == run_id)
            )
            assert run is not None and run.status == "pending" and run.summary_json is None, (topology, run)
            assert step is not None and step.status == "ready" and step.result_json is None, (topology, step)
            assert event_count == 0, (topology, event_count)


async def test_real_worker_projection_faults_rollback_catalog_and_progress(environment) -> None:
    from sqlalchemy import event, func, select
    from sqlalchemy.exc import OperationalError

    from app.database import engine
    from app.models import CatalogProduct, TaskRun, TaskStep, TaskStepEvent
    from app.task_runtime import catalog_export_workers
    from app.task_runtime.catalog_export_workers import register_catalog_export_workers
    from app.task_runtime.scheduler import drain_ready_steps

    for stage in ("topology_query", "terminal_event_flush", "commit"):
        catalog_id = (await _seed_catalog_export_products(environment, [
            {"item_code": f"R1-REAL-ROLLBACK-{stage}"},
        ]))[0]
        run_id, step_id = await _prepare_catalog_run(environment, [catalog_id])
        upload_calls: list[Path] = []
        original_upload = catalog_export_workers.upload_private_file
        installed: list[tuple[object, str, object]] = []
        commit_state = {"raised": False}
        terminal_event_state = {"raised": False}
        topology_state = {"raised": False}

        def fail_projector_query(_connection, _cursor, statement, parameters, _context, _executemany):
            normalized = " ".join(statement.lower().split())
            if (
                not topology_state["raised"]
                and "select task_groups.id" in normalized
                and "where task_groups.task_run_id" in normalized
            ):
                topology_state["raised"] = True
                raise OperationalError(statement, parameters, RuntimeError("R1 injected topology query failure"))

        def fail_terminal_event_flush(_connection, _cursor, statement, parameters, _context, _executemany):
            if "insert into task_step_events" not in statement.lower():
                return
            rendered = repr(parameters)
            if (
                not terminal_event_state["raised"]
                and "导出文件生成完成" in rendered
                and "status" in rendered
                and "progress" not in rendered
            ):
                terminal_event_state["raised"] = True
                raise OperationalError(statement, parameters, RuntimeError("R1 injected terminal event flush failure"))

        def fail_commit_once(_connection):
            if commit_state["raised"]:
                return
            commit_state["raised"] = True
            raise OperationalError("COMMIT", {}, RuntimeError("R1 injected commit failure"))

        def fake_upload(path: Path, object_key: str) -> dict:
            upload_calls.append(Path(path))
            if stage == "topology_query":
                event.listen(engine.sync_engine, "before_cursor_execute", fail_projector_query)
                installed.append((engine.sync_engine, "before_cursor_execute", fail_projector_query))
            elif stage == "terminal_event_flush":
                event.listen(engine.sync_engine, "before_cursor_execute", fail_terminal_event_flush)
                installed.append((engine.sync_engine, "before_cursor_execute", fail_terminal_event_flush))
            else:
                event.listen(engine.sync_engine, "commit", fail_commit_once)
                installed.append((engine.sync_engine, "commit", fail_commit_once))
            return {"object_key": object_key, "url": f"https://oss.invalid/{object_key}"}

        register_catalog_export_workers()
        catalog_export_workers.upload_private_file = fake_upload
        try:
            with external_socket_guard():
                await drain_ready_steps()
        finally:
            for target, event_name, listener in installed:
                event.remove(target, event_name, listener)
            catalog_export_workers.upload_private_file = original_upload

        assert len(upload_calls) == 1 and upload_calls[0].is_file(), (stage, upload_calls)
        async with environment.session_factory() as db:
            run = await db.get(TaskRun, run_id)
            step = await db.get(TaskStep, step_id)
            catalog = await db.get(CatalogProduct, catalog_id)
            progress_events = await db.scalar(
                select(func.count(TaskStepEvent.id)).where(
                    TaskStepEvent.task_run_id == run_id,
                    TaskStepEvent.event_type == "progress",
                )
            )
            terminal_payload_events = await db.scalar(
                select(func.count(TaskStepEvent.id)).where(
                    TaskStepEvent.task_run_id == run_id,
                    TaskStepEvent.message == "导出文件生成完成",
                )
            )
            assert run is not None and run.status == "failed" and run.summary_json is None, (stage, run)
            assert step is not None and step.status == "failed" and step.result_json is None, (stage, step)
            assert step.error_message and step.error_message.startswith("OperationalError:"), (stage, step)
            assert catalog is not None and catalog.exported_at is None, (stage, catalog)
            assert catalog.export_task_id is None and catalog.export_file_path is None, (stage, catalog)
            assert progress_events == 0, (stage, progress_events)
            assert terminal_payload_events == 0, (stage, terminal_payload_events)


async def test_catalog_export_terminal_projection_rolls_back_atomically(environment) -> None:
    from sqlalchemy import event, func, select

    from app.api.schemas import OfflineTaskCatalogExportRequest
    from app.models import CatalogProduct, TaskGroup, TaskRun, TaskStep, TaskStepEvent
    from app.services.offline_tasks import _catalog_export_result_payload
    from app.task_planners.catalog_export import create_catalog_export_runs
    from app.task_runtime.registry import TaskWorkerOutcome
    from app.task_runtime.scheduler import _project_single_step_worker_outcome

    for stage in ("catalog_flush", "terminal_before_flush", "terminal_after_flush", "commit"):
        catalog_id = (await _seed_catalog_export_products(environment, [
            {"item_code": f"R1-ROLLBACK-{stage}"},
        ]))[0]
        async with environment.session_factory() as db:
            runs, errors = await create_catalog_export_runs(
                db,
                OfflineTaskCatalogExportRequest(catalog_product_ids=[catalog_id]),
                created_by="r1_catalog_rollback_test",
                auto_start=False,
            )
            assert not errors, errors
            run_id = runs[0].id
            group = (
                await db.execute(select(TaskGroup).where(TaskGroup.task_run_id == run_id))
            ).scalar_one()
            step = (
                await db.execute(select(TaskStep).where(TaskStep.task_run_id == run_id))
            ).scalar_one()
            run = await db.get(TaskRun, run_id)
            assert run is not None
            run.status = "running"
            group.status = "running"
            step.status = "running"
            await db.commit()
            group_id = group.id
            step_id = step.id

        artifact = environment.data_dir / "exports" / "rollback" / f"{stage}.zip"
        artifact.parent.mkdir(parents=True, exist_ok=True)
        artifact.write_bytes(stage.encode())
        payload = _catalog_export_result_payload(
            category="Sofas & Couches",
            categories=["Sofas & Couches"],
            template_name="CHAIR_SOFA.xlsm",
            template_path="/fixture/CHAIR_SOFA.xlsm",
            catalog_ids=[catalog_id],
            report_rows=[{
                "商品资料ID": catalog_id,
                "商品ID": catalog_id + 2000,
                "商品Code": f"R1-ROLLBACK-{stage}",
                "状态": "已导出",
            }],
            created_at=datetime.now(),
            filename=artifact.name,
            file_path=str(artifact),
            file_size=artifact.stat().st_size,
        )
        outcome = TaskWorkerOutcome(
            payload=payload,
            terminal_status="succeeded",
            event_type="status",
            event_message="导出文件生成完成",
            propagate_single_step_run=True,
        )

        async with environment.session_factory() as db:
            run = await db.get(TaskRun, run_id)
            group = await db.get(TaskGroup, group_id)
            step = await db.get(TaskStep, step_id)
            catalog = await db.get(CatalogProduct, catalog_id)
            assert run is not None and group is not None and step is not None and catalog is not None
            catalog.exported_at = datetime.now()
            catalog.export_task_id = run_id
            catalog.export_file_path = str(artifact)

            target = db.sync_session
            event_name: str
            listener = None
            if stage == "catalog_flush":
                event_name = "after_flush"

                def listener(_session, _flush_context):
                    raise RuntimeError("R1 injected after CatalogProduct flush")
            else:
                db.sync_session.autoflush = False
                if stage == "terminal_before_flush":
                    event_name = "before_flush"

                    def listener(_session, _flush_context, _instances):
                        raise RuntimeError("R1 injected after terminal assignments")
                elif stage == "terminal_after_flush":
                    event_name = "after_flush"

                    def listener(_session, _flush_context):
                        raise RuntimeError("R1 injected after terminal event flush")
                else:
                    event_name = "commit"
                    target = db.sync_session.get_bind()

                    def listener(_connection):
                        raise RuntimeError("R1 injected database commit failure")

            event.listen(target, event_name, listener)
            try:
                try:
                    await _project_single_step_worker_outcome(
                        db,
                        run=run,
                        group=group,
                        step=step,
                        outcome=outcome,
                    )
                except RuntimeError as exc:
                    assert "R1 injected" in str(exc), (stage, exc)
                else:
                    raise AssertionError(f"fault injection did not fire: {stage}")
                await db.rollback()
            finally:
                event.remove(target, event_name, listener)

        async with environment.session_factory() as observer:
            run = await observer.get(TaskRun, run_id)
            group = await observer.get(TaskGroup, group_id)
            step = await observer.get(TaskStep, step_id)
            catalog = await observer.get(CatalogProduct, catalog_id)
            terminal_events = await observer.scalar(
                select(func.count(TaskStepEvent.id)).where(
                    TaskStepEvent.task_run_id == run_id,
                    TaskStepEvent.event_type.in_(("status", "warning", "error")),
                )
            )
            assert run is not None and run.status == "running" and run.summary_json is None, (stage, run)
            assert group is not None and group.status == "running" and group.summary_json is None, (stage, group)
            assert step is not None and step.status == "running" and step.result_json is None, (stage, step)
            assert catalog is not None and catalog.exported_at is None, (stage, catalog)
            assert catalog.export_task_id is None and catalog.export_file_path is None, (stage, catalog)
            assert terminal_events == 0, (stage, terminal_events)


async def main() -> None:
    from testing.r1_mysql import isolated_r1_mysql

    async with isolated_r1_mysql(ROOT) as environment:
        test_catalog_export_response_normalizer_contract()
        test_catalog_export_response_provenance_contract()
        test_single_validated_catalog_export_outcome_contract()
        await test_catalog_export_response_normalization_across_apis(environment)
        await test_unsafe_valid_artifact_fails_all_catalog_consumers(environment)
        await test_canonical_payload_contract(environment)
        await test_pre_r1_artifact_availability_contract(environment)
        await test_material_predicate_python_mysql_type_equivalence(environment)
        await test_export_center_shared_effective_projection_contract(environment)
        await test_invalid_step_payload_falls_back_across_apis(environment)
        test_worker_outcome_contract_types()
        await test_progress_update_is_flush_only(environment)
        await test_sync_catalog_export_api_commits_builder_mutations(environment)
        await test_real_worker_late_business_failure_rolls_back_row_mutations(environment)
        await test_effective_legacy_partial_list_api(environment)
        await test_summary_total_parser_and_legacy_projection_contract(environment)
        await test_pre_r1_task_run_remote_download_and_action_contract(environment)
        await test_catalog_export_single_selector_and_consumption_contract(environment)
        await test_real_builder_db_error_is_not_business_partial(environment)
        await test_real_worker_upload_system_error_rolls_back_builder_state(environment)
        await test_real_worker_projection_faults_rollback_catalog_and_progress(environment)
        await test_real_mixed_worker_scheduler_outcome(environment)
        await test_real_done_and_business_failed_outcomes(environment)
        await test_catalog_export_download_and_export_center_api(environment)
        await test_pre_r1_offline_task_export_and_executor_reuse(environment)
        await test_catalog_export_zip_recovery_preserves_outcome(environment)
        await test_catalog_export_outcome_rejects_invalid_topology(environment)
        await test_catalog_export_terminal_projection_rolls_back_atomically(environment)
        print(f"R1 catalog export canonical payload checks passed: {environment.database_name}")


if __name__ == "__main__":
    asyncio.run(main())
